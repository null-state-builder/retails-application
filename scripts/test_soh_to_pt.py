"""The SOH-only PT builder: one upload-ready file with no invented bill and a reasoned HSN."""

from __future__ import annotations

import shutil
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from scripts.soh_to_pt import convert, gender_group, main
from scripts.test_normalize_pt_files import book

SOH_HEADER = ["Item Name", "Brand", "Size", "Supplier", "Barcode", "Design No", "Category", "Gender", "Fit", "Season",
              "Sale", "Tqty", "Mrp", "Rate", "Amount"]
PT_HEADER = ["Bill Date", "BillNo", "Item", "Brand", "Gender", "Barcode", "HSN", "Qty", "P Rate", "MRP"]


def soh(barcode: Any, qty: int, *, item: str = "SHIRT", brand: str = "KILLER", gender: str = "MALE", mrp: Any = 1999,
        rate: Any = 999.5, season: Any = "SS25", design: Any = "D1") -> list[Any]:
    return [item, brand, "M", "KDPS", barcode, design, "CASUAL WEAR", gender, "SLIM", season, -1, qty, mrp, rate,
            None if rate is None else rate * qty]


class SohToPtTests(unittest.TestCase):
    def setUp(self) -> None:
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir)
        self.pt = self.dir / "pt.xlsx"
        book(self.pt, PT_HEADER, [[None, "S-1", "SHIRT", "KILLER", "MALE", "8909000000001", "62052090", 1, 900, 1999],
                                  [None, "S-1", "JEANS", "KILLER", "MALE", "8909000000009", "62034290", 1, 900, 1999]])
        self.soh = self.dir / "soh.xlsx"
        book(self.soh, SOH_HEADER, [
            soh(8909000000001, 2),  # the same barcode is in the PT
            soh(8909000000002, 1, item="JEANS"),  # same brand, item and gender
            soh(8909000000003, 1, item="JEANS", gender="FEMALE"),  # men's evidence never reaches women's jeans
            soh(8909000000004, 1, item="KURTI SET", gender=".", design=datetime(2026, 5, 1)),
            soh(8909000000005, 1, item="KURTI SET", gender="FEAMLE", season=datetime(2025, 6, 27)),
            soh(8909000000006, 0), soh(8909000000007, -2), soh(8909000000008, 1, mrp=0), soh(8909000000010, 1, rate=2500),
            [None] * 10 + [-2, 8, None, None, 5000],  # the totals row
        ])

    def test_rows_hsn_and_exclusions(self) -> None:
        result = convert(self.soh, [self.pt])
        rows = {r["Barcode"]: r for r in result.rows}
        self.assertEqual(sorted(rows), [f"890900000000{n}" for n in range(1, 6)])
        self.assertEqual(result.stats["left_out"], {"no stock (Tqty 0)": 1, "negative Tqty": 1, "no MRP": 1,
                                                    "Rate above MRP": 1})
        self.assertEqual(rows["8909000000001"]["HSN"], "62052090")
        self.assertEqual(rows["8909000000002"]["HSN"], "62034290")
        self.assertEqual(rows["8909000000003"]["HSN"], "6204")
        self.assertEqual(rows["8909000000004"]["HSN"], "6204")  # "." takes the item's usual gender
        first = rows["8909000000001"]
        self.assertEqual((first["BillNo"], first["Bill Date"]), (None, None))  # no invented bill
        self.assertEqual((first["Qty"], first["P Rate"], first["MRP"], first["Season"]), (2, 999.5, 1999, "SS-25"))
        self.assertEqual(rows["8909000000004"]["Design"], "2026-05-01")
        self.assertEqual((rows["8909000000005"]["Season"], rows["8909000000005"]["Legacy Season"]), ("SS-25", "2025-06-27"))
        self.assertIn("  sale        totals             -2  rows             -9   <- differs", result.report)

    def test_gender_groups(self) -> None:
        cases = {"MALE": "M", "KIDS MALE": "M", "KIDSM": "M", "KIDEM": "M", "FEMALE": "F", "FMAMLE": "F",
                 "FEMALE KIDS": "F", "KIDSF": "F", "UNISEX": None, ".": None, "MIX": None, None: None}
        self.assertEqual({raw: gender_group(raw) for raw in cases}, cases)

    def test_the_written_file_passes_the_pt_upload_reader(self) -> None:
        out = self.dir / "out"
        self.assertEqual(main(["--soh", str(self.soh), "--hsn-from", str(self.pt), "--out", str(out), "--name", "x.xlsx"]), 0)
        sheet = load_workbook(out / "x.xlsx").active
        assert sheet is not None
        barcode = [c.value for c in sheet[1]].index("Barcode")
        self.assertTrue(all(isinstance(row[barcode], str) for row in sheet.iter_rows(min_row=2, values_only=True)))
        self.assertIn("problems 0, price conflicts 0, unreadable 0", (out / "normalize-report.txt").read_text())

    def test_an_unmapped_item_fails_the_run(self) -> None:
        book(self.soh, SOH_HEADER, [soh(8909000000011, 1, item="MYSTERY THING", gender="UNISEX")])
        out = self.dir / "out"
        self.assertEqual(main(["--soh", str(self.soh), "--out", str(out)]), 1)
        self.assertIn("8909000000011 (row 2) MYSTERY THING", (out / "normalize-report.txt").read_text())


if __name__ == "__main__":
    unittest.main()
