"""Combine a test store's PT workbooks into one normalised PT file (store test run).

The PT upload at /goods/opening takes at most 40 files, and the store's PT
files carry no season and disagree with the legacy SOH in places. This reads
every PT workbook in a folder plus the store's SOH and writes one workbook the
upload reads unchanged, and a report of every change. The inputs are only read.

The SOH is the truth for each barcode's description, season and selling price;
the PT supplies quantity, cost (P RATE), HSN and tax:

* a bill already read from another file is dropped (the same bill saved twice);
* Item, Brand, Size, Design, Sub Category, Gender and Fit come from the SOH
  (a blank or ``0`` SOH value keeps the PT's); brands are not merged;
* Season is the SOH's, as ``SS-YY`` / ``AW-YY`` (``ptmapper.legacy_season``),
  with the legacy text kept in ``Legacy Season``;
* MRP is the SOH's (the PT's where the SOH's is below cost); P RATE is the
  barcode's latest bill's, on every line;
* an HSN that is not 4, 6 or 8 digits takes the HSN most used by the same brand
  and item in these files (else the same item);
* barcodes are written as text.

    backend/.venv/bin/python scripts/normalize_pt_files.py --pt "docs/wardrobe-ratu-road/PT FILE" \\
        --soh docs/wardrobe-ratu-road/SOH/SOH_WARDROBE_28-10-2026.xlsx --out docs/wardrobe-ratu-road/normalized

The outputs hold cost data: keep them out of git.
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "backend") not in sys.path:
    sys.path.insert(0, str(ROOT / "backend"))

from ptmapper.legacy_season import normalize_season

HEADER = ["Bill Date", "BillNo", "Item", "Brand", "Size", "Design", "Sub Category", "Gender", "Fit", "Season",
          "Legacy Season", "Barcode", "HSN", "Qty", "Basic", "P Rate", "MRP", "Input Tax", "Output Tax", "Bill Amount",
          "Source File"]
#: PT column → its SOH column, for the description taken from the SOH.
FROM_SOH = {"Item": "itemname", "Brand": "brand", "Size": "size", "Design": "designno", "Sub Category": "category",
            "Gender": "gender", "Fit": "fit"}
PT_NEEDS = {"billdate", "billno", "barcode", "qty", "prate", "mrp", "hsn"}
SOH_NEEDS = {"barcode", "season", "mrp"}
EXAMPLES = 15


def _key(value: Any) -> str:
    return "" if value is None else str(value).strip().casefold().replace(" ", "")


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, Decimal) and value == value.to_integral_value():
        return str(int(value))
    return str(value).strip()


def number(value: Any) -> Decimal | None:
    raw = text(value).replace(",", "")
    if not raw:
        return None
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


def valid_hsn(value: str) -> bool:
    return value.isdigit() and len(value) in {4, 6, 8}


def _table(path: Path, needs: set[str]) -> list[dict[str, Any]]:
    """Rows of the first sheet carrying ``needs`` in its header, as {normalised header: value}."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            header: list[str] | None = None
            rows: list[dict[str, Any]] = []
            for number_, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                if header is None:
                    names = [_key(v) for v in values]
                    if needs <= set(names):
                        header = names
                    elif number_ > 20:
                        break
                    continue
                cells = {name: values[i] for i, name in enumerate(header) if name and i < len(values)}
                if any(text(v) for v in cells.values()):
                    rows.append({**cells, "_row": number_})
            if header is not None:
                return rows
    finally:
        workbook.close()
    raise SystemExit(f"{path.name}: no sheet has the columns {', '.join(sorted(needs))}.")


@dataclass
class Result:
    rows: list[dict[str, Any]]
    report: list[str] = field(default_factory=list)
    stats: dict[str, Any] = field(default_factory=dict)


def normalize(pt_dir: Path, soh_path: Path) -> Result:
    files = sorted(p for p in pt_dir.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        raise SystemExit(f"No .xlsx files in {pt_dir}.")
    soh: dict[str, dict[str, Any]] = {}
    for row in _table(soh_path, SOH_NEEDS):
        barcode = text(row.get("barcode"))
        if barcode:
            soh[barcode] = row

    # 1. Read every PT line; a bill already read from another file is the same bill saved twice.
    lines: list[dict[str, Any]] = []
    bill_file: dict[str, str] = {}
    dropped: dict[str, Counter[str]] = defaultdict(Counter)
    for path in files:
        for row in _table(path, PT_NEEDS):
            bill = text(row.get("billno"))
            first = bill_file.setdefault(bill, path.name)
            if first != path.name:
                dropped[path.name][f"{bill} (already in {first})"] += 1
                continue
            barcode = text(row.get("barcode"))
            if not barcode and not text(row.get("qty")):
                continue
            lines.append({"pt": row, "file": path.name, "barcode": barcode})

    def when(line: dict[str, Any]) -> tuple[Any, ...]:
        moment = line["pt"].get("billdate")
        moment = moment.date() if isinstance(moment, datetime) else moment
        return (moment if isinstance(moment, date) else date.min, text(line["pt"].get("billno")), line["file"],
                line["pt"]["_row"])

    lines.sort(key=when)

    # 2. Per barcode: the latest bill's P RATE; HSN most used by brand+item, else item.
    latest_rate: dict[str, Any] = {}
    rates: dict[str, set[str]] = defaultdict(set)
    for line in lines:
        rate = line["pt"].get("prate")
        latest_rate[line["barcode"]] = rate
        rates[line["barcode"]].add(text(rate))

    def described(line: dict[str, Any], column: str) -> str:
        source = soh.get(line["barcode"])
        pt_value = text(line["pt"].get(_key(column)))
        if source is None:
            return pt_value
        value = text(source.get(FROM_SOH[column]))
        return pt_value if value in {"", "0"} else value

    hsn_by: dict[tuple[str, ...], Counter[str]] = defaultdict(Counter)
    for line in lines:
        hsn = text(line["pt"].get("hsn"))
        if valid_hsn(hsn):
            brand, item = described(line, "Brand"), described(line, "Item")
            hsn_by[(brand, item)][hsn] += 1
            hsn_by[(item,)][hsn] += 1

    # 3. Build the output rows.
    out: list[dict[str, Any]] = []
    changed: dict[str, list[str]] = defaultdict(list)
    seasons: dict[tuple[str, str], set[str]] = defaultdict(set)
    hsn_fixes: list[str] = []
    for line in lines:
        pt, barcode = line["pt"], line["barcode"]
        source = soh.get(barcode)
        row: dict[str, Any] = {"Bill Date": pt.get("billdate"), "BillNo": text(pt.get("billno")), "Barcode": barcode,
                               "Qty": pt.get("qty"), "Basic": pt.get("basic"), "Input Tax": pt.get("inputtax"),
                               "Output Tax": pt.get("outputtax"), "Bill Amount": pt.get("billamount"),
                               "Source File": f"{line['file']} row {pt['_row']}"}
        for column in FROM_SOH:
            row[column] = described(line, column)
            before = text(pt.get(_key(column)))
            if row[column] != before:
                changed[column].append(f"{barcode}: {before!r} → {row[column]!r}")
        legacy = text(source.get("season")) if source else ""
        row["Legacy Season"] = legacy
        row["Season"] = normalize_season(legacy) or ""
        seasons[(legacy, row["Season"] or "unknown historical")].add(barcode)
        row["P Rate"] = latest_rate[barcode]
        mrp = number(source.get("mrp")) if source else None
        cost = number(row["P Rate"])
        if mrp is not None and cost is not None and mrp < cost:
            changed["MRP kept from the PT (the SOH MRP is below cost)"].append(
                f"{barcode}: SOH {text(mrp)} < P RATE {text(cost)}; kept {text(pt.get('mrp'))}")
            mrp = None
        row["MRP"] = mrp if mrp is not None else pt.get("mrp")
        if number(pt.get("mrp")) != number(row["MRP"]):
            changed["MRP"].append(f"{barcode}: {text(pt.get('mrp'))} → {text(row['MRP'])}")
        if len(rates[barcode]) > 1 and text(pt.get("prate")) != text(row["P Rate"]):
            changed["P Rate"].append(f"{barcode} ({line['file']}): {text(pt.get('prate'))} → {text(row['P Rate'])}")
        hsn = text(pt.get("hsn"))
        if not valid_hsn(hsn):
            options = hsn_by.get((row["Brand"], row["Item"])) or hsn_by.get((row["Item"],))
            if options:
                fixed = options.most_common(1)[0][0]
                hsn_fixes.append(f"{barcode} {row['Brand']} {row['Item']} ({line['file']}): {hsn} → {fixed}")
                hsn = fixed
        row["HSN"] = hsn
        out.append(row)

    # 4. Report.
    barcodes = {row["Barcode"] for row in out}
    pieces = sum(int(number(row["Qty"]) or 0) for row in out)
    report = [f"PT files read: {len(files)}; lines kept: {len(out)}; barcodes: {len(barcodes)}; pieces: {pieces}",
              f"SOH barcodes: {len(soh)}; PT barcodes not in the SOH: {len(barcodes - set(soh))}", ""]
    report.append(f"Duplicate bills dropped ({sum(sum(c.values()) for c in dropped.values())} lines):")
    report += [f"  {name}: {bill}, {count} lines" for name, bills in dropped.items() for bill, count in bills.items()]
    report += ["", f"HSN fixed ({len(hsn_fixes)}):", *(f"  {fix}" for fix in hsn_fixes)]
    for column in ["MRP", "MRP kept from the PT (the SOH MRP is below cost)", "P Rate", *FROM_SOH]:
        items = changed.get(column, [])
        more = f" (first {EXAMPLES})" if len(items) > EXAMPLES else ""
        report += ["", f"{column} changed on {len(items)} lines{more}:"]
        report += [f"  {item}" for item in items[:EXAMPLES]]
    report += ["", "Seasons (legacy text → season: barcodes):"]
    report += [f"  {legacy or '(blank)'!s:32} → {season}: {len(codes)}"
               for (legacy, season), codes in sorted(seasons.items(), key=lambda kv: (kv[0][1], kv[0][0]))]
    per_season: Counter[str] = Counter()
    for row in out:
        per_season[row["Season"] or "unknown historical"] += int(number(row["Qty"]) or 0)
    report += ["", "Pieces per season:", *(f"  {season}: {qty}" for season, qty in sorted(per_season.items()))]
    if barcodes - set(soh):
        report += ["", "PT barcodes not in the SOH (kept with their PT values):",
                   *(f"  {b}" for b in sorted(barcodes - set(soh)))]
    stats = {"files": len(files), "lines": len(out), "barcodes": len(barcodes), "pieces": pieces,
             "dropped_lines": sum(sum(c.values()) for c in dropped.values()), "hsn_fixes": len(hsn_fixes),
             "seasons": dict(per_season)}
    return Result(out, report, stats)


def write_workbook(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PT"
    sheet.append(HEADER)
    for row in rows:
        sheet.append([row.get(column) for column in HEADER])
    barcode = HEADER.index("Barcode") + 1
    for cells in sheet.iter_rows(min_row=2, min_col=barcode, max_col=barcode):
        cells[0].number_format = "@"
    workbook.save(path)


def self_check(path: Path) -> dict[str, Any]:
    """What the PT upload's own reader makes of the file."""
    from ptmapper.pt_opening import parse_pt_files

    metadata, rows, problems = parse_pt_files([(path.name, path.read_bytes())])
    opening = metadata["pt_opening"]
    return {"barcodes": len(rows), "pieces": metadata["quantity"], "pt_rows": opening["pt_rows"],
            "problems": len(problems), "conflicts": opening["conflict_count"],
            "unreadable": opening["unreadable_count"], "problem_list": opening["problems"][:20]}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--pt", type=Path, required=True, help="folder of PT .xlsx files")
    parser.add_argument("--soh", type=Path, required=True, help="the store's legacy SOH .xlsx")
    parser.add_argument("--out", type=Path, required=True, help="folder for the combined file and report")
    parser.add_argument("--name", default="WARDROBE_PT_combined.xlsx")
    args = parser.parse_args(argv)
    result = normalize(args.pt, args.soh)
    args.out.mkdir(parents=True, exist_ok=True)
    target = args.out / args.name
    write_workbook(result.rows, target)
    check = self_check(target)
    lines = [*result.report, "", "Self-check with the PT upload's reader:",
             (f"  barcodes {check['barcodes']}, pieces {check['pieces']}, rows {check['pt_rows']}, "
             f"problems {check['problems']}, price conflicts {check['conflicts']}, unreadable {check['unreadable']}"),
             *(f"  problem: {p['barcode']} {p['reason']}" for p in check["problem_list"])]
    (args.out / "normalize-report.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines[:3] + lines[-2 - len(check["problem_list"]):]))
    print(f"Wrote {target} and {args.out / 'normalize-report.txt'}")
    return 1 if check["problems"] or check["unreadable"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
