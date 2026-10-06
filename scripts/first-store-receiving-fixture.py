"""Explicit fictional inputs and read-only reconciliation for the receiving sibling.

No accepted inventory, origin, journal or business decision is seeded here.
All receipt, damage and SOH state changes belong to normal browser commands.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import secrets
import stat
import uuid
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
RECORD = ROOT / ".local/first-store-receiving-credentials.json"
ARTIFACTS = ROOT / ".local/first-store-receiving-browser-inputs"


def initialise() -> tuple[Any, Any, dict[str, Any]]:
    if (os.environ.get("KDPS_PROOF_MODE") != "1"
            or not os.environ.get("KDPS_REHEARSAL_DB", "").startswith("kdps_rehearsal_receiving_")):
        raise RuntimeError("Use the owned receiving proof wrapper.")
    if stat.S_IMODE(RECORD.stat().st_mode) != 0o600:
        raise RuntimeError("The receiving proof credential record must be private.")
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
    import django
    django.setup()
    from accounts.registration_models import InstallationRegistration
    from accounts.registration_services import deployment_key
    from django.db import connection

    record = json.loads(RECORD.read_text())
    expected = record["receiving_rehearsal"]
    with connection.cursor() as cursor:
        cursor.execute("SELECT current_database(), current_user, system_identifier::text FROM pg_control_system()")
        identity = cursor.fetchone()
    if identity != (expected["database"], "kdps_proof", expected["system_identifier"]):
        raise RuntimeError("The receiving database identity changed; nothing was run.")
    claim = InstallationRegistration.objects.select_related("tenant").get(deployment_key=deployment_key(), completed_at__isnull=False)
    if claim.tenant is None or not claim.tenant.synthetic or claim.tenant.code != "ALPHA":
        raise RuntimeError("Only the synthetic ALPHA sibling is permitted.")
    spec = importlib.util.spec_from_file_location("receiving_api_helper", ROOT / "scripts/first-store-onboarding-proof.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("The canonical API helper is unavailable.")
    proof = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(proof)
    proof.RECORD = RECORD
    for who in ("owner", "admin", "manager", "count_checker"):
        if not record[who]["email"].endswith(".example.test"):
            raise RuntimeError("Only generated fictional proof people are permitted.")
    return proof, claim, record


def snapshot(claim: Any, record: dict[str, Any]) -> dict[str, Any]:
    from core.canonical import content_hash, normalise
    from core.goods_fields import bounds
    from finledger.models import CashLedgerEntry
    from inbound.goods_models import Arrival, CountSession, GoodsGrn
    from masters.goods_models import SiteGuard
    from outbound.goods_models import DamageReport
    from outbound.goods_soh_models import SohReconciliation
    from ptmapper.soh_models import SohImport
    from sell.models import Sale, SaleTender
    from sell.services.goods_stock import read_shelf
    from stockledger.goods_models import JournalBatch, Origin, Position, ValueLeg

    positions = [{"lot_id": str(row.lot_id), "origin_id": str(row.origin_id) if row.origin_id else None,
                  "sku_id": str(row.sku_id) if row.sku_id else None, "qty": bounds(row.portion)[1] - bounds(row.portion)[0],
                  "condition": row.condition, "boundary": row.boundary, "location_id": str(row.location_id),
                  "unit_cost_paise": str(row.origin.unit_cost) if row.origin else None}
                 for row in Position.objects.filter(site_id=claim.first_store_id).select_related("origin").order_by("lot_id", "portion")]
    sales = list(Sale.objects.values("id", "doc_number", "net_paise", "idempotency_uuid").order_by("pk"))
    guard = SiteGuard.objects.get(site_id=claim.first_store_id)
    data = {"synthetic_only": True, "rehearsal_identity": record["receiving_rehearsal"],
            "site_id": claim.first_store_id,
            "sellable_qty": sum(read_shelf(guard.site).quantities.values()), "positions": positions,
            "stock_cost_paise": str(sum(row["qty"] * int(row["unit_cost_paise"] or 0)
                                         for row in positions if row["boundary"] == "physical")),
            "origins": list(Origin.objects.values("id", "opening_qty", "unit_cost", "mrp").order_by("pk")),
            "journals": list(JournalBatch.objects.values("id", "posting_kind", "event_at").order_by("pk")),
            "value_legs": list(ValueLeg.objects.values("id", "amount").order_by("pk")),
            "arrivals": Arrival.objects.count(), "counts": CountSession.objects.count(), "grns": GoodsGrn.objects.count(),
            "damage_reports": list(DamageReport.objects.values("id", "state", "quantity", "reporter_id", "reviewer_id").order_by("pk")),
            "soh_sources": list(SohImport.objects.values("id", "state", "source_hash", "revision").order_by("pk")),
            "soh_counts": list(SohReconciliation.objects.values("id", "state", "journal_batch_id").order_by("pk")),
            "freeze_id": str(guard.freeze_id) if guard.freeze_id else None,
            "sales": sales, "sale_hash": content_hash(normalise(sales)),
            "tenders": list(SaleTender.objects.values("id", "amount_paise").order_by("pk")),
            "cash_hash": content_hash(normalise(list(CashLedgerEntry.objects.values().order_by("pk"))))}
    return normalise(data)


def prepare(proof: Any, claim: Any, record: dict[str, Any]) -> dict[str, Any]:
    from accounts.goods_models import HumanIdentity, RoleAssignment
    from accounts.models import Role, User
    from approvals.goods_models import ApprovalRequest
    from core.commands import database_now
    from core.numbering import prepare_series
    from masters.goods_models import ConfigVersion, Location
    from masters.models import Store
    from stockledger.goods_models import Origin

    def actor(who: str) -> Any:
        client = proof.login(record, who)
        proof.request(client, "post", "/api/auth/step-up", {"password": record[who]["password"]})
        return client

    def publish(key: str, kind: str, payload: dict[str, Any], scope: dict[str, Any]) -> ConfigVersion:
        admin = actor("admin")
        effective_from = record.setdefault("receiving_fixture", {}).setdefault("effective_from", database_now().isoformat())
        proof.save(record)
        draft = proof.request(admin, "post", "/api/goods-v1/masters/configurations", proof.command(record,
            f"receiving-{key}-create", kind=kind, scope=scope, effective_from=effective_from, payload=payload))
        existing = ConfigVersion.objects.filter(tenant_id=claim.tenant_id, draft_id=draft["id"]).first()
        if existing is not None:
            return existing
        proof.request(admin, "post", f"/api/goods-v1/masters/configurations/{draft['id']}/submit", proof.command(record,
            f"receiving-{key}-submit", expected_revision=draft["revision"], reviewed_hash=draft["content_hash"]))
        approval = ApprovalRequest.objects.get(subject_kind="configuration", subject_key=f"configdraft:{draft['id']}")
        proof.request(actor("owner"), "post", f"/api/goods-v1/approvals/{approval.pk}/decide", proof.command(record,
            f"receiving-{key}-approve", expected_revision=approval.revision, reviewed_hash=approval.reviewed_hash, decision="approve"))
        return ConfigVersion.objects.get(tenant_id=claim.tenant_id, draft_id=draft["id"])

    fixture = record.setdefault("receiving_fixture", {})
    person = record.setdefault("warehouse", {"email": f"warehouse-{uuid.uuid4().hex[:12]}@receiving.example.test", "password": secrets.token_urlsafe(24) + "A1!"})
    proof.save(record)
    user = User.objects.filter(tenant_id=claim.tenant_id, email=person["email"]).first()
    if user is None:
        human = HumanIdentity.objects.create(tenant_id=claim.tenant_id, staff_code="RECEIVING-PROOF", display_name="Fictional warehouse preparer")
        user = User.objects.create_user(username=person["email"], email=person["email"], password=person["password"],
            tenant_id=claim.tenant_id, human=human, full_name=human.display_name, must_change_password=False)
        for code in ("warehouse", "it_admin"):
            role = Role.objects.get(tenant_id=claim.tenant_id, code=code)
            RoleAssignment.objects.create(tenant_id=claim.tenant_id, human=human, role=role,
                all_sites=False, site_ids=[claim.first_store_id], all_brands=True, effective_from=database_now())
    if not user.is_active or not user.check_password(person["password"]):
        raise RuntimeError("The generated warehouse proof identity changed; preserved.")
    owner = proof.login(record, "owner")
    proof.request(owner, "post", "/api/auth/step-up", {"password": record["owner"]["password"]})
    if not fixture.get("vendor_id"):
        vendor = proof.request(owner, "post", "/api/goods-v1/vendors", proof.command(record, "receiving-vendor", code="RECEIVING-PROOF", name="Fictional receiving proof supplier"))
        fixture["vendor_id"] = str(vendor["id"])
    site = Store.objects.get(tenant_id=claim.tenant_id, pk=claim.first_store_id, code="FIRST")
    for kind in ("RPT", "GRN", "HLD", "REL", "CNT"):
        prepare_series(claim.tenant_id, site.gstin.legal_entity, kind)
    origin = Origin.objects.filter(sku__aliases__value="ALPHA000123").first()
    if origin is None:
        raise RuntimeError("The browser's accepted opening origin is absent.")
    fixture.update(synthetic_only=True, site_id=str(site.pk), brand_id=str(origin.sku.style.brand_id), sku_id=str(origin.sku_id),
        origin_id=str(origin.pk), floor_id=str(Location.objects.get(site=site, kind="floor").pk), barcode="ALPHA000123")
    publish("shortage-policy-v2", "approval", {"action": "receipt.disposition.decide", "roles": ["owner"], "purpose": "accept_shortage",
        "require_distinct": True, "site_ids": [], "brand_ids": [], "qty_max": 100, "step_up": True,
        "unknown_value": "quantity_only"}, {"scope_kind": "tenant", "purposes": ["accept_shortage"]})
    proof.save(record)
    return {**fixture, "snapshot": snapshot(claim, record)}


def supplier_identity(proof: Any, record: dict[str, Any]) -> dict[str, Any]:
    """Append the explicit supplier code after the full-store source checks.

    The full-store reader deliberately refuses multiple active aliases, even
    when they name the same SKU. Do not add this before testing that contract.
    """
    from core.commands import database_now

    fixture = record["receiving_fixture"]
    effective_from = fixture.setdefault("supplier_identity_effective_from", database_now().isoformat())
    proof.save(record)
    owner = proof.login(record, "owner")
    proof.request(owner, "post", "/api/auth/step-up", {"password": record["owner"]["password"]})
    alias = proof.request(owner, "post", "/api/goods-v1/masters/aliases", proof.command(record, "receiving-supplier-alias",
        sku_id=fixture["sku_id"], issuer_key="RECEIVING-PROOF", alias_type="barcode", value=fixture["barcode"],
        site_id=fixture["site_id"], effective_from=effective_from))
    return {"synthetic_only": True, "alias_id": alias["master"]["id"], "stock_changed": False}


def artifact(claim: Any, record: dict[str, Any], name: str, quantity: int, barcode: str) -> dict[str, Any]:
    from openpyxl import Workbook, load_workbook
    from tests.first_store_goods import source_cutoff

    if not name.replace("-", "").isalnum():
        raise RuntimeError("Use a simple fictional scenario name.")
    if barcode not in ("ALPHA000123", "FICTIONAL-ZERO-CATALOGUE") or (barcode != "ALPHA000123" and quantity != 0):
        raise RuntimeError("Only the accepted fictional identity or an unstocked zero catalogue row is permitted.")
    artifacts = ARTIFACTS / record["receiving_rehearsal"]["database"]
    artifacts.mkdir(parents=True, exist_ok=True)
    path = artifacts / f"{name}.xlsx"
    cutoff = source_cutoff().isoformat()
    if path.exists():
        metadata = json.loads((artifacts / f"{name}.json").read_text())
        prior = load_workbook(path, read_only=True, data_only=True)
        try:
            assert prior.active is not None
            row = next(prior.active.iter_rows(min_row=2, max_row=2, values_only=True))
            if (metadata["quantity"] != quantity or row[3] != barcode or row[6] != quantity
                    or metadata["sha256"] != hashlib.sha256(path.read_bytes()).hexdigest()):
                raise RuntimeError("The prior source artifact changed; preserve it and use a new scenario.")
        finally:
            prior.close()
        return metadata
    book = Workbook()
    assert book.active is not None
    book.active.append(["ItemName", "Brand", "Size", "Barcode", "Category", "Season", "Tqty", "Mrp", "Rate", "Amount"])
    book.active.append([f"Fictional {name} counted shirt", "Fictional source brand", "M", barcode, "Shirt", "Fictional source season", quantity, 1000, 500, max(quantity, 0) * 500])
    book.save(path)
    evidence = artifacts / f"{name}-valuation.xlsx"
    value_book = Workbook()
    assert value_book.active is not None
    value_book.active.append(["Fictional source reconciliation", "barcode", "qty", "basic INR", "value INR"])
    value_book.active.append([name, barcode, quantity, 500, quantity * 500])
    value_book.save(evidence)
    csv = artifacts / f"{name}.csv"
    csv.write_text(f"barcode,observed_qty,observed_condition,reason\n{barcode},{max(quantity, 0)},good,Generated fictional {name} physical observation only\n")
    metadata = {"synthetic_only": True, "file": str(path), "cutoff_at": cutoff, "quantity": quantity,
                "physical_file": str(csv), "evidence_file": str(evidence), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    (artifacts / f"{name}.json").write_text(json.dumps(metadata, indent=2))
    return metadata


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=["prepare", "supplier-identity", "snapshot", "artifact"])
    parser.add_argument("name", nargs="?", default="")
    parser.add_argument("quantity", nargs="?", type=int, default=1)
    parser.add_argument("barcode", nargs="?", default="ALPHA000123")
    args = parser.parse_args()
    proof, claim, record = initialise()
    from core.tenancy import tenant_context
    with tenant_context(claim.tenant_id):
        if args.action == "prepare":
            result = prepare(proof, claim, record)
        elif args.action == "supplier-identity":
            result = supplier_identity(proof, record)
        elif args.action == "artifact":
            result = artifact(claim, record, args.name, args.quantity, args.barcode)
        else:
            result = snapshot(claim, record)
        print(json.dumps(result, sort_keys=True))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:  # noqa: BLE001 -- never expose fictional credentials in tracebacks.
        print(str(error) if isinstance(error, RuntimeError) else f"Receiving proof failed ({type(error).__name__}); preserved for inspection.")
        raise SystemExit(1) from None
