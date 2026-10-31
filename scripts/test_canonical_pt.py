"""The canonical PT format: one meaning per column, the bill's cost kept, nothing guessed."""

from __future__ import annotations

import shutil
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from scripts.canonical_pt import HEADER, convert, item_of, main, price_tier, read_fit, season_of, size_of
from scripts.normalize_pt_files import HEADER as OLD_HEADER
from scripts.test_normalize_pt_files import book


def old(barcode: str, *, item: str = "SHIRT", category: str = "CASUAL WEAR", gender: str = "MALE",
        fit: str = "SLIM", size: str = "M", season: str = "AUTUMN WINTER(Oct-26)", hsn: str = "62052090",
        qty: int = 1, basic: float | None = 1000.0, prate: float = 1100.0, mrp: int = 1999) -> list[object]:
    row = dict.fromkeys(OLD_HEADER)
    row.update({"Item": item, "Brand": "KILLER", "Size": size, "Design": "D1", "Sub Category": category,
                "Gender": gender, "Fit": fit, "Legacy Season": season, "Barcode": barcode, "HSN": hsn, "Qty": qty,
                "Basic": basic, "P Rate": prate, "MRP": mrp})
    return [row[column] for column in OLD_HEADER]


class CanonicalPtTests(unittest.TestCase):
    def setUp(self) -> None:
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir)
        self.pt = self.dir / "pt.xlsx"
        book(self.pt, OLD_HEADER, [
            old("1"),  # a bill line: BASIC 1000, P Rate 1100 → LANDED 10%
            old("2", item="WINTER KURTI SET", category="SETP", gender="FEAMLE", fit="LM", size="FS", hsn="6204",
                basic=None, prate=1200.0, season="2026-05-24"),  # an SOH line: P RATE only
            old("3", item="BOY SET", category="UTE", gender=".", fit="PE RG Mainline", size="6-12M"),
        ])

    def test_rows(self) -> None:
        rows = {r["BARCODE"]: r for r in convert(self.pt).rows}
        bill = rows["1"]
        self.assertEqual((bill["BASIC"], bill["LANDED %"], bill["P RATE"]), (Decimal("1000.00"), Decimal("10.00"),
                                                                            Decimal("1100.00")))
        self.assertEqual((bill["DIVISION"], bill["CATEGORY"], bill["FIT"], bill["CONSTRUCTION"]),
                         ("Apparel", "Topwear", "Slim", "Woven"))
        self.assertEqual((bill["SEASON"], bill["ARRIVAL MONTH"], bill["OCCASION"]), ("AW-26", "2026-10", "Casual"))
        soh = rows["2"]
        self.assertEqual((soh["BASIC"], soh["LANDED %"], soh["P RATE"]), (Decimal("1000.00"), Decimal("20"),
                                                                          Decimal("1200.00")))
        self.assertEqual((soh["ITEM"], soh["CATEGORY"], soh["OCCASION"], soh["GENDER"], soh["AGE GROUP"]),
                         ("KURTI SET", "Set", "Winter", "Female", "Adult"))
        self.assertEqual((soh["PRICE TIER"], soh["MARGIN BAND"], soh["SIZE"]), ("Premium", "Low", "FREE SIZE"))
        self.assertEqual((soh["CONSTRUCTION"], soh["SEASON"], soh["ARRIVAL MONTH"]), ("", "SS-26", "2026-05"))
        kid = rows["3"]  # BOY SET names its audience; a size in months is an infant's
        self.assertEqual((kid["ITEM"], kid["GENDER"], kid["AGE GROUP"]), ("SET", "Male", "Infant"))
        self.assertEqual((kid["FIT"], kid["STYLE NAME"], kid["PRICE TIER"]), ("Regular", "PE RG Mainline", "Economy"))

    def test_value_readers(self) -> None:
        self.assertEqual(item_of("TROUSERS", "x"), ("TROUSER", {}))
        self.assertEqual(item_of("MYSTERY", "x"), ("", {}))
        self.assertEqual([price_tier(c) for c in ("SAREEE", "KURM", "SAREE", "TOP", "CASUAL WEAR")],
                         ["Economy", "Medium", "", "", ""])
        self.assertEqual([read_fit(f) for f in ("HALF SHIRT", "SILK SAREE", "3 PCE SET", "VHM", "ROVAR", "XYZ")],
                         [{"SLEEVE": "Half"}, {"FABRIC": "Silk"}, {"PIECES": 3}, {"MARGIN BAND": "Very high"},
                          {"STYLE NAME": "ROVER"}, {"unknown": "XYZ"}])
        self.assertEqual([size_of(s) for s in ("4-5Y", "34B", "55CM", "274 X 274 CM", "16*24", "NA", "XXXL")],
                         [("4-5 Y", "Age (years)"), ("34B", "Cup"), ("55 CM", "Length (cm)"),
                          ("274 X 274 CM", "Dimensions"), ("16-24", "Range"), ("", ""), ("3XL", "Alpha")])
        self.assertEqual(season_of("AW67"), ("", ""))  # a legacy typo stays unknown
        self.assertEqual(season_of("AUTUMN WINTER(Jan-26)"), ("AW-25", "2026-01"))

    def test_the_written_file(self) -> None:
        out = self.dir / "out"
        self.assertEqual(main(["--pt", str(self.pt), "--out", str(out), "--name", "x.xlsx"]), 0)
        workbook = load_workbook(out / "x.xlsx")
        self.assertEqual(workbook.sheetnames, ["PT", "Lists", "Source values", "About"])
        self.assertEqual([c.value for c in workbook["PT"][1]], HEADER)
        self.assertIn("problems 0, price conflicts 0, unreadable 0", (out / "x-report.txt").read_text())

    def test_an_unlisted_item_fails_the_run(self) -> None:
        book(self.pt, OLD_HEADER, [old("9", item="MYSTERY")])
        self.assertEqual(main(["--pt", str(self.pt), "--out", str(self.dir / "out")]), 1)


if __name__ == "__main__":
    unittest.main()
