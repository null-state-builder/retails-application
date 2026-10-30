"""The PT normaliser: one upload-ready file, the SOH's description, season and MRP."""

from __future__ import annotations

import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from scripts.normalize_pt_files import main, normalize

PT_HEADER = ["", "Bill Date", "BillNo", "Item", "Brand", "Size", "Design", "Sub Category", "Gender", "Fit", "Barcode",
             "HSN", "Qty", "Basic", "P Rate", "MRP", "Input Tax", "Output Tax", "Bill Amount"]
SOH_HEADER = ["Item Name", "Brand", "Size", "Supplier", "Barcode", "Design No", "Category", "Gender", "Fit", "Season",
              "Tqty", "Mrp", "Rate"]


def pt(day: int, bill: str, barcode: Any, qty: int, *, item: str = "SHIRT", brand: str = "KILLER", hsn: Any = 62052000,
       prate: float = 1376.38, mrp: int = 2199) -> list[Any]:
    return [None, date(2026, 9, day), bill, item, brand, "M", "D1", "CASUAL WEAR", "MALE", "SLIM", barcode, hsn,
            qty, 1310.84, prate, mrp, 5, 5, prate * qty]


def soh(barcode: Any, season: str, mrp: int, *, item: str = "SHIRT", brand: str = "KILLER", fit: str = "SLIM",
        size: Any = "M") -> list[Any]:
    return [item, brand, size, "KDPS", barcode, "D1", "CASUAL WEAR", "MALE", fit, season, 1, mrp, 1376.38]


def book(path: Path, header: list[str], rows: list[list[Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


class NormalizePtFilesTests(unittest.TestCase):
    def setUp(self) -> None:
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir)
        self.pt = self.dir / "pt"
        self.pt.mkdir()
        book(self.pt / "1522.xlsx", PT_HEADER, [pt(21, "S-1522", 8909241888700, 2, hsn=2707122),
                                                pt(21, "S-1522", "SM1217624-13", 1, prate=900.0)])
        book(self.pt / "lp1522.xlsx", PT_HEADER, [pt(21, "S-1522", 8909241888700, 2)])  # the same bill again
        book(self.pt / "1600.xlsx", PT_HEADER, [pt(25, "S-1600", 8909241888700, 1, prate=1300.0),
                                                pt(25, "S-1600", 8909241889999, 1, item="T-SHIRT", hsn=61091000),
                                                pt(25, "S-1600", 8909241887777, 1)])
        self.soh = self.dir / "soh.xlsx"
        book(self.soh, SOH_HEADER, [soh(8909241888700, "AUTUMN WINTER(Jan-26)", 2299, fit="ANKLE"),
                                    soh("SM1217624-13", "SS25", 1999, brand="V KILLER", size=0),
                                    soh(8909241889999, "", 999, item="T-SHIRT")])

    def test_one_file_with_the_soh_truth_and_no_duplicate_bill(self) -> None:
        result = normalize(self.pt, self.soh)
        rows = {(r["Barcode"], r["Source File"].split()[0]): r for r in result.rows}
        self.assertEqual(result.stats["dropped_lines"], 1)  # lp1522.xlsx repeats bill S-1522
        self.assertNotIn(("8909241888700", "lp1522.xlsx"), rows)
        first = rows[("8909241888700", "1522.xlsx")]
        self.assertEqual(first["Barcode"], "8909241888700")  # a number barcode becomes text
        self.assertEqual((first["Season"], first["Legacy Season"]), ("AW-25", "AUTUMN WINTER(Jan-26)"))
        self.assertEqual((first["Fit"], str(first["MRP"])), ("ANKLE", "2299"))  # description and MRP from the SOH
        self.assertEqual(first["P Rate"], 1300.0)  # the latest bill's cost on every line
        self.assertEqual(first["HSN"], "62052000")  # the 7-digit HSN takes the brand and item's
        other = rows[("SM1217624-13", "1522.xlsx")]
        self.assertEqual((other["Brand"], other["Size"], other["Season"]), ("V KILLER", "M", "SS-25"))
        self.assertEqual(rows[("8909241889999", "1600.xlsx")]["Season"], "")  # blank → unknown historical
        self.assertEqual(rows[("8909241887777", "1600.xlsx")]["MRP"], 2199)  # not in the SOH: the PT's own

    def test_the_written_file_passes_the_pt_upload_reader(self) -> None:
        out = self.dir / "out"
        self.assertEqual(main(["--pt", str(self.pt), "--soh", str(self.soh), "--out", str(out), "--name", "x.xlsx"]), 0)
        sheet = load_workbook(out / "x.xlsx").active
        assert sheet is not None
        header = [c.value for c in sheet[1]]
        barcode = header.index("Barcode")
        self.assertTrue(all(isinstance(row[barcode], str) for row in sheet.iter_rows(min_row=2, values_only=True)))
        report = (out / "normalize-report.txt").read_text()
        self.assertIn("problems 0, price conflicts 0, unreadable 0", report)
        self.assertIn("lp1522.xlsx: S-1522 (already in 1522.xlsx), 1 lines", report)


if __name__ == "__main__":
    unittest.main()
