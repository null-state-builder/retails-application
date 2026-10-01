"""The KDPS master sheet: read it, and work out what importing it would change.

KDPS keeps its product lists in the first sheet ("Master Sheet") of the PT file
workbook: SEASON, BRAND, COLOR, GENDER, SUB CATEGORY, TYPE, ITEM, FIT and SIZE are
independent lists (columns A-I), GST % (J) is a reference list, and K/L carry each
ITEM's suggested SUB CATEGORY and TYPE (row-aligned with ITEM).

This module is pure: it reads bytes and plain snapshots and returns plain JSON. The
import package (``masters.master_sheet_services``) stores both, and its approval
writes them through the product's own writers. The synthetic rulebook seed
(``ptmapper.goods_rulebook_seed``) reads the same sheet through
:func:`read_master_sheets`.
"""

from __future__ import annotations

import io
import re
import unicodedata
from collections.abc import Iterable
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl

from core.refusals import Refusal
from masters.goods_identity_services import normalise_text, vocabulary_value_id

SHEET_NAME = "Master Sheet"

#: Master sheet column -> goods-v1 vocabulary dimension. BRAND (column 1) names
#: brand masters, not a vocabulary; GST % (column 9) is reference only.
DIM_COLS = {
    0: "season",
    2: "colour",
    3: "gender",
    4: "sub_category",
    5: "type",
    6: "item",
    7: "fit",
    8: "size",
}
#: Columns 10 and 11 of the master sheet: ITEM's suggested SUB CATEGORY and TYPE.
HELPER_COLS = {10: "sub_category", 11: "type"}
BRAND_COL = 1
GST_COL = 9
DIMENSIONS = tuple(DIM_COLS.values())
#: The header each column must carry, by position (compared without case or spacing).
HEADERS = (
    "SEASON",
    "BRAND",
    "COLOR",
    "GENDER",
    "SUB CATEGORY",
    "TYPE",
    "ITEM",
    "FIT",
    "SIZE",
    "GST %",
    "SUB CATEGORY",
    "TYPE",
)
#: The sheet's own heading for each list.
COLUMN_LABELS = {
    "season": "SEASON",
    "brand": "BRAND",
    "colour": "COLOR",
    "gender": "GENDER",
    "sub_category": "SUB CATEGORY",
    "type": "TYPE",
    "item": "ITEM",
    "fit": "FIT",
    "size": "SIZE",
}
#: The longest value key/label a vocabulary carries, and a brand master's name.
MAX_VALUE = 100
MAX_BRAND = 120
MAX_BRAND_CODE = 32
MAX_SEASON_CODE = 24
MAX_ROWS = 5000

GST_NOTE = "Tax comes from Tax settings by HSN; the GST % column is not imported."

# Problem codes. A blocking problem leaves only that value out; nothing else stops.
TOO_LONG = "TOO_LONG"
DUPLICATE_IN_SHEET = "DUPLICATE_IN_SHEET"
WHITESPACE_FIXED = "WHITESPACE_FIXED"
LOOKS_LIKE = "LOOKS_LIKE"
HELPER_UNKNOWN = "HELPER_UNKNOWN"
HELPER_MULTIPLE = "HELPER_MULTIPLE"

# Warnings about other configuration; each must be acknowledged before submitting.
PENDING_DRAFT = "PENDING_DRAFT"
PINNED_PROFILE = "PINNED_PROFILE"
ALLOW_LIST = "ALLOW_LIST"


def _sheet_invalid(message: str) -> Refusal:
    return Refusal("MASTER_SHEET_INVALID", message, status=422)


# ----------------------------------------------------------------------------- text


def cell_text(val: Any) -> str:
    """A cell as the text a person reads: ``28.0`` is ``28``, NBSPs are spaces."""
    if val is None:
        return ""
    if isinstance(val, bool):
        return str(val).upper()
    if isinstance(val, float | Decimal):
        number = Decimal(str(val))
        if number == number.to_integral_value():
            return str(int(number))
        return format(number.normalize(), "f")
    if isinstance(val, int):
        return str(val)
    return " ".join(unicodedata.normalize("NFC", str(val)).split())


def match_key(text: str) -> str:
    """The form two texts are compared in: NFC, whitespace collapsed, case folded."""
    return normalise_text(text)


def compact(text: str) -> str:
    """Letters and digits only, case folded: ``TOM BOY`` and ``TOMBOY`` agree."""
    folded = unicodedata.normalize("NFKC", text).casefold()
    return "".join(ch for ch in folded if ch.isalnum())


def slug(text: str, limit: int) -> str:
    """A lower-case code of letters, digits and hyphens, at most ``limit`` long."""
    folded = (
        unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    )
    return re.sub(r"[^a-z0-9]+", "-", folded).strip("-")[:limit].strip("-")


def pick_valid(cell: str, valid: set[str]) -> str:
    """From a helper cell that may hold several '/'-separated options
    ('FORMAL / CASUAL/ PARTY WEAR'), the first that is a real value (exact, else a
    value the token starts: 'CASUAL' -> 'CASUAL WEAR')."""
    options = helper_options(cell, valid)
    return options[0] if options else ""


def helper_options(cell: str, valid: Iterable[str]) -> list[str]:
    """Every real value a helper cell names, in the cell's order.

    A token matches a value exactly (without case), else the first value it starts
    ('CASUAL' -> 'CASUAL WEAR'): ``'LUGGAGE/ACCESSORIES'`` -> ``['LUGGAGE', 'ACCESSORIES']``.
    """
    ordered = sorted(valid)  # sorted: deterministic when a token prefixes several
    by_key = {match_key(value): value for value in ordered}
    out: list[str] = []
    for token in re.split(r"[/,]", cell or ""):
        key = match_key(token)
        if not key:
            continue
        found = by_key.get(key) or next(
            (value for value in ordered if match_key(value).startswith(key)), ""
        )
        if found and found not in out:
            out.append(found)
    return out


_MONTHS = {
    m: i
    for i, m in enumerate(
        [
            "JAN",
            "FEB",
            "MAR",
            "APR",
            "MAY",
            "JUN",
            "JUL",
            "AUG",
            "SEP",
            "OCT",
            "NOV",
            "DEC",
        ],
        1,
    )
}
_SEASON_RE = re.compile(r"\((\w{3})-(\d{2})\)\s*$")


def sort_key(dimension: str, text: str) -> tuple[Any, ...]:
    """A natural order for new values: seasons by month, numbers by size."""
    if dimension == "season":
        found = _SEASON_RE.search(text.upper())
        if found and found.group(1) in _MONTHS:
            return (0, int(found.group(2)), _MONTHS[found.group(1)], text)
    parts = re.split(r"(\d+(?:\.\d+)?)", text.upper())
    return (
        1,
        *[
            (0, float(p), "") if re.fullmatch(r"\d+(?:\.\d+)?", p) else (1, 0.0, p)
            for p in parts
            if p
        ],
    )


def season_code(label: str) -> str:
    """A Season master code for a sheet season: ``SPRING SUMMER(Jan-25)`` -> ``ss-jan-25``."""
    found = _SEASON_RE.search(label.upper())
    if found and found.group(1) in _MONTHS:
        prefix = "".join(word[0] for word in label[: found.start()].split()).lower()
        return f"{prefix or 's'}-{found.group(1).lower()}-{found.group(2)}"[
            :MAX_SEASON_CODE
        ]
    return slug(label, MAX_SEASON_CODE)


# ----------------------------------------------------------------------------- the seed reader


@dataclass
class MasterSheet:
    """The KDPS master sheet's vocabularies, ITEM helper and brand names."""

    values: dict[str, set[str]] = field(default_factory=dict)
    item_helper: dict[str, tuple[str, str]] = field(default_factory=dict)
    brands: set[str] = field(default_factory=set)
    read: list[str] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)


def read_master_sheets(paths: Iterable[Path]) -> MasterSheet:
    """Union every sheet's values; for the ITEM helper the first sheet wins."""
    sheet = MasterSheet(values={dim: set() for dim in DIM_COLS.values()})
    for path in paths:
        if not path.exists():
            sheet.missing.append(str(path))
            continue
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = [list(r) for r in workbook[SHEET_NAME].iter_rows(values_only=True)]
        workbook.close()
        sheet.read.append(str(path))
        for row in rows[1:]:
            cells = [cell_text(c) for c in row]
            for index, dimension in DIM_COLS.items():
                if index < len(cells) and cells[index]:
                    sheet.values[dimension].add(cells[index])
            if len(cells) > 1 and cells[1]:
                sheet.brands.add(cells[1])
            item = cells[6] if len(cells) > 6 else ""
            if item and item not in sheet.item_helper:
                helper = tuple(cells[i] if i < len(cells) else "" for i in HELPER_COLS)
                sheet.item_helper[item] = (helper[0], helper[1])
    return sheet


# ----------------------------------------------------------------------------- the upload reader


def _header_key(text: Any) -> str:
    return re.sub(r"\s+", " ", cell_text(text)).upper()


def _pick_sheet(workbook: Any) -> Any:
    named = [
        ws for ws in workbook.worksheets if _header_key(ws.title) == SHEET_NAME.upper()
    ]
    if named:
        return named[0]
    matching = []
    for ws in workbook.worksheets:
        first = next(ws.iter_rows(max_row=1, values_only=True), ())
        if tuple(_header_key(c) for c in list(first)[: len(HEADERS)]) == HEADERS:
            matching.append(ws)
    if len(matching) == 1:
        return matching[0]
    raise _sheet_invalid(
        f'The workbook has no sheet named "{SHEET_NAME}". Upload the KDPS PT file '
        "(its first sheet holds the lists) or a copy of that sheet."
    )


def read_master_sheet(data: bytes) -> dict[str, Any]:
    """Read one uploaded workbook's master sheet into plain JSON (see :func:`plan_changes`).

    The header row must name A-L as the KDPS sheet does; every column is its own
    list, read in sheet order with the row each value came from. Problems are noted
    per value: a blocking one leaves only that value out.
    """
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
    except Exception as exc:  # openpyxl raises many types for a damaged file
        raise _sheet_invalid(
            "The file could not be read as an Excel workbook."
        ) from exc
    try:
        ws = _pick_sheet(workbook)
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        title = str(ws.title)
    finally:
        workbook.close()
    if not rows:
        raise _sheet_invalid(f'The "{SHEET_NAME}" sheet is empty.')
    header = [_header_key(c) for c in rows[0][: len(HEADERS)]]
    header += [""] * (len(HEADERS) - len(header))
    wrong = [
        f"column {chr(65 + i)} should be {want} (found {got or 'blank'})"
        for i, (want, got) in enumerate(zip(HEADERS, header, strict=True))
        if want != got
    ]
    if wrong:
        raise _sheet_invalid(
            "The master sheet's header row does not match the KDPS layout: "
            + "; ".join(wrong)
            + "."
        )
    body = rows[1:]
    while body and not any(cell_text(c) for c in body[-1]):
        body.pop()
    if len(body) > MAX_ROWS:
        raise _sheet_invalid(f"The master sheet has more than {MAX_ROWS} rows.")

    problems: list[dict[str, Any]] = []
    columns: dict[str, list[dict[str, Any]]] = {
        name: [] for name in (*DIMENSIONS, "brand")
    }
    seen: dict[str, dict[str, int]] = {name: {} for name in columns}

    def note(
        column: str, row: int, text: str, code: str, message: str, blocking: bool
    ) -> None:
        problems.append(
            {
                "column": column,
                "row": row,
                "text": text,
                "code": code,
                "message": message,
                "blocking": blocking,
            }
        )

    def take(column: str, raw: Any, row: int, limit: int) -> None:
        text = cell_text(raw)
        if not text:
            return
        if isinstance(raw, str) and raw != text:
            note(
                column,
                row,
                text,
                WHITESPACE_FIXED,
                "Extra or unusual spaces were removed.",
                False,
            )
        if len(text) > limit:
            note(
                column,
                row,
                text,
                TOO_LONG,
                f"Longer than {limit} characters; this value is left out.",
                True,
            )
            return
        key = match_key(text)
        first = seen[column].get(key)
        if first is not None:
            note(
                column,
                row,
                text,
                DUPLICATE_IN_SHEET,
                f"Already listed on row {first}; kept once.",
                False,
            )
            return
        seen[column][key] = row
        columns[column].append({"text": text, "row": row})

    helper: list[dict[str, Any]] = []
    gst: list[str] = []
    for offset, raw_row in enumerate(body):
        row = offset + 2
        cells = list(raw_row) + [None] * (len(HEADERS) - len(raw_row))
        for index, dimension in DIM_COLS.items():
            take(dimension, cells[index], row, MAX_VALUE)
        take("brand", cells[BRAND_COL], row, MAX_BRAND)
        item = cell_text(cells[6])
        if item and len(item) <= MAX_VALUE:
            sub, kind = cell_text(cells[10]), cell_text(cells[11])
            if sub or kind:
                helper.append(
                    {"item": item, "row": row, "sub_category": sub, "type": kind}
                )
        tax = cell_text(cells[GST_COL])
        if tax and tax not in gst:
            gst.append(tax)
    return {
        "sheet": title,
        "rows_read": len(body),
        "columns": columns,
        "helper": helper,
        "gst": gst,
        "problems": problems,
    }


# ----------------------------------------------------------------------------- the plan


#: Lists whose values differ by design in a digit or two (``28``/``29``, ``Jan-25``/
#: ``Jan-26``): only a spacing or punctuation twin is worth a second look there.
_EXACT_ONLY = frozenset({"brand", "size", "season"})


def _one_letter_more(a: str, b: str) -> bool:
    """One of ``a``/``b`` is the other with one extra character (``GREE``/``GREEN``).

    A changed letter is not counted: KURTA/KURTI and SHIRT/SKIRT are different things.
    """
    if abs(len(a) - len(b)) != 1:
        return False
    if len(a) > len(b):
        a, b = b, a
    i = 0
    while i < len(a) and a[i] == b[i]:
        i += 1
    return a[i:] == b[i + 1 :]


def _words(text: str) -> list[str]:
    return [w for w in re.split(r"[^0-9a-z]+", text.casefold()) if w]


def _close(a: str, b: str) -> bool:
    """Two different texts a person may have meant as one: ``GREE``/``GREEN``.

    One letter missing or extra in a word of four or more letters; not a whole extra
    word (``SHIRT``/``T-SHIRT``) and not a different number.
    """
    ka, kb = compact(a), compact(b)
    if not ka or not kb or min(len(ka), len(kb)) < 4 or not _one_letter_more(ka, kb):
        return False
    if re.sub(r"\d", "", ka) == re.sub(r"\d", "", kb):
        return False
    wa, wb = _words(a), _words(b)
    return not (set(wa) < set(wb) or set(wb) < set(wa))


def _looks_like(
    dimension: str, added: list[dict[str, Any]], others: list[str]
) -> list[dict[str, Any]]:
    """Each added value that is a near twin of a value already known or listed earlier."""
    out: list[dict[str, Any]] = []
    pool = list(others)
    for entry in added:
        text = entry["text"]
        key = compact(text)
        twin = next(
            (other for other in pool if other != text and compact(other) == key), ""
        )
        if not twin and dimension not in _EXACT_ONLY:
            twin = next(
                (other for other in pool if other != text and _close(text, other)), ""
            )
        if twin:
            out.append(
                {
                    "column": dimension,
                    "row": entry["row"],
                    "text": text,
                    "code": LOOKS_LIKE,
                    "message": f"{text} and {twin} differ by one letter or by spacing. "
                    "Leave one out if it is a typo.",
                    "blocking": False,
                }
            )
        pool.append(text)
    return out


def _value_index(values: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}
    for value in values:
        for text in (value["value_key"], value["label"]):
            index.setdefault(match_key(text), value)
    return index


def plan_changes(
    parsed: dict[str, Any], current: dict[str, Any], selections: dict[str, Any]
) -> dict[str, Any]:
    """What approving this sheet would change, as plain JSON.

    ``current`` is the tenant's snapshot (see ``master_sheet_services.snapshot``):
    each dimension's published values, live brands, Season masters, effective ITEM
    rules and the configuration that pins or restricts lists. ``selections`` are the
    reviewer's choices: values to retire or leave out, brands and seasons to skip,
    the option picked for an ITEM with several, and acknowledged warnings.
    """
    retire = {str(v) for v in selections.get("retire", [])}
    skip_values = {
        dim: {match_key(t) for t in texts}
        for dim, texts in (selections.get("skip_values") or {}).items()
    }
    skip_brands = {match_key(t) for t in selections.get("skip_brands", [])}
    skip_seasons = {match_key(t) for t in selections.get("skip_seasons", [])}
    choices = selections.get("rule_choices") or {}
    acknowledged = set(selections.get("acknowledged", []))
    problems = list(parsed.get("problems", []))

    dimensions: list[dict[str, Any]] = []
    after: dict[str, list[str]] = {}
    for dimension in DIMENSIONS:
        values = current["values"].get(dimension, [])
        index = _value_index(values)
        in_sheet: set[str] = set()
        added: list[dict[str, Any]] = []
        present: list[dict[str, Any]] = []
        back: list[dict[str, Any]] = []
        left_out: list[dict[str, Any]] = []
        for entry in parsed["columns"].get(dimension, []):
            key = match_key(entry["text"])
            found = index.get(key)
            if found is not None:
                in_sheet.add(found["value_id"])
                if found["retired"] and found["value_id"] not in retire:
                    back.append(
                        {"value_id": found["value_id"], "label": found["label"]}
                    )
                else:
                    present.append(
                        {"value_id": found["value_id"], "label": found["label"]}
                    )
            elif key in skip_values.get(dimension, set()):
                left_out.append(entry)
            else:
                added.append(entry)
        not_in_sheet = [
            {
                "value_id": v["value_id"],
                "label": v["label"],
                "retire": v["value_id"] in retire,
            }
            for v in values
            if not v["retired"] and v["value_id"] not in in_sheet
        ]
        problems.extend(_looks_like(dimension, added, [v["label"] for v in values]))
        payload_values: list[dict[str, Any]] = []
        back_ids = {b["value_id"] for b in back}
        for v in values:
            retired = v["retired"]
            if v["value_id"] in back_ids:
                retired = False
            elif v["value_id"] in retire and not v["retired"]:
                retired = True
            payload_values.append(
                {
                    "value_key": v["value_key"],
                    "label": v["label"],
                    "sort_order": v["sort_order"],
                    "retired": retired,
                }
            )
        start = max((v["sort_order"] for v in values), default=-1) + 1
        for offset, entry in enumerate(added):
            payload_values.append(
                {
                    "value_key": entry["text"],
                    "label": entry["text"],
                    "sort_order": start + offset,
                    "retired": False,
                }
            )
        changed = bool(added or back or any(n["retire"] for n in not_in_sheet))
        after[dimension] = [v["label"] for v in payload_values if not v["retired"]]
        dimensions.append(
            {
                "dimension": dimension,
                "column": COLUMN_LABELS[dimension],
                "current_version_id": current["version_ids"].get(dimension),
                "changed": changed,
                "added": added,
                "left_out": left_out,
                "present": present,
                "back_in_use": back,
                "not_in_sheet": not_in_sheet,
                "values": payload_values if changed else [],
                "counts": {
                    "added": len(added),
                    "present": len(present),
                    "back_in_use": len(back),
                    "not_in_sheet": len(not_in_sheet),
                    "retire": sum(1 for n in not_in_sheet if n["retire"]),
                    "left_out": len(left_out),
                },
            }
        )

    brands = _plan_brands(parsed, current, skip_brands, problems)
    seasons = _plan_seasons(parsed, current, skip_seasons, after["season"])
    rules = _plan_rules(
        parsed, current, after, choices, retire_items=_retired_items(dimensions)
    )
    problems.extend(rules.pop("problems"))
    warnings = _warnings(current, dimensions, acknowledged)
    for d in dimensions:
        d["problems"] = [p for p in problems if p["column"] == d["dimension"]]
        d["counts"]["problems"] = len(d["problems"])
    brands["problems"] = [p for p in problems if p["column"] == "brand"]
    rules["problems"] = [p for p in problems if p["column"] == "item_rule"]

    changes = (
        sum(1 for d in dimensions if d["changed"])
        + len([b for b in brands["to_create"] if b["include"]])
        + len([s for s in seasons["to_create"] if s["include"]])
        + len(rules["new"])
        + len(rules["changed"])
    )
    return {
        "dimensions": dimensions,
        "brands": brands,
        "seasons": seasons,
        "item_rules": rules,
        "ignored_columns": [
            {"column": "GST %", "values": parsed.get("gst", []), "reason": GST_NOTE}
        ],
        "warnings": warnings,
        "change_count": changes,
        "retires": any(d["counts"]["retire"] for d in dimensions),
        "ready": changes > 0 and all(w["acknowledged"] for w in warnings),
    }


def _retired_items(dimensions: list[dict[str, Any]]) -> set[str]:
    item = next(d for d in dimensions if d["dimension"] == "item")
    return {match_key(n["label"]) for n in item["not_in_sheet"] if n["retire"]}


def _plan_brands(
    parsed: dict[str, Any],
    current: dict[str, Any],
    skip: set[str],
    problems: list[dict[str, Any]],
) -> dict[str, Any]:
    known: dict[str, dict[str, Any]] = {}
    for brand in current["brands"]:
        for text in (brand["code"], brand["name"]):
            known.setdefault(compact(text), brand)
    codes = {b["code"].lower() for b in current["brands"]}
    to_create: list[dict[str, Any]] = []
    present: set[int] = set()
    planned: dict[str, str] = {}
    for entry in parsed["columns"].get("brand", []):
        found = known.get(compact(entry["text"]))
        if found is not None:
            present.add(found["id"])
            continue
        twin = planned.get(compact(entry["text"]))
        if twin is not None:
            problems.append(
                {
                    "column": "brand",
                    "row": entry["row"],
                    "text": entry["text"],
                    "code": DUPLICATE_IN_SHEET,
                    "message": f"Same brand as {twin} (spacing or punctuation differs); kept once.",
                    "blocking": False,
                }
            )
            continue
        planned[compact(entry["text"])] = entry["text"]
        base = slug(entry["text"], MAX_BRAND_CODE) or "brand"
        code, n = base, 2
        while code in codes:
            suffix = f"-{n}"
            code = f"{base[: MAX_BRAND_CODE - len(suffix)]}{suffix}"
            n += 1
        codes.add(code)
        to_create.append(
            {
                "name": entry["text"],
                "code": code,
                "row": entry["row"],
                "include": match_key(entry["text"]) not in skip,
            }
        )
    problems.extend(
        _looks_like(
            "brand",
            [{"text": b["name"], "row": b["row"]} for b in to_create],
            [b["name"] for b in current["brands"]],
        )
    )
    in_sheet = {compact(e["text"]) for e in parsed["columns"].get("brand", [])}
    not_in_sheet = [
        {"id": b["id"], "name": b["name"]}
        for b in current["brands"]
        if b["id"] not in present and compact(b["name"]) not in in_sheet
    ]
    return {
        "to_create": to_create,
        "present_count": len(present),
        "not_in_sheet": not_in_sheet,
        "terms_note": "New brands start with their commercial terms to be set on Brands > Terms.",
    }


def _plan_seasons(
    parsed: dict[str, Any], current: dict[str, Any], skip: set[str], labels: list[str]
) -> dict[str, Any]:
    names = {match_key(s["name"]) for s in current["seasons"]} | {
        match_key(s["code"]) for s in current["seasons"]
    }
    codes = {s["code"].lower() for s in current["seasons"]}
    sheet = {match_key(e["text"]) for e in parsed["columns"].get("season", [])}
    to_create: list[dict[str, Any]] = []
    for label in labels:
        key = match_key(label)
        if key in names or key not in sheet:
            continue
        base = season_code(label) or "season"
        code, n = base, 2
        while code in codes:
            suffix = f"-{n}"
            code = f"{base[: MAX_SEASON_CODE - len(suffix)]}{suffix}"
            n += 1
        codes.add(code)
        to_create.append({"name": label, "code": code, "include": key not in skip})
    return {"to_create": to_create}


def _plan_rules(
    parsed: dict[str, Any],
    current: dict[str, Any],
    after: dict[str, list[str]],
    choices: dict[str, Any],
    *,
    retire_items: set[str],
) -> dict[str, Any]:
    rules = current.get("item_rules", {})
    items = {match_key(label) for label in after["item"]}
    new: list[dict[str, Any]] = []
    changed: list[dict[str, Any]] = []
    same = 0
    problems: list[dict[str, Any]] = []
    for row in parsed.get("helper", []):
        key = match_key(row["item"])
        if key not in items or key in retire_items:
            continue
        entry: dict[str, Any] = {"item": row["item"], "row": row["row"]}
        status: set[str] = set()
        for dimension in ("sub_category", "type"):
            cell = row[dimension]
            options = helper_options(cell, after[dimension])
            if not options:
                if cell:
                    problems.append(
                        {
                            "column": "item_rule",
                            "row": row["row"],
                            "text": f"{row['item']}: {cell}",
                            "code": HELPER_UNKNOWN,
                            "message": f"{cell} is not a {COLUMN_LABELS[dimension]} value; "
                            "no suggestion is set for this ITEM.",
                            "blocking": False,
                        }
                    )
                continue
            existing = rules.get(f"{dimension}\x1f{key}")
            wanted = (choices.get(key) or {}).get(dimension)
            if wanted not in options and existing and existing["target_label"] in options:
                wanted = existing["target_label"]  # a past pick stands until changed
            chosen = wanted if wanted in options else options[0]
            if len(options) > 1:
                problems.append(
                    {
                        "column": "item_rule",
                        "row": row["row"],
                        "text": f"{row['item']}: {cell}",
                        "code": HELPER_MULTIPLE,
                        "message": f"{cell} names {len(options)} values; {chosen} is used. "
                        "Pick another if it fits better.",
                        "blocking": False,
                    }
                )
            target = str(
                vocabulary_value_id(
                    dimension, _value_key(current, after, dimension, chosen)
                )
            )
            entry[dimension] = {
                "value": chosen,
                "options": options,
                "target_id": target,
                "from": existing["target_label"] if existing else None,
                "rule_id": existing["rule_id"] if existing else None,
            }
            if existing is None:
                status.add("new")
            elif existing["target_id"] != target:
                status.add("changed")
        if not status:
            if any(dim in entry for dim in ("sub_category", "type")):
                same += 1
            continue
        (changed if "changed" in status else new).append(entry)
    return {"new": new, "changed": changed, "same_count": same, "problems": problems}


def _value_key(
    current: dict[str, Any], after: dict[str, list[str]], dimension: str, label: str
) -> str:
    """The value key a label carries once the import applies (a new value's key is its text)."""
    for value in current["values"].get(dimension, []):
        if match_key(value["label"]) == match_key(label) or match_key(
            value["value_key"]
        ) == match_key(label):
            return str(value["value_key"])
    return label


def _warnings(
    current: dict[str, Any], dimensions: list[dict[str, Any]], acknowledged: set[str]
) -> list[dict[str, Any]]:
    changed = {d["dimension"]: d for d in dimensions if d["changed"]}
    out: list[dict[str, Any]] = []

    def add(code: str, dimension: str, message: str) -> None:
        wid = f"{code}:{dimension}"
        out.append(
            {
                "id": wid,
                "code": code,
                "dimension": dimension,
                "message": message,
                "acknowledged": wid in acknowledged,
            }
        )

    for dimension in sorted(set(current.get("pending_drafts", [])) & set(changed)):
        add(
            PENDING_DRAFT,
            dimension,
            f"A {COLUMN_LABELS[dimension]} list change is already waiting in Configuration. "
            "Approving this import first means that draft must be redone.",
        )
    for pin in current.get("pinned_profiles", []):
        if pin["dimension"] in changed:
            add(
                PINNED_PROFILE,
                pin["dimension"],
                f"A PT profile is fixed to the current {COLUMN_LABELS[pin['dimension']]} list. "
                "After this import, new PTs on that profile stop until the profile is updated.",
            )
    for dimension, allowed in (current.get("allow_lists") or {}).items():
        d = changed.get(dimension)
        if d and allowed and d["added"]:
            add(
                ALLOW_LIST,
                dimension,
                f"The product identity profile only allows listed {COLUMN_LABELS[dimension]} "
                "values; new ones need adding there before products use them.",
            )
    return out
