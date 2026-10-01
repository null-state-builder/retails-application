"""The KDPS PT file, generated from the lists in force (download from Product lists).

Two sheets, laid out as KDPS's own PT file: "Master Sheet" holds every list
(columns A-L, the import's own layout) and "Work Sheet" is a blank PT with the
same dropdowns and formulas the staff sheets carry. Uploading the generated Master
Sheet back changes nothing; a filled Work Sheet uploads as a PT.
"""

from __future__ import annotations

import io
import uuid
from datetime import datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from masters.goods_identity_models import GovernanceState, SourceCrosswalk
from masters.goods_identity_services import ITEM_ISSUER, vocabulary
from masters.master_sheet import DIM_COLS, HEADERS
from masters.models import Brand

WORK_SHEET = "Work Sheet"
WORK_ROWS = 1000
FIRST_ROW = 3
#: The GST % reference list the KDPS sheet carries (tax itself comes from Tax settings).
GST_REFERENCE = [5, 12, 18, "TAX FREE"]
WORK_HEADERS = [
    "SEASON",
    "BRAND",
    "COLOR",
    "GENDER",
    "SUB CATEGORY",
    "TYPE",
    "ITEM",
    "FIT",
    "SIZE",
    "BARCODE",
    "DESIGN",
    "HSN",
    "QTY",
    "MRP",
    "BASIC",
    "P RATE",
    "INPUT TAX",
    "OUTPUT TAX",
    "NAG",
    "MARGIN",
    None,
    "SUGGESTED SUB CATEGORY",
    "SUGGESTED TYPE",
]
_HEAD = Font(bold=True)
_FILL = PatternFill("solid", fgColor="FFF4EFE7")


def _tax_formula(r: int, amount: str, limit: int) -> str:
    """The staff sheet's GST rule (input on BASIC, output on MRP)."""
    return (
        f'=IF(M{r}="","",IF(E{r}="FABRIC",5,IF(G{r}="SAREE",5,'
        f'IF(OR(G{r}="BELT",G{r}="LADIES PURSE",G{r}="WALLET"),18,IF(G{r}="","",'
        f'IF(AND(F{r}="LUGGAGE",{amount}{r}<>""),18,IF({amount}{r}="","",'
        f"IF({amount}{r}<={limit},5,18))))))))"
    )


def template_bytes(
    tenant_id: uuid.UUID, now: datetime, p_rate_factor: Decimal
) -> bytes:
    lists = vocabulary(tenant_id, now)
    columns: dict[int, list[Any]] = {}
    for index, dimension in DIM_COLS.items():
        values = sorted(
            (v for v in lists.get(dimension, []) if not v.retired),
            key=lambda v: v.sort_order,
        )
        columns[index] = [v.label for v in values]
    columns[1] = list(
        Brand.objects.filter(is_active=True)
        .order_by("name")
        .values_list("name", flat=True)
    )
    columns[9] = list(GST_REFERENCE)
    labels = {str(v.id): v.label for values in lists.values() for v in values}
    rules: dict[tuple[str, str], str] = {}
    for kind, source, target in SourceCrosswalk.objects.filter(
        tenant_id=tenant_id,
        kind__in=["sub_category", "type"],
        issuer_key=ITEM_ISSUER,
        governance_state=GovernanceState.EFFECTIVE,
        retired_at__isnull=True,
    ).values_list("kind", "source_key", "target_key"):
        rules.setdefault(
            (kind, " ".join(source.split()).casefold()), labels.get(target, "")
        )
    items = columns[6]
    columns[10] = [rules.get(("sub_category", item.casefold()), "") for item in items]
    columns[11] = [rules.get(("type", item.casefold()), "") for item in items]

    book = Workbook()
    master = book.active
    assert master is not None
    master.title = "Master Sheet"
    master.append(list(HEADERS))
    height = max((len(v) for v in columns.values()), default=0)
    for r in range(height):
        master.append(
            [
                columns[c][r] if r < len(columns.get(c, [])) else None
                for c in range(len(HEADERS))
            ]
        )
    for cell in master[1]:
        cell.font, cell.fill = _HEAD, _FILL
    for c in range(1, len(HEADERS) + 1):
        master.column_dimensions[get_column_letter(c)].width = 24
    master.freeze_panes = "A2"

    work = book.create_sheet(WORK_SHEET)
    work.append([])
    work.append(WORK_HEADERS)
    for cell in work[2]:
        cell.font, cell.fill = _HEAD, _FILL
    last = FIRST_ROW + WORK_ROWS - 1
    factor = format(p_rate_factor.normalize(), "f")
    for r in range(FIRST_ROW, last + 1):
        work[f"P{r}"] = f'=IF(O{r}<>"",O{r}*{factor},"")'
        work[f"Q{r}"] = _tax_formula(r, "O", 2500)
        work[f"R{r}"] = _tax_formula(r, "N", 2625)
        work[f"S{r}"] = f"=M{r}"
        work[f"T{r}"] = f'=IFERROR((N{r}-P{r})*100/N{r},"")'
        work[f"V{r}"] = (
            f'=IF($G{r}=""," ",IFERROR(VLOOKUP($G{r},\'Master Sheet\'!$G$2:$L${height + 1},'
            f'5,0),"WRONG ITEM"))'
        )
        work[f"W{r}"] = (
            f'=IF($G{r}=""," ",IFERROR(VLOOKUP($G{r},\'Master Sheet\'!$G$2:$L${height + 1},'
            f'6,0),"PLEASE RECTIFY"))'
        )
    for c in range(len(DIM_COLS) + 1):
        letter = get_column_letter(c + 1)
        size = len(columns.get(c, []))
        if not size:
            continue
        check = DataValidation(
            type="list",
            formula1=f"'Master Sheet'!${letter}$2:${letter}${size + 1}",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Not in the list",
            error=f"Pick a {HEADERS[c]} from the Master Sheet.",
        )
        check.add(f"{letter}{FIRST_ROW}:{letter}{last}")
        work.add_data_validation(check)
    for c in range(1, len(WORK_HEADERS) + 1):
        work.column_dimensions[get_column_letter(c)].width = 16
    work.freeze_panes = "A3"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
