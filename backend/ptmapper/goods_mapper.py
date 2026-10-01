"""Brand-file mapping for goods-v1 PT work (store and warehouse operations PRD §5.4, OPS-14).

A brand's own PT file is read, its sheet and header row found, its columns matched
to the twenty-two KDPS columns through the brand-file profiles (``profiles.py``),
and each describing value cleaned (colour, size, season, fit, gender) and taken
through the one governed rulebook (:mod:`ptmapper.goods_rulebook`).

For every KDPS column of a row, :func:`map_row` answers the value *and where it
came from*:

* ``file`` - the file's own value, valid as it stands;
* ``rule`` - a confirmed rule (or the product's own mechanical clean-up, such as
  ``2XL`` -> ``XXL`` or a season read from the invoice date) gave it;
* ``suggestion`` - no rule applies; close matches wait for a person. The value is
  **not** filled - ``value`` stays empty and ``suggestions`` lists the choices;
* ``none`` - nothing to offer.

Derived columns in a brand file (P RATE, taxes, margin) are carried as checks
only; OUTPUT TAX, NAG, MARGIN and the calculations are the product's, never the
file's. Nothing here writes to the database: OPS-15 stores the result and OPS-16
draws it. No AI or language-model assistance. This module imports nothing from
the legacy PT mapper's models or engine.

Typical use::

    rulebook = Rulebook.load(tenant_id, now)
    mapped = map_file(content, "PETER ENGLAND.xlsx", "", rulebook,
                      MappingContext(brand_id=brand.pk, invoice_date=date(2026, 9, 1)))
    for row in mapped.rows:
        row.cells["COLOR"].value, row.cells["COLOR"].origin
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any

import openpyxl

from masters.goods_identity_services import (
    DEFAULT_SOURCE,
    ITEM_ISSUER,
    KEYWORD_ISSUER,
    brand_issuer,
)
from ptmapper.goods_rulebook import (
    CONFLICTING_RULES,
    FILE,
    NAMED_IN_TEXT,
    NONE,
    RULE,
    SUGGESTION,
    Candidate,
    Rule,
    Rulebook,
    RuleValue,
    close_matches,
    issuers_for,
)
from ptmapper.profiles import (
    ALIASES,
    FILENAME_HINTS,
    GENERIC_PROFILE,
    HEADER_KEYWORDS,
    KDPS_COLUMNS,
    MASTER_SHEET_NAME,
    PROFILES,
    WORK_SHEET_HEADERS,
)

MAX_ROWS = 8000

#: KDPS describing column -> the governed vocabulary dimension it takes values from.
#: BRAND is not a vocabulary: it names a brand master (a ``brand`` crosswalk).
COLUMN_DIMENSIONS: dict[str, str] = {
    "SEASON": "season",
    "COLOR": "colour",
    "GENDER": "gender",
    "SUB CATEGORY": "sub_category",
    "TYPE": "type",
    "ITEM": "item",
    "FIT": "fit",
    "SIZE": "size",
}

#: Why a ``rule``-origin value needed no stored rule: the product's own clean-up.
CLEANED = "cleaned"
INVOICE_DATE = "invoice_date"
RECEIPT_DATE = "receipt_date"
RECEIPT_BRAND = "receipt_brand"
PROFILE_BRAND = "profile_brand"
CONFIRMED_RULE = "confirmed_rule"
DESCRIPTION_WORD = "description_word"
KEYWORD_DEFAULT = "keyword_default"
BRAND_DEFAULT = "brand_default"
ITEM_HELPER = "item_helper"


class UnsupportedFormat(Exception):
    pass


# ----------------------------------------------------------------------------- results


@dataclass(frozen=True)
class Suggestion:
    """A close match for a person to accept; never applied on its own."""

    value_id: str
    value: str
    label: str
    reason: str

    def as_dict(self) -> dict[str, str]:
        return {
            "value_id": self.value_id,
            "value": self.value,
            "label": self.label,
            "reason": self.reason,
        }


@dataclass(frozen=True)
class MappedCell:
    """One KDPS column of one mapped row.

    ``value`` is the vocabulary value key for a describing column, the brand ID
    for BRAND, and the file's text for the free columns. ``value_id`` is the
    vocabulary value (or brand) ID. ``source`` is the file's own text behind the
    cell. With origin ``suggestion`` the value is empty and ``suggestions``
    carries the choices.
    """

    value: str | None
    origin: str
    source: str = ""
    value_id: str | None = None
    label: str | None = None
    rule_id: str | None = None
    reason: str | None = None
    suggestions: tuple[Suggestion, ...] = ()

    def as_dict(self) -> dict[str, Any]:
        return {
            "value": self.value,
            "origin": self.origin,
            "source": self.source,
            "value_id": self.value_id,
            "label": self.label,
            "rule_id": self.rule_id,
            "reason": self.reason,
            "suggestions": [s.as_dict() for s in self.suggestions],
        }


@dataclass(frozen=True)
class MappedRow:
    line_no: int
    source_row: int
    cells: dict[str, MappedCell]

    def as_dict(self) -> dict[str, Any]:
        return {
            "line_no": self.line_no,
            "source_row": self.source_row,
            "cells": {column: self.cells[column].as_dict() for column in KDPS_COLUMNS},
        }


@dataclass(frozen=True)
class MappingContext:
    """What the caller knows beyond the file.

    ``brand_id`` is the receipt's brand: it fills BRAND when the file names none
    and scopes brand rules. ``issuer_key`` is the vendor whose file this is, for
    brand crosswalks. ``invoice_date`` is the season fallback when the file gives
    neither a season nor a date.
    """

    brand_id: int | None = None
    issuer_key: str | None = None
    invoice_date: date | None = None


@dataclass(frozen=True)
class BrandFile:
    filename: str
    sheet: str
    header_row: int
    headers: list[str]
    profile: dict[str, Any]
    records: list[tuple[int, dict[str, Any]]]
    truncated: bool
    brand_guess: str


@dataclass(frozen=True)
class MappedFile:
    profile_code: str
    profile_name: str
    archetype: str
    sheet: str
    header_row: int
    headers: list[str]
    source_rows: int
    truncated: bool
    rows: list[MappedRow] = field(default_factory=list)


# ----------------------------------------------------------------------------- helpers


def norm(v: Any) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).strip()).upper()


def raw_str(v: Any) -> str:
    return "" if v is None else str(v).strip()


def clean_code(v: Any) -> str:
    """Stringify a barcode/HSN without a spurious trailing .0 from float cells."""
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def number_text(v: Any) -> str:
    """A numeric cell as exact decimal text ('' when blank or not a number)."""
    if v is None or isinstance(v, bool):
        return ""
    text = str(v).replace(",", "").strip().rstrip("%").strip()
    if not text:
        return ""
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return ""
    if not amount.is_finite():
        return ""
    if amount == amount.to_integral_value():
        return str(amount.quantize(Decimal(1)))
    return format(amount.normalize(), "f")


_DATE_FORMATS = [
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%Y",
    "%Y%m%d",
    "%m/%d/%Y",
]


def parse_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    # Excel serial date (common in .xlsb exports where dates arrive as numbers).
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < float(v) < 80000:
        return (datetime(1899, 12, 30) + timedelta(days=int(v))).date()
    s = raw_str(v)
    if not s:
        return None
    s = s.split(".")[0] if s.isdigit() and len(s) > 8 else s
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def season_label(d: date) -> str:
    """KDPS's season name for a month: ``SPRING SUMMER(Jan-26)``, ``AUTUMN WINTER(Aug-26)``."""
    prefix = "SPRING SUMMER" if d.month <= 6 else "AUTUMN WINTER"
    return f"{prefix}({d.strftime('%b')}-{d.strftime('%y')})"


# --------------------------------------------------------------------------- normalisers
# Mechanical value normalisers: turn a brand's raw string into the KDPS canonical
# form deterministically so it lands directly on an approved value. They encode no
# business judgement (shade -> bucket, category -> item): that lives in the governed
# rulebook. Pure functions, unit-tested without a database.

# Alpha size variants -> KDPS master size. Master has XS/S/M/L/XL/XXL/3XL/4XL/5XL/6XL
# and "FREE SIZE" (no "2XL"/"XXXL"), so collapse those here.
_ALPHA_SIZE_ALIASES = {
    "2XL": "XXL",
    "XXXL": "3XL",
    "XXXXL": "4XL",
    "XXXXXL": "5XL",
    "XXXXXXL": "6XL",
    "FS": "FREE SIZE",
    "F": "FREE SIZE",
    "FREE": "FREE SIZE",
    "FREESIZE": "FREE SIZE",
    "ONESIZE": "FREE SIZE",
    "OS": "FREE SIZE",
    "1MTR": "FREE SIZE",
    "1MTR.": "FREE SIZE",
}
_AGE_Y_RE = re.compile(r"^(\d+)\s*-\s*(\d+)\s*(?:Y|YR|YRS|YEAR|YEARS)\.?$")
_AGE_M_RE = re.compile(r"^(\d+)\s*-\s*(\d+)\s*(?:M|MO|MTH|MTHS|MONTH|MONTHS)\.?$")
_PAREN_RE = re.compile(r"\(([^)]+)\)")
_ALPHA_SIZE_RE = re.compile(r"\d*X*[SML]")  # S, M, L, XS, XL, XXL, 3XL, ...


def _canon_alpha_size(s: str) -> str:
    u = re.sub(r"\s+", " ", s.strip().upper())
    return _ALPHA_SIZE_ALIASES.get(u.replace(" ", ""), u)


def _is_alpha_size(s: str) -> bool:
    return bool(_ALPHA_SIZE_RE.fullmatch(s.replace(" ", "")))


def _size_once(s: str) -> str:
    """One normalisation pass over an already upper/cleaned size string."""
    # "100 CMS" / "96 CMS." -> drop the (possibly repeated) plural CMS suffix to the
    # bare number (a bare "96 CM" stays: the master keeps "96 CM" distinct from "96").
    s = re.sub(r"(?:\s*CMS\.?)+\s*$", "", s)
    if not s:
        return ""
    if s.replace(" ", "") in _ALPHA_SIZE_ALIASES:  # FS, 2XL, FREE, 1 MTR. ...
        return _ALPHA_SIZE_ALIASES[s.replace(" ", "")]
    m = _PAREN_RE.search(s)  # a measurement + an alpha size, either side in the parens
    if m:
        base = _canon_alpha_size(s[: m.start()])  # "XL (105 CMS)" -> base "XL"
        inner = _canon_alpha_size(m.group(1))  # "1.14M(2XL)" -> inner "XXL"
        if _is_alpha_size(base):
            return base
        if _is_alpha_size(inner):
            return inner
        return base or inner  # neither side alpha -> reduced further on the next pass
    if "/" in s:  # "36/XS" = the SAME size in two notations -> take the alpha (XS)
        parts = [p.strip() for p in s.split("/") if p.strip()]
        alpha = [p for p in parts if re.search(r"[A-Za-z]", p)]
        numeric = [p for p in parts if re.fullmatch(r"\d+", p)]
        if len(alpha) == 1 and numeric:  # num/alpha equivalence ("44/XXL" -> "XXL")
            return _canon_alpha_size(alpha[0])
        return s  # ambiguous ("S/M" = two sizes, "36/38") -> unchanged -> a person decides
    m = _AGE_Y_RE.match(s)  # "7-8Y" / "8-9 YEARS" -> "7-8 Y"
    if m:
        return f"{int(m.group(1))}-{int(m.group(2))} Y"
    m = _AGE_M_RE.match(s)  # "0-6M" / "12-18 MONTHS" -> "0-6 M"
    if m:
        return f"{int(m.group(1))}-{int(m.group(2))} M"
    return _canon_alpha_size(s)


def normalize_size(raw: Any) -> str:
    """Brand size string -> KDPS master size (best effort; unresolved stays for a person).

    Handles float artifacts ('44.0'->'44'), free-size words, metric-with-alpha
    ('1.14M(2XL)'->'XXL', '96CM(M)'->'M'), slash duals ('36/XS'->'XS'), and age
    bands ('7-8Y'/'8-9 YEARS'->'7-8 Y'). Plain numbers/bra sizes pass through.

    Idempotent: a paren/slash can expose a further-reducible token (e.g. '(105 CMS)'),
    so passes repeat to a fixed point - f(f(x)) == f(x), so re-mapping never drifts.
    """
    s = clean_code(raw).upper().strip()
    for _ in range(4):
        nxt = _size_once(s)
        if nxt == s:
            return nxt
        s = nxt
    return s


_COLOR_LEAD_CODE = re.compile(r"^\s*\d+\s*[-_]\s*")  # "88-BEIGE", "16 - BLUE"
_COLOR_TRAIL_NUM = re.compile(r"\s*\d+\s*(?:/\s*\d+)?\s*$")  # "BLACK73", "WHITE74/1"
_COLOR_TRAIL_TOK = re.compile(r"\s+(?:DN|DBY|DB|MEL|MELANGE)\s*$")  # denim/melange wash codes


def normalize_color(raw: Any) -> str:
    """Strip a brand colour string down to its shade word (mechanical only).

    Leading numeric codes ('88-BEIGE'->'BEIGE'), trailing wash/lot codes
    ('BLACK73/1'->'BLACK', 'BLUE DN89'->'BLUE'). The shade -> bucket judgement
    ('MID BLUE'->'BLUE') is a governed colour rule, not this function.
    """
    s0 = re.sub(r"\s+", " ", raw_str(raw).upper()).strip()
    if not s0:
        return ""
    # Strip leading codes and trailing lot numbers / wash codes to a fixed point, so
    # the result is idempotent ("L BLU DN88" -> "L BLU", "BLACK73/1" -> "BLACK").
    s = s0
    prev = None
    while prev != s:
        prev = s
        s = _COLOR_LEAD_CODE.sub("", s)
        s = _COLOR_TRAIL_NUM.sub("", s)
        s = _COLOR_TRAIL_TOK.sub("", s)
    s = re.sub(r"\s+", " ", s).strip()
    # An all-code input ("501", "123-") must not vanish to blank: keep it for a person.
    return s or s0


_COLOR_CODE_PREFIX = re.compile(r"^\s*\d+\s*-\s*")


def color_source(raw: Any, profile: dict[str, Any]) -> str:
    """The COLOR source text, with a '<code>-NAME' prefix stripped for flagged profiles.

    Ginesys CATEGORY3 encodes colour as '16-BLUE'; plain names and unflagged
    profiles pass through unchanged.
    """
    text = raw_str(raw)
    if profile.get("flags", {}).get("color_strip_code_prefix"):
        return _COLOR_CODE_PREFIX.sub("", text)
    return text


_GENDER_KEYWORDS = (
    # (keyword, KDPS gender) - order matters: kids/female checked before adult/male
    (" GIRLS ", "KIDS FEMALE"),
    (" GIRL ", "KIDS FEMALE"),
    (" BOYS ", "KIDS MALE"),
    (" BOY ", "KIDS MALE"),
    (" INFANT ", "UNISEX"),
    (" WOMENS ", "FEMALE"),
    (" WOMEN ", "FEMALE"),
    (" WOMAN ", "FEMALE"),
    (" LADIES ", "FEMALE"),
    (" LADIE ", "FEMALE"),
    (" FEMALE ", "FEMALE"),
    (" MENS ", "MALE"),
    (" MEN ", "MALE"),
    (" MAN ", "MALE"),
    (" GENTS ", "MALE"),
    (" MALE ", "MALE"),
    (" KIDS ", "UNISEX"),
    (" KID ", "UNISEX"),
    (" JUNIOR ", "UNISEX"),
    (" UNISEX ", "UNISEX"),
)


def gender_from_text(text: str) -> str:
    """A KDPS gender named by words in a description ('' when none).

    Separators are flattened so 'SOCKS-MENS' and 'Senior Girls Top' both match.
    The mapper offers this only as a suggestion; a governed gender rule fills.
    """
    flat = " " + re.sub(r"[^A-Z0-9]+", " ", norm(text)).strip() + " "
    for kw, g in _GENDER_KEYWORDS:
        if kw in flat:
            return g
    return ""


_FIT_CODE_TOKEN = re.compile(r"^[A-Z]{2,3}$")


def fit_code_candidates(raw: Any) -> list[str]:
    """Candidates for a coded FIT TYPE column (profile flag ``fit_code_tokens``).

    Peter England / ABFRL encode fit as '<line> <family> <style...>' - e.g.
    'PJ RG OCTANEMIDSTR': token 2 is the fit family (RG = Regular, SL = Slim).
    Order: the full value first, then token 2, then a trailing word.
    """
    s = norm(raw)
    if not s:
        return []
    candidates = [s]
    tokens = s.split()
    # The coded form always has >= 3 tokens; a plain 2-word fit ("SLIM FIT") stays one.
    if len(tokens) >= 3 and _FIT_CODE_TOKEN.fullmatch(tokens[1]):
        candidates.append(tokens[1])
    if len(tokens) >= 3 and tokens[-1].isalpha():
        candidates.append(tokens[-1])
    return candidates


def clean_for_dimension(dimension: str, text: str) -> list[str]:
    """The mechanical clean-ups of ``text`` for one dimension (may be empty)."""
    if dimension == "size":
        cleaned = [normalize_size(text)]
    elif dimension == "colour":
        cleaned = [normalize_color(text)]
    elif dimension == "fit":
        cleaned = fit_code_candidates(text)[1:]
    else:
        cleaned = []
    return [c for c in dict.fromkeys(cleaned) if c and c != text]


# ----------------------------------------------------------------------------- reading


def _read_xlsx(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out: list[tuple[str, list[list[Any]]]] = []
    for ws in wb.worksheets:
        rows: list[list[Any]] = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            rows.append(list(r))
            if i >= MAX_ROWS:
                break
        out.append((ws.title, rows))
    wb.close()
    return out


def _read_csv(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    text = None
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:  # pragma: no cover - latin-1 decodes every byte string
        raise UnsupportedFormat("Could not decode the CSV text.")
    rows: list[list[Any]] = list(csv.reader(io.StringIO(text)))[: MAX_ROWS + 1]
    return [("csv", rows)]


def _read_xls(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    """Legacy OLE2 .xls via xlrd; Excel serial dates -> datetime."""
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    out: list[tuple[str, list[list[Any]]]] = []
    for sh in book.sheets():
        rows: list[list[Any]] = []
        for i in range(min(sh.nrows, MAX_ROWS + 1)):
            row: list[Any] = []
            for j in range(sh.ncols):
                cell = sh.cell(i, j)
                raw = cell.value
                v: str | float | datetime = raw
                if cell.ctype == xlrd.XL_CELL_DATE and isinstance(raw, (int, float)):
                    try:
                        v = xlrd.xldate_as_datetime(float(raw), book.datemode)
                    except Exception:  # noqa: BLE001 - a bad date cell stays its raw number
                        pass
                row.append(v)
            rows.append(row)
        out.append((sh.name, rows))
    return out


def _read_xlsb(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    """Binary .xlsb (e.g. Madura SAP export) via pyxlsb."""
    from pyxlsb import open_workbook as open_xlsb

    out: list[tuple[str, list[list[Any]]]] = []
    with open_xlsb(io.BytesIO(content)) as wb:
        for name in wb.sheets:
            rows: list[list[Any]] = []
            with wb.get_sheet(name) as sheet:
                for i, row in enumerate(sheet.rows()):
                    rows.append([c.v for c in row])
                    if i >= MAX_ROWS:
                        break
            out.append((name, rows))
    return out


def read_sheets(
    content: bytes, filename: str, content_type: str = ""
) -> list[tuple[str, list[list[Any]]]]:
    """Every sheet of a .xlsx, .xls, .xlsb or .csv brand file as rows of cell values."""
    name = (filename or "").lower()
    is_zip = content[:4] == b"PK\x03\x04"  # xlsx / xlsb are both ZIP containers
    is_ole = content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy OLE2 .xls
    # .xlsb is a ZIP too, so it must be matched by extension before the ZIP sniff.
    if name.endswith(".xlsb"):
        return _read_xlsb(content)
    # The real container wins over a (possibly wrong) extension.
    if is_zip:
        return _read_xlsx(content)
    if is_ole:
        return _read_xls(content)
    if name.endswith(".xlsx") or "spreadsheetml" in (content_type or ""):
        return _read_xlsx(content)
    if name.endswith(".xls"):
        return _read_xls(content)
    if name.endswith(".csv") or "csv" in (content_type or ""):
        return _read_csv(content)
    raise UnsupportedFormat(
        "Unsupported file type. Upload a .xlsx, .xls, .xlsb or .csv brand PT file."
    )


# ----------------------------------------------------------------------------- detection


def score_header(row: list[Any]) -> float:
    s = 0.0
    for c in row:
        cv = raw_str(c).lower()
        if not cv:
            continue
        if cv in HEADER_KEYWORDS:
            s += 1
        elif len(cv) < 40 and any(k in cv for k in HEADER_KEYWORDS):
            s += 0.5
    return s


def detect_header(rows: list[list[Any]]) -> int:
    """The index of the row that reads most like a header (within the first 25)."""
    best_i, best = 0, -1.0
    for i in range(min(25, len(rows))):  # printed-invoice item tables can start deep
        sc = score_header(rows[i])
        if sc > best:
            best, best_i = sc, i
    return best_i


def _has_data(rows: list[list[Any]]) -> bool:
    return any(any(raw_str(c) for c in r) for r in rows[:30])


def _work_sheet_rows(rows: list[list[Any]]) -> int:
    """How many item rows a KDPS work sheet holds (0: not a work sheet, or empty).

    Its formula columns carry values on every row, so only a row with a barcode or
    a design counts as filled.
    """
    header_idx = detect_header(rows)
    headers = [norm(c) for c in rows[header_idx]] if rows else []
    if not all(h in headers for h in WORK_SHEET_HEADERS):
        return 0
    barcode = headers.index("BARCODE")
    design = headers.index("DESIGN") if "DESIGN" in headers else -1
    return sum(
        1
        for r in rows[header_idx + 1 :]
        if (barcode < len(r) and raw_str(r[barcode]))
        or (0 <= design < len(r) and raw_str(r[design]))
    )


def choose_sheet(sheets: list[tuple[str, list[list[Any]]]]) -> tuple[str, list[list[Any]]]:
    # The KDPS PT file: one staff work sheet, never its Master Sheet. Two filled work
    # sheets are two PTs, so the file is refused rather than one picked silently.
    work = [
        (name, rows, filled)
        for name, rows in sheets
        if norm(name) != MASTER_SHEET_NAME
        for filled in [_work_sheet_rows(rows)]
        if filled
    ]
    if len(work) > 1:
        names = ", ".join(name.strip() for name, _rows, _n in work)
        raise UnsupportedFormat(
            f"This file has {len(work)} filled work sheets ({names}). "
            "Upload one work sheet at a time."
        )
    if work:
        return work[0][0], work[0][1]
    # Prefer a sheet whose name matches a profile's sheet_contains token.
    for p in PROFILES:
        sc = p["match"].get("sheet_contains")
        if sc:
            for name, rows in sheets:
                if sc in name.upper() and _has_data(rows):
                    return name, rows
    # Else the non-empty sheet with the best-scoring header row.
    best = None
    best_score = -1.0
    for name, rows in sheets:
        if not _has_data(rows):
            continue
        hi = detect_header(rows)
        sc = score_header(rows[hi]) + min(len(rows), 100) * 0.001
        if sc > best_score:
            best_score, best = sc, (name, rows)
    return best or sheets[0]


def profile_by_code(code: str) -> dict[str, Any]:
    for p in PROFILES:
        if p["code"] == code:
            return p
    return GENERIC_PROFILE


def identify_profile(filename: str, header_set: set[str], sheet_name: str) -> dict[str, Any]:
    """The brand-file profile: a filename hint whose columns fit, else a header fingerprint."""
    up = (filename or "").upper()
    for kw, code in FILENAME_HINTS.items():
        if kw in up:
            p = profile_by_code(code)
            need = p.get("match", {}).get("header_has")
            # A filename hint counts only if the file's columns fit that profile.
            if not need or all(h in header_set for h in need):
                return p
    for p in PROFILES:
        m = p["match"]
        if not m:
            continue
        if m.get("sheet_contains") and m["sheet_contains"] not in (sheet_name or "").upper():
            continue
        if m.get("header_has") and not all(h in header_set for h in m["header_has"]):
            continue
        return p
    return GENERIC_PROFILE


def build_records(rows: list[list[Any]], header_idx: int) -> list[tuple[int, dict[str, Any]]]:
    """(1-based sheet row, {HEADER: value}) for every data row under the header."""
    headers = [norm(c) for c in rows[header_idx]]
    records: list[tuple[int, dict[str, Any]]] = []
    blanks_in_a_row = 0
    for offset, r in enumerate(rows[header_idx + 1 :]):
        if not any(raw_str(c) for c in r):
            blanks_in_a_row += 1
            if blanks_in_a_row >= 3:
                break
            continue
        blanks_in_a_row = 0
        first = raw_str(r[0]).lower() if r else ""
        if first.startswith("total") or first.startswith("grand total"):
            break
        rec: dict[str, Any] = {}
        for j, h in enumerate(headers):
            if h and j < len(r):
                rec.setdefault(h, r[j])
        records.append((header_idx + offset + 2, rec))
    return records


def read_brand_file(content: bytes, filename: str, content_type: str = "") -> BrandFile:
    """Read a brand file: its sheet, header row, profile and data records."""
    sheets = read_sheets(content, filename, content_type)
    if not sheets:
        raise UnsupportedFormat("This file has no sheets.")
    sheet_name, rows = choose_sheet(sheets)
    if not any(any(raw_str(c) for c in r) for r in rows):
        raise UnsupportedFormat("This file has no data rows to map.")
    header_idx = detect_header(rows)
    headers = [norm(c) for c in rows[header_idx]]
    profile = identify_profile(filename, {h for h in headers if h}, sheet_name)
    # A brand guess from the file name (alpha words only) for files with no brand column.
    stem = (filename or "").rsplit(".", 1)[0]
    guess = norm(" ".join(t for t in re.split(r"[ _\-]+", stem) if t.isalpha()))
    return BrandFile(
        filename=filename,
        sheet=sheet_name,
        header_row=header_idx,
        headers=[h for h in headers if h],
        profile=profile,
        records=build_records(rows, header_idx),
        truncated=len(rows) > MAX_ROWS,
        brand_guess=guess,
    )


# ----------------------------------------------------------------------------- mapping


def _pick(rec: dict[str, Any], role: str, profile: dict[str, Any]) -> Any:
    override = profile.get("overrides", {}).get(role)
    if override:
        if isinstance(override, list):
            return [rec.get(norm(c)) for c in override]
        return rec.get(norm(override))
    for alias in ALIASES.get(role, []):
        if alias in rec and raw_str(rec[alias]):
            return rec[alias]
    return None


def _description(rec: dict[str, Any], profile: dict[str, Any]) -> str:
    cols = profile.get("overrides", {}).get("DESC_SRC") or ALIASES["DESC_SRC"]
    parts = [raw_str(rec.get(norm(c))) for c in cols]
    return " ".join(p for p in parts if p)


def _hit(value: RuleValue, origin: str, source: str, **extra: Any) -> MappedCell:
    return MappedCell(
        value=value.value_key if value.dimension != "brand" else value.id,
        origin=origin,
        source=source,
        value_id=value.id,
        label=value.label,
        **extra,
    )


def _from_rule(rule: Rule, source: str, reason: str = CONFIRMED_RULE) -> MappedCell:
    return _hit(rule.target, RULE, source, rule_id=rule.id, reason=reason)


def _offer(source: str, candidates: Sequence[Candidate]) -> MappedCell:
    if not candidates:
        return MappedCell(value=None, origin=NONE, source=source)
    return MappedCell(
        value=None,
        origin=SUGGESTION,
        source=source,
        suggestions=tuple(
            Suggestion(
                value_id=c.value.id,
                value=c.value.value_key if c.value.dimension != "brand" else c.value.id,
                label=c.value.label,
                reason=c.reason,
            )
            for c in candidates
        ),
    )


def _conflicts(values: Sequence[RuleValue]) -> list[Candidate]:
    return [Candidate(value=v, reason=CONFLICTING_RULES, score=1.0) for v in values]


def _merge(*groups: Sequence[Candidate]) -> list[Candidate]:
    seen: set[str] = set()
    out: list[Candidate] = []
    for group in groups:
        for candidate in group:
            if candidate.value.id not in seen:
                seen.add(candidate.value.id)
                out.append(candidate)
    return out


class _Resolution:
    """A cell's value, or the suggestions collected on the way to not finding one."""

    def __init__(self, cell: MappedCell | None, offers: list[Candidate]) -> None:
        self.cell = cell
        self.offers = offers


def _attribute(
    rulebook: Rulebook,
    dimension: str,
    raw: str,
    cleaned: Sequence[str],
    issuers: Sequence[str],
) -> _Resolution:
    """The file's own text for one column, through the rulebook.

    Order: the text as it stands is an approved value (file) -> a mechanical
    clean-up of it is (rule, ``cleaned``) -> a confirmed rule for the text or a
    clean-up, brand rules before ``*`` (rule) -> close matches (suggestions).
    """
    raw = raw.strip()
    if not raw or not rulebook.governed(dimension):
        return _Resolution(None, [])
    value = rulebook.exact(dimension, raw)
    if value is not None:
        return _Resolution(_hit(value, FILE, raw), [])
    for text in cleaned:
        value = rulebook.exact(dimension, text)
        if value is not None:
            return _Resolution(_hit(value, RULE, raw, reason=CLEANED), [])
    conflicts: list[Candidate] = []
    for text in (raw, *cleaned):
        rule, clash = rulebook.lookup(dimension, text, issuers)
        if rule is not None:
            return _Resolution(_from_rule(rule, raw), [])
        conflicts.extend(_conflicts(clash))
    offers = _merge(conflicts, close_matches([raw, *cleaned], rulebook.choices(dimension)))
    return _Resolution(None, offers)


def _settle(resolution: _Resolution, source: str) -> MappedCell:
    if resolution.cell is not None:
        return resolution.cell
    return _offer(source, resolution.offers)


def _free(text: str, reason: str | None = None) -> MappedCell:
    return (
        MappedCell(value=text, origin=FILE, source=text, reason=reason)
        if text
        else MappedCell(value=None, origin=NONE)
    )


@dataclass
class _Row:
    """The working state of one row while its columns are mapped."""

    rec: dict[str, Any]
    profile: dict[str, Any]
    rulebook: Rulebook
    context: MappingContext
    description: str
    cells: dict[str, MappedCell] = field(default_factory=dict)
    item_rule: Rule | None = None

    def get(self, role: str) -> Any:
        return _pick(self.rec, role, self.profile)

    def text(self, role: str) -> str:
        return raw_str(self.get(role))


def _map_brand(row: _Row, brand_guess: str) -> tuple[MappedCell, int | None]:
    rulebook, context = row.rulebook, row.context
    cell_text = row.text("BRAND_SRC")
    profile_brand = raw_str(row.profile.get("brand_const"))
    raw = cell_text or profile_brand or ("" if context.brand_id is not None else brand_guess)
    if raw:
        value = rulebook.exact("brand", raw)
        if value is not None:
            origin, reason = (FILE, None) if cell_text else (RULE, PROFILE_BRAND)
            return _hit(value, origin, raw, reason=reason), int(value.id)
        rule, clash = rulebook.lookup(
            "brand", raw, issuers_for(None, issuer_key=context.issuer_key)
        )
        if rule is not None:
            return _from_rule(rule, raw), int(rule.target.id)
        if cell_text or profile_brand:
            # The file names a brand nobody has mapped: a person decides. The
            # receipt's brand is not put in its place, nor are its rules applied.
            offers = _merge(_conflicts(clash), close_matches([raw], rulebook.brands))
            return _offer(raw, offers), None
    if context.brand_id is not None:
        for brand in rulebook.brands:
            if brand.id == str(context.brand_id):
                return _hit(brand, RULE, raw, reason=RECEIPT_BRAND), context.brand_id
    return MappedCell(value=None, origin=NONE, source=raw), context.brand_id


def _map_season(row: _Row, issuers: Sequence[str]) -> MappedCell:
    """SEASON: an explicit season wins - a season is a name, never a date - then the
    invoice date's season, then the receipt date's."""
    rulebook = row.rulebook
    code = row.text("SEASON_SRC")
    found = _attribute(rulebook, "season", code, [], issuers)
    if found.cell is not None:
        return found.cell
    raw_date = row.get("DATE_SRC")
    for when, reason in (
        (parse_date(raw_date), INVOICE_DATE),
        (row.context.invoice_date, RECEIPT_DATE),
    ):
        if when is None:
            continue
        value = rulebook.exact("season", season_label(when))
        if value is not None:
            return _hit(value, RULE, code or raw_str(raw_date), reason=reason)
    return _offer(code, found.offers)


def _map_item(row: _Row, issuers: Sequence[str]) -> MappedCell:
    """ITEM: the longest confirmed keyword found as whole words in the description."""
    rulebook, desc = row.rulebook, row.description
    if not desc or not rulebook.governed("item"):
        return MappedCell(value=None, origin=NONE, source=desc)
    rule, clash = rulebook.keyword("item", desc, issuers)
    if rule is not None:
        row.item_rule = rule
        return _from_rule(rule, desc, reason=DESCRIPTION_WORD)
    named = [
        Candidate(value=v, reason=NAMED_IN_TEXT, score=0.8) for v in rulebook.named_in("item", desc)
    ]
    return _offer(desc, _merge(_conflicts(clash), named))


def _keyword_default(row: _Row, dimension: str) -> MappedCell | None:
    """The default the item keyword carries for another column (``keyword`` issuer)."""
    if row.item_rule is None:
        return None
    rule, _clash = row.rulebook.lookup(dimension, row.item_rule.source_key, [KEYWORD_ISSUER])
    return _from_rule(rule, row.description, reason=KEYWORD_DEFAULT) if rule else None


def _map_gender(row: _Row, issuers: Sequence[str], brand_id: int | None) -> MappedCell:
    """GENDER: the gender column, then a confirmed gender word (or an approved
    gender named outright) in the description, then the item keyword's default,
    then the brand's default. A gender word no rule confirms ('LADIES' with no
    rule for it) is only ever a suggestion."""
    rulebook, desc = row.rulebook, row.description
    raw = row.text("GENDER_SRC")
    found = _attribute(rulebook, "gender", raw, [], issuers)
    if found.cell is not None:
        return found.cell
    if not rulebook.governed("gender"):
        return MappedCell(value=None, origin=NONE, source=raw)
    offers = list(found.offers)
    if desc:
        rule, clash = rulebook.keyword("gender", desc, issuers)
        if rule is not None:
            return _from_rule(rule, desc, reason=DESCRIPTION_WORD)
        offers = _merge(offers, _conflicts(clash))
        named = rulebook.named_in("gender", desc)
        longest = max((len(v.value_key) for v in named), default=0)
        spoken = [v for v in named if len(v.value_key) == longest]
        if len(spoken) == 1:
            return _hit(spoken[0], RULE, desc, reason=DESCRIPTION_WORD)
    default = _keyword_default(row, "gender")
    if default is not None:
        return default
    if brand_id is not None:
        rule, _clash = rulebook.lookup("gender", DEFAULT_SOURCE, [brand_issuer(brand_id)])
        if rule is not None:
            return _from_rule(rule, raw or desc, reason=BRAND_DEFAULT)
    inferred = gender_from_text(desc) if desc else ""
    value = rulebook.exact("gender", inferred) if inferred else None
    if value is not None:
        offers = _merge(offers, [Candidate(value=value, reason=NAMED_IN_TEXT, score=0.8)])
    return _offer(raw or desc, offers)


def _map_fit(row: _Row, issuers: Sequence[str]) -> MappedCell:
    raw = row.text("FIT_SRC")
    coded = bool(row.profile.get("flags", {}).get("fit_code_tokens"))
    cleaned = fit_code_candidates(raw)[1:] if coded else []
    found = _attribute(row.rulebook, "fit", raw, cleaned, issuers)
    if found.cell is not None:
        return found.cell
    default = _keyword_default(row, "fit")
    return default if default is not None else _offer(raw, found.offers)


def _item_helper(row: _Row, dimension: str) -> MappedCell | None:
    item = row.cells.get("ITEM")
    if item is None or item.value is None:
        return None
    rule, _clash = row.rulebook.lookup(dimension, item.value, [ITEM_ISSUER])
    return _from_rule(rule, item.value, reason=ITEM_HELPER) if rule else None


def _map_category(row: _Row, dimension: str) -> tuple[MappedCell, MappedCell]:
    """(SUB CATEGORY or TYPE, its SUGGESTED column): the item keyword's own value,
    else the master's ITEM helper; the SUGGESTED column is always the helper."""
    helper = _item_helper(row, dimension)
    empty = MappedCell(value=None, origin=NONE)
    value = _keyword_default(row, dimension) or helper or empty
    return value, helper or empty


def _map_explicit(
    row: _Row, dimension: str, role: str, issuers: Sequence[str]
) -> MappedCell | None:
    """A column the file fills itself (a KDPS work sheet): its value, a suggestion for
    text no rule knows, or None for a blank cell (the caller's fallback applies)."""
    raw = row.text(role)
    if not raw:
        return None
    return _settle(_attribute(row.rulebook, dimension, raw, [], issuers), raw)


def _basic(row: _Row, qty: str) -> MappedCell:
    if row.profile.get("flags", {}).get("basic_from_taxable_per_unit"):
        taxable = number_text(row.rec.get("TAXABLE_AMOUNT"))
        if taxable and qty and Decimal(qty) != 0:
            per_unit = (Decimal(taxable) / Decimal(qty)).quantize(Decimal("0.01"))
            return _free(number_text(per_unit), reason="taxable_amount_per_unit")
    return _free(number_text(row.get("BASIC_SRC")))


def map_row(
    rec: dict[str, Any],
    profile: dict[str, Any],
    rulebook: Rulebook,
    context: MappingContext | None = None,
    *,
    line_no: int = 1,
    source_row: int = 0,
    brand_guess: str = "",
) -> MappedRow | None:
    """Map one brand-file record to the twenty-two KDPS columns, each with its origin.

    ``rec`` is ``{HEADER: value}`` with upper-cased headers, as
    :func:`read_brand_file` gives it. None when the row carries no item (no
    barcode or design) or no quantity or price - a note or a subtotal line.
    """
    context = context or MappingContext()
    row = _Row(
        rec=rec,
        profile=profile,
        rulebook=rulebook,
        context=context,
        description=_description(rec, profile),
    )
    barcode = clean_code(row.get("BARCODE"))
    design = row.text("DESIGN")
    qty = number_text(row.get("QTY"))
    mrp = number_text(row.get("MRP"))
    if not (barcode or design) or not (qty or mrp):
        return None
    brand, brand_id = _map_brand(row, brand_guess)
    issuers = issuers_for(brand_id)
    cells = row.cells
    cells["SEASON"] = _map_season(row, issuers)
    cells["BRAND"] = brand
    colour_raw = color_source(row.get("COLOR_SRC"), profile)
    cells["COLOR"] = _settle(
        _attribute(
            rulebook, "colour", colour_raw, clean_for_dimension("colour", colour_raw), issuers
        ),
        colour_raw,
    )
    explicit = bool(profile.get("flags", {}).get("explicit_attributes"))
    named_item = _map_explicit(row, "item", "ITEM_SRC", issuers) if explicit else None
    cells["ITEM"] = named_item or _map_item(row, issuers)
    cells["GENDER"] = _map_gender(row, issuers, brand_id)
    sub, suggested_sub = _map_category(row, "sub_category")
    kind, suggested_type = _map_category(row, "type")
    if explicit:
        # A work sheet names its SUB CATEGORY and TYPE; the ITEM's suggestion only
        # fills a blank one (and is still shown in its SUGGESTED column).
        sub = _map_explicit(row, "sub_category", "SUBCAT_SRC", issuers) or sub
        kind = _map_explicit(row, "type", "TYPE_SRC", issuers) or kind
    cells["SUB CATEGORY"] = sub
    cells["TYPE"] = kind
    cells["FIT"] = _map_fit(row, issuers)
    size_raw = clean_code(row.get("SIZE_SRC"))
    cells["SIZE"] = _settle(
        _attribute(rulebook, "size", size_raw, clean_for_dimension("size", size_raw), issuers),
        size_raw,
    )
    cells["BARCODE"] = _free(barcode)
    cells["DESIGN"] = _free(design)
    cells["HSN"] = _free(clean_code(row.get("HSN")))
    cells["QTY"] = _free(qty)
    cells["MRP"] = _free(mrp)
    cells["BASIC"] = _basic(row, qty)
    # A brand file's rate and tax are checks against the product's own calculation.
    cells["P RATE"] = _free(number_text(row.get("PRATE_SRC")))
    cells["INPUT TAX"] = _free(number_text(row.get("TAX_SRC")))
    for derived in ("OUTPUT TAX", "NAG", "MARGIN"):
        cells[derived] = MappedCell(value=None, origin=NONE)
    cells["SUGGESTED SUB CATEGORY"] = suggested_sub
    cells["SUGGESTED TYPE"] = suggested_type
    return MappedRow(
        line_no=line_no,
        source_row=source_row,
        cells={column: cells[column] for column in KDPS_COLUMNS},
    )


def map_file(
    content: bytes,
    filename: str,
    content_type: str,
    rulebook: Rulebook,
    context: MappingContext | None = None,
) -> MappedFile:
    """Read a brand file and map every item row (see :func:`map_row`). Writes nothing."""
    brand_file = read_brand_file(content, filename, content_type)
    return map_records(brand_file, rulebook, context)


def map_records(
    brand_file: BrandFile, rulebook: Rulebook, context: MappingContext | None = None
) -> MappedFile:
    profile = brand_file.profile
    rows: list[MappedRow] = []
    for source_row, rec in brand_file.records:
        mapped = map_row(
            rec,
            profile,
            rulebook,
            context,
            line_no=len(rows) + 1,
            source_row=source_row,
            brand_guess=brand_file.brand_guess,
        )
        if mapped is not None:
            rows.append(mapped)
    return MappedFile(
        profile_code=str(profile["code"]),
        profile_name=str(profile["name"]),
        archetype=str(profile["archetype"]),
        sheet=brand_file.sheet,
        header_row=brand_file.header_row,
        headers=brand_file.headers,
        source_rows=len(brand_file.records),
        truncated=brand_file.truncated,
        rows=rows,
    )
