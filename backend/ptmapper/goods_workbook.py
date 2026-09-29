"""Exact numbers in goods-v1 XLSX files (overall PRD §15.2.1 rule 6).

openpyxl formats every number it writes through a float, so a money cell can gain
digits on the way out; here a money cell is written as its exact decimal text.

An XLSX numeric cell *is* a binary double (xsd:double); its stored text is only one
spelling of it, and writers differ (``95.24`` or ``95.23999999999999`` for the same
typed 95.24). A numeric cell is therefore read as the ``Decimal`` of the shortest
decimal that identifies that double - exactly what was typed or generated, with no
arithmetic done in floats. A whole number is read from its digits directly. Callers
then apply their own precision rule to that decimal. Formulas are never evaluated:
only cached values are read.
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from decimal import Decimal
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet._reader import WorkSheetParser
from openpyxl.xml.constants import SHEET_MAIN_NS

from core.goods_money import amount_text

_VALUE_TAG = f"{{{SHEET_MAIN_NS}}}v"


@dataclass(frozen=True)
class Amount:
    """A money cell to write: integer paise, stored as exact rupee text."""

    paise: int


class _ExactNumbers(WorkSheetParser):
    """openpyxl's sheet parser, keeping a numeric cell's stored text as a ``Decimal``."""

    def parse_cell(self, element: Any) -> dict[str, Any]:
        cell: dict[str, Any] = super().parse_cell(element)
        if cell["data_type"] == "n" and cell["value"] is not None:
            cell["value"] = stored_number(element.findtext(_VALUE_TAG))
        return cell


def stored_number(text: str) -> Decimal:
    """The decimal a numeric cell holds: its digits, or the shortest form of its double."""
    if text.lstrip("-").isdigit():
        return Decimal(text)
    return Decimal(repr(float(text)))


def read_rows(data: bytes, *, max_rows: int | None = None) -> Iterator[tuple[int, tuple[Any, ...]]]:
    """``(row number, values)`` for each stored row of the first sheet, numbers exact.

    Stops after ``max_rows`` rows, so an oversized file is never read in full.

    Raises whatever openpyxl raises for an unreadable file; callers map it to their
    own refusal.
    """
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        # types-openpyxl omits these private reader internals. This parser is
        # intentionally coupled to the pinned openpyxl implementation.
        with sheet._get_source() as source:  # type: ignore[attr-defined]
            parser = _ExactNumbers(
                source,
                sheet._shared_strings,  # type: ignore[attr-defined]
                data_only=True,
                epoch=workbook.epoch,
                date_formats=workbook._date_formats,  # type: ignore[attr-defined]
                timedelta_formats=workbook._timedelta_formats,  # type: ignore[attr-defined]
            )
            for count, (number, cells) in enumerate(parser.parse(), start=1):
                if max_rows is not None and count > max_rows:
                    return
                values: list[Any] = [None] * max((cell["column"] for cell in cells), default=0)
                for cell in cells:
                    values[cell["column"] - 1] = cell["value"]
                yield number, tuple(values)
    finally:
        workbook.close()


def workbook_bytes(title: str, rows: Iterable[list[Any]]) -> bytes:
    """A one-sheet workbook; ``Amount`` cells hold exact numeric text, ``None`` stays empty."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title
    for row_number, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            if value is None:
                continue
            cell = sheet.cell(row=row_number, column=column)
            if isinstance(value, Amount):
                assert isinstance(cell, Cell)
                cell.value = 0  # a numeric cell ...
                cell._value = amount_text(value.paise)  # type: ignore[attr-defined]  # exact stored text
                cell.number_format = "0.00"
            else:
                cell.value = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
