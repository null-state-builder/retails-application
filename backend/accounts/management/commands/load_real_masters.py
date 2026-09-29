"""Load real KDPS master data — legal entity, GSTINs, all seventeen sites,
vendors and brands — into an already-bootstrapped, non-synthetic tenant.

Run: `python manage.py load_real_masters --tenant-code <code>` (create that
tenant first with `bootstrap_deployment`). Idempotent — safe to re-run.

Every site is born on the goods-v1 stock contract, with its protected system
locations, and nothing else: it stays `planned` with no readiness approved, the
way a site created on the Organisation screen is born (store and warehouse
operations PRD §3.1, ticket OPS-13). Opening it for goods is the readiness
approval's job, not a load's — and real opening stock stays blocked by OQ-54
whatever this command does. The load never *converts* a site: one that already
sits on the legacy contract or already has legacy stock history refuses the
whole load, because the goods-v1 design (§5.4) moves such a site only through a
separately approved reconciliation record.

Source files under `docs/data-from-kdps/`, and what each one does NOT give us:
  - `Q&A-req-recieved/LIST OF ALL STORES.xlsx` — store name + location only.
    No state, no GSTIN, no store/warehouse distinction — `CONFIRMED_SITES`
    below supplies those, and every entry is checked against a matching row
    in this file; the whole load refuses if one is missing rather than
    guessing. The file's EMAIL column is never read.
  - `Q&A-req-recieved/SUPPLIER BRAND DETAILS.xlsx` — brand -> supplier name
    pairs only. No vendor GSTIN, address or payment terms, and no brand
    commercial model (ownership / return terms) — those are left at the
    model's own default and are NOT a confirmed business decision.

Neither real GSTINs nor real brand commercial terms exist in the source data
we have. See `docs/features/data-migration/migration-plan.md` §7 for both as
open questions for Anand — this command does not invent them, it loads a
clearly-marked placeholder and says so loudly on every run.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from masters.goods_models import SiteGuard, Tenant
from masters.goods_services import ensure_system_locations
from masters.models import Brand, Gstin, LegalEntity, Store
from vendors.models import Vendor

#: Deliberately not a real GSTIN format/checksum, so nobody mistakes this row
#: for a real registration.
PLACEHOLDER_GSTIN = {
    "bihar": "PENDINGGSTINBH",
    "jharkhand": "PENDINGGSTINJH",
}

#: The state follows the site's LOCATION (Anand, 23 September 2026). The store
#: list itself names no state.
STATE_BY_LOCATION = {
    "BHAGALPUR": "bihar",
    "KAHELGAON": "bihar",
    "BANKA": "bihar",
    "GAYA": "bihar",
    "KANKARBAGH": "bihar",
    "SAHEBGUNJ": "jharkhand",
    "DUMKA": "jharkhand",
    "BOKARO": "jharkhand",
    "SINGH MORE": "jharkhand",
    "HAZARIBAGH": "jharkhand",
    "DEOGARH": "jharkhand",
    "RANCHI": "jharkhand",
}


@dataclass(frozen=True)
class SiteSpec:
    code: str
    store_type: str
    display_name: str
    sheet_name: str  # STORE NAME exactly as it appears in the store-list file
    sheet_location: str  # LOCATION exactly as it appears there

    @property
    def state(self) -> str:
        """The state, "bihar" or "jharkhand", read off the location - never stated per site."""
        return STATE_BY_LOCATION[self.sheet_location]


_STORE = Store.StoreType.STORE

#: Every row of the store list, in its own SL NO order: sixteen stores and the
#: Ranchi warehouse. The first five codes (DEO, DUM, BANKA, BGP, RAN-WH) were
#: the 21 September 2026 pilot and keep their codes; the other twelve were added
#: on 23 September 2026, when every site moved to goods-v1 (OPS-13). Where one
#: location has several stores, each store beyond the pilot's own carries a brand
#: suffix (DEO-LEE, BANKA-JKY); Kankarbagh had no pilot store, so both of its do.
CONFIRMED_SITES: tuple[SiteSpec, ...] = (
    SiteSpec("BGP", _STORE, "Vaishnavi Bhagalpur", "VAISHNAVI", "BHAGALPUR"),  # 1
    SiteSpec("KHG", _STORE, "Vaishnavi Kahelgaon", "VAISHNAVI", "KAHELGAON"),  # 2
    SiteSpec("SBG", _STORE, "MBO Sahebgunj", "MBO", "SAHEBGUNJ"),  # 3
    SiteSpec("DUM", _STORE, "MBO Dumka", "MBO", "DUMKA"),  # 4
    SiteSpec("BKR", _STORE, "MBO Bokaro", "MBO", "BOKARO"),  # 5
    SiteSpec("SGM", _STORE, "Vaishnavi Singh More", "VAISHNAVI", "SINGH MORE"),  # 6
    SiteSpec(
        "HZB", _STORE, "Jainsons Lifestyle Hazaribagh", "JAINSONS-LIFESTYLE", "HAZARIBAGH"
    ),  # 7
    SiteSpec("DEO", _STORE, "Vaishnavi Deoghar", "VAISHNAVI", "DEOGARH"),  # 8
    SiteSpec("DEO-LEE", _STORE, "Lee Deoghar", "LEE", "DEOGARH"),  # 9
    SiteSpec("DEO-AS", _STORE, "Allen Solly Deoghar", "ALLEN SOLLY", "DEOGARH"),  # 10
    SiteSpec("DEO-SPY", _STORE, "Spykar Deoghar", "SPYKAR", "DEOGARH"),  # 11
    SiteSpec("BANKA", _STORE, "Vaishnavi Banka", "VAISHNAVI", "BANKA"),  # 12
    SiteSpec("BANKA-JKY", _STORE, "Jockey Banka", "JOCKEY", "BANKA"),  # 13
    SiteSpec("GAYA", _STORE, "MBO Gaya", "MBO", "GAYA"),  # 14
    SiteSpec("KKB-PE", _STORE, "Peter England Kankarbagh", "PETER ENGLAND", "KANKARBAGH"),  # 15
    SiteSpec("KKB-MBO", _STORE, "MBO Kankarbagh", "MBO", "KANKARBAGH"),  # 16
    SiteSpec(
        "RAN-WH", Store.StoreType.WAREHOUSE, "Ranchi Central Warehouse", "WAREHOUSE", "RANCHI"
    ),  # 17
)


def _slugify(value: str, max_len: int) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.strip().lower()).strip("-")
    return slug[:max_len] or "x"


def _unique_slug(value: str, max_len: int, taken: set[str], warnings: list[str]) -> str:
    base = _slugify(value, max_len)
    if base not in taken:
        taken.add(base)
        return base
    n = 2
    while True:
        suffix = f"-{n}"
        slug = base[: max_len - len(suffix)] + suffix
        if slug not in taken:
            taken.add(slug)
            warnings.append(f"{value!r} collided on code {base!r}, assigned {slug!r} instead.")
            return slug
        n += 1


def _load_store_names(path: Path) -> dict[tuple[str, str], str]:
    """(STORE NAME, LOCATION), both upper-cased, -> STORE NAME as written."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["STORE LIST"]
    out: dict[tuple[str, str], str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, location = row[1], row[2]
        if not name or not location:
            continue
        key = (str(name).strip().upper(), str(location).strip().upper())
        out[key] = str(name).strip()
    return out


def _load_brand_rows(path: Path) -> tuple[list[tuple[str, str]], list[str]]:
    """Returns (brand, supplier) pairs, plus brand names whose supplier cell is
    blank — those still get a Brand row, just no vendor link (§ brand w/o supplier)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["BRAND"]
    pairs: list[tuple[str, str]] = []
    unlinked_brands: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        brand, supplier = row[0], row[1]
        if not brand:
            continue
        brand = str(brand).strip()
        if supplier:
            pairs.append((brand, str(supplier).strip()))
        else:
            unlinked_brands.append(brand)
    return pairs, unlinked_brands


def _legacy_history_codes(codes: list[str]) -> list[str]:
    """Existing sites this load would *convert* to goods-v1 rather than create on it.

    A site already on goods-v1 is simply there, and a site with no guard and no
    legacy stock (say, one an earlier five-site run of this command made) has
    nothing to convert. A site whose guard is on the legacy contract, or that
    holds legacy stock rows, has legacy operational history - and moving it is
    a separately approved reconciliation, never a load's side effect.
    """
    from stockledger.models import StockLedgerEntry

    out = []
    for store in Store.objects.filter(code__in=codes).order_by("code"):
        guard = SiteGuard.objects.filter(site=store).first()
        if guard is not None and guard.stock_contract == SiteGuard.StockContract.GOODS_V1:
            continue
        if guard is not None or StockLedgerEntry.objects.filter(store=store).exists():
            out.append(store.code)
    return out


class Command(BaseCommand):
    help = (
        "Load real KDPS masters (legal entity, GSTINs, all seventeen sites on the "
        "goods-v1 stock contract, vendors and brands) into an already-bootstrapped "
        "real tenant. Idempotent."
    )

    def add_arguments(self, parser: Any) -> None:
        repo_root = Path(settings.BASE_DIR).parent.parent
        data = repo_root / "docs" / "data-from-kdps"
        parser.add_argument(
            "--tenant-code",
            required=True,
            help="Code of the real tenant to load into (created by bootstrap_deployment).",
        )
        parser.add_argument(
            "--stores-path", default=str(data / "Q&A-req-recieved" / "LIST OF ALL STORES.xlsx")
        )
        parser.add_argument(
            "--vendors-path",
            default=str(data / "Q&A-req-recieved" / "SUPPLIER BRAND DETAILS.xlsx"),
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Parse and validate the source files only; write nothing to the database.",
        )

    def handle(self, *args: Any, **opts: Any) -> None:
        from core.tenancy import tenant_context

        stores_path = Path(opts["stores_path"])
        vendors_path = Path(opts["vendors_path"])
        if not stores_path.exists():
            raise CommandError(f"Store list not found at {stores_path}")
        if not vendors_path.exists():
            raise CommandError(f"Vendor/brand list not found at {vendors_path}")

        sheet_names = _load_store_names(stores_path)
        missing = [
            s for s in CONFIRMED_SITES if (s.sheet_name, s.sheet_location) not in sheet_names
        ]
        if missing:
            raise CommandError(
                "These confirmed sites are not in the store list file — refusing the "
                "whole load rather than guessing: "
                + ", ".join(f"{s.code} ({s.sheet_name}/{s.sheet_location})" for s in missing)
            )
        # A row the file gained is not guessed into a site either - it is named, so
        # it cannot be dropped in silence.
        confirmed = {(s.sheet_name, s.sheet_location) for s in CONFIRMED_SITES}
        unlisted = sorted(key for key in sheet_names if key not in confirmed)
        unlisted_note = (
            "Rows in the store list that are not confirmed sites, so NOT loaded: "
            + ", ".join(f"{name}/{location}" for name, location in unlisted)
            if unlisted
            else ""
        )

        pairs, unlinked_brands = _load_brand_rows(vendors_path)
        if not pairs and not unlinked_brands:
            raise CommandError(f"No brand rows found in {vendors_path}")

        if opts["dry_run"]:
            brand_names = {b for b, _ in pairs} | set(unlinked_brands)
            supplier_names = {s for _, s in pairs}
            if unlisted_note:
                self.stdout.write(self.style.WARNING(unlisted_note))
            self.stdout.write(
                self.style.SUCCESS(
                    f"Dry run OK: all {len(CONFIRMED_SITES)} confirmed sites found in the "
                    f"store list; {len(brand_names)} distinct brands ({len(unlinked_brands)} "
                    f"with no supplier named) / {len(supplier_names)} distinct suppliers "
                    "parsed. Nothing written."
                )
            )
            return

        try:
            tenant = Tenant.objects.get(code=opts["tenant_code"])
        except Tenant.DoesNotExist as exc:
            raise CommandError(
                f"No tenant with code {opts['tenant_code']!r}. Run bootstrap_deployment first."
            ) from exc
        if tenant.synthetic:
            raise CommandError(
                f"Tenant {tenant.code!r} is marked synthetic — refusing to load real "
                "master data onto a demo tenant."
            )

        warnings: list[str] = [unlisted_note] if unlisted_note else []
        with tenant_context(tenant.pk):
            converting = _legacy_history_codes([s.code for s in CONFIRMED_SITES])
            if converting:
                raise CommandError(
                    "These sites already exist with legacy stock history or on the legacy "
                    "stock contract. Loading would convert them to goods-v1, which needs "
                    "its own approved reconciliation, not a master load — refusing the "
                    "whole load: " + ", ".join(converting)
                )
            with transaction.atomic():
                _entity, gstins = self._load_entity_gstins()
                stores = self._load_stores(tenant, gstins)
                vendors, brands = self._load_vendors_brands(pairs, unlinked_brands, warnings)

        for w in warnings:
            self.stdout.write(self.style.WARNING(w))
        self.stdout.write(
            self.style.SUCCESS(
                f"Loaded into tenant {tenant.code!r}: 1 legal entity, {len(gstins)} GSTINs "
                f"(placeholder — not real), {len(stores)} sites on the goods-v1 stock "
                f"contract (planned, no readiness approved), {len(vendors)} vendors, "
                f"{len(brands)} brands."
            )
        )
        self.stdout.write(
            self.style.WARNING(
                "GSTINs are placeholder values — replace before any real document is "
                "raised. Every brand's ownership/return-terms is left at the model "
                "default (owned / no returns) — unconfirmed, not a real commercial term. "
                "See docs/features/data-migration/migration-plan.md §7."
            )
        )

    def _load_entity_gstins(self) -> tuple[LegalEntity, dict[str, Gstin]]:
        entity, _ = LegalEntity.objects.update_or_create(
            code="kdps", defaults={"name": "KDPS Lifestyle Pvt Ltd"}
        )
        bihar, _ = Gstin.objects.update_or_create(
            gstin=PLACEHOLDER_GSTIN["bihar"],
            defaults={"legal_entity": entity, "state_code": "10", "state_name": "Bihar"},
        )
        jhk, _ = Gstin.objects.update_or_create(
            gstin=PLACEHOLDER_GSTIN["jharkhand"],
            defaults={"legal_entity": entity, "state_code": "20", "state_name": "Jharkhand"},
        )
        return entity, {"bihar": bihar, "jharkhand": jhk}

    def _load_stores(self, tenant: Tenant, gstins: dict[str, Gstin]) -> dict[str, Store]:
        """Each site, its goods-v1 guard and its six protected system locations.

        The guard is created once and never rewritten: a re-run finds it and leaves
        it alone, so readiness somebody approved since the last load stays approved,
        and nothing here ever approves any. Lifecycle and every readiness flag keep
        the values a site is born with.
        """
        out: dict[str, Store] = {}
        for spec in CONFIRMED_SITES:
            store, _ = Store.objects.update_or_create(
                code=spec.code,
                defaults={
                    "name": spec.display_name,
                    "store_type": spec.store_type,
                    "gstin": gstins[spec.state],
                    "city": spec.sheet_location.title(),
                },
            )
            SiteGuard.objects.get_or_create(
                site=store,
                defaults={
                    "tenant_id": tenant.pk,
                    "lifecycle": SiteGuard.Lifecycle.PLANNED,
                    "stock_contract": SiteGuard.StockContract.GOODS_V1,
                },
            )
            ensure_system_locations(tenant.pk, store)
            out[spec.code] = store
        return out

    def _load_vendors_brands(
        self, pairs: list[tuple[str, str]], unlinked_brands: list[str], warnings: list[str]
    ) -> tuple[dict[str, Vendor], dict[str, Brand]]:
        by_supplier: dict[str, list[str]] = {}
        for brand_name, supplier_name in pairs:
            by_supplier.setdefault(supplier_name, []).append(brand_name)

        brand_slugs: set[str] = set()
        vendor_slugs: set[str] = set()
        brands: dict[str, Brand] = {}
        vendors: dict[str, Vendor] = {}

        if unlinked_brands:
            warnings.append(
                f"{len(unlinked_brands)} brands have no supplier named in the source file "
                f"— created with no vendor link: {', '.join(sorted(unlinked_brands))}"
            )
        for name in sorted({b for b, _ in pairs} | set(unlinked_brands)):
            code = _unique_slug(name, 32, brand_slugs, warnings)
            brand, _ = Brand.objects.update_or_create(code=code, defaults={"name": name[:120]})
            brands[name] = brand

        for supplier_name in sorted(by_supplier):
            code = _unique_slug(supplier_name, 32, vendor_slugs, warnings)
            vendor, _ = Vendor.objects.update_or_create(
                code=code, defaults={"name": supplier_name[:160]}
            )
            vendor.brands.set([brands[b] for b in by_supplier[supplier_name]])
            vendors[supplier_name] = vendor

        return vendors, brands
