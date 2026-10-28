"""Store test run: a store opened from its PT files, then brought to the legacy SOH.

The owner's process (28 Oct 2026): set up the store empty, upload its KDPS PT
workbooks, "Approve and load" (no scanning, cost = P RATE), then upload the
legacy SOH at the stock sync. Synthetic tenant only.
"""
from __future__ import annotations

import hashlib
import io
import uuid
from decimal import Decimal
from typing import Any

import pytest
from django.utils import timezone
from openpyxl import Workbook

from accounts.goods_demo import seed_working_calendar
from accounts.goods_setup import bootstrap_deployment
from accounts.models import User
from accounts.registration_services import deployment_key
from accounts.test_store_loader import Loader, OpeningDriver, Plan, access_for, act, opening_auto_in_force
from approvals.goods_models import ApprovalDecision
from core.kernel_models import JobState, OutboxIntent
from core.outbox import drain
from core.refusals import Refusal
from core.tenancy import tenant_context
from files.goods_services import stage_upload
from masters.models import Season, Store
from ptmapper import pt_opening_services as quick
from ptmapper import stock_sync_services as sync_services
from ptmapper.pt_opening import evidence_workbook, parse_pt_files
from ptmapper.soh_models import SohImport
from ptmapper.soh_parser import parse_soh
from sell.services.goods_stock import read_shelf

pytestmark = pytest.mark.django_db

PT_HEADER = ["SEASON", "BRAND", "COLOR", "GENDER", "SUB CATEGORY", "TYPE", "ITEM", "FIT", "SIZE", "BARCODE", "DESIGN",
             "HSN", "QTY", "MRP", "BASIC", "P RATE", "INPUT TAX", "OUTPUT TAX", "NAG", "MARGIN"]
OWNER = "syn.owner@synthetic.demo"


def pt_row(barcode: Any, qty: Any, mrp: Any = 1399, prate: Any = Decimal("1091.22"), hsn: Any = 61091000,
           brand: str = "STATUS QUO", size: str = "M") -> list[Any]:
    return ["SPRING SUMMER(May-26)", brand, "TEAL", "MALE", "CASUAL WEAR", "TOP WEAR", "T-SHIRT", "SLIM", size,
            barcode, "SQI-RN-26193", hsn, qty, mrp, Decimal("909.35"), prate, 5, 5, qty, 22]


def pt_book(sheets: dict[str, list[list[Any]]]) -> bytes:
    book = Workbook()
    first = book.active
    assert first is not None
    book.remove(first)
    master = book.create_sheet("Master Sheet")
    master.append(["SEASON", "BRAND", "COLOR", "SIZE"])
    master.append(["SPRING SUMMER(May-26)", "STATUS QUO", "TEAL", "M"])
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        sheet.append([sum(r[12] for r in rows if isinstance(r[12], int)), None, None])  # the work sheet's totals row
        sheet.append(PT_HEADER)
        for row in rows:
            sheet.append([float(v) if isinstance(v, Decimal) else v for v in row])
        sheet.append([None] * 7 + ["SLIM"])  # a fill-down stray cell, not a piece
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def soh_book(stock: dict[str, int]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Item Name", "Brand", "Size", "Barcode", "Category", "Season", "Tqty", "Mrp", "Rate", "Amount"])
    for code, qty in stock.items():
        sheet.append(["T-SHIRT", "STATUS QUO", "M", code, "CASUAL WEAR", "SPRING SUMMER(May-26)", qty, 1399, 900, 900 * qty])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# -- reader --------------------------------------------------------------------


def test_reader_adds_up_barcodes_across_sheets_and_files_and_lists_problems() -> None:
    one = pt_book({'"NARESH" Work Sheet': [pt_row(8905699672984, 2), pt_row(8905699673004, 1, prate=1500)],
                   '"ANKIT" Work Sheet': [pt_row(8905699672984, 1), pt_row(None, 3), pt_row(111, 1, hsn=6109)]})
    two = pt_book({"Sheet1": [pt_row("8905699672984", 4, mrp=1499), pt_row(222, 1, hsn="61X")]})
    metadata, rows, problems = parse_pt_files([("a.xlsx", one), ("b.xlsx", two)])
    by = {row.barcode: row for row in rows}
    assert by["8905699672984"].quantity == 7  # 2 + 1 + 4, numeric and text barcodes are one barcode
    assert by["8905699672984"].mrp == 149900 and by["8905699672984"].rate == 109122  # the last row wins ...
    assert metadata["pt_opening"]["conflicts"][0]["barcode"] == "8905699672984"  # ... and the MRP conflict is listed
    assert by["8905699672984"].source["hsn"] == "61091000" and by["8905699672984"].source["itemname"] == "T-SHIRT"
    assert problems == {"8905699673004": "PT: P RATE is above MRP", "222": "PT: HSN is not 4, 6 or 8 digits"}
    assert by["8905699673004"].source["problem"] == "PT: P RATE is above MRP"
    assert by["111"].source.get("problem") is None  # a 4-digit HSN is fine
    assert metadata["pt_opening"]["unreadable"] == [{"where": 'a.xlsx › "ANKIT" Work Sheet › row 4', "reason": "No barcode"}]
    assert [s["sheet"] for f in metadata["pt_opening"]["files"] for s in f["sheets"]] == [
        '"NARESH" Work Sheet', '"ANKIT" Work Sheet', "Sheet1"]  # Master Sheet has no PT header
    assert metadata["quantity"] == 7 + 1 + 1 + 1 and metadata["rate_meaning"] == "pt_prate"
    assert parse_soh(evidence_workbook(rows))[0]["quantity"] == metadata["quantity"]


def test_reader_refuses_a_file_with_no_pt_sheet() -> None:
    book = Workbook()
    assert book.active is not None
    book.active.append(["Barcode", "Tqty", "Mrp"])
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(Refusal) as refused:
        parse_pt_files([("soh.xlsx", buffer.getvalue())])
    assert refused.value.code == "PT_SOURCE_INVALID"


# -- the store -----------------------------------------------------------------


def stand_in_auto_decision(run: Any, request_id: Any) -> Any:
    """``decide_without_second_person`` until the approvals change lands: the maker asks, the
    test-store opening policy must carry the flag, and the tenant owner's decision is recorded."""
    from accounts.goods_demo import owner_user
    from accounts.test_store_loader import opening_auto_in_force
    from approvals.goods_models import ApprovalRequest
    from approvals.goods_services import decide
    from masters.goods_models import Tenant

    request = ApprovalRequest.objects.get(pk=request_id)
    if str(request.maker_id) != str(run.principal.human_id):
        raise Refusal("ACTION_DENIED", "Only the maker may ask for a decision with no second person.")
    if not opening_auto_in_force(run.tenant_id):
        raise Refusal("AUTO_APPROVAL_NOT_ALLOWED", "This opening needs a second person.")
    owner = access_for(owner_user(Tenant.objects.get(pk=run.tenant_id)).email)
    return act(owner, request.requested_action, lambda inner: decide(
        inner, access=owner, request_id=request_id, decision="approve", reviewed_hash=request.reviewed_hash,
        reason_code=None))


def use_auto_decisions(monkeypatch: pytest.MonkeyPatch) -> None:
    from approvals import goods_services

    if not hasattr(goods_services, "decide_without_second_person"):
        monkeypatch.setattr(goods_services, "decide_without_second_person", stand_in_auto_decision, raising=False)


def build_store(*, approval: bool) -> Store:
    Season.objects.get_or_create(historical_unknown=True,
                                 defaults={"code": "UNKNOWN-HIST", "name": "Unknown historical season", "status": "closed"})
    tenant = bootstrap_deployment(deployment_key=deployment_key(), code="KDPS-TEST", name="KDPS Lifestyle (test)",
                                  timezone_name="Asia/Kolkata", currency="INR", locale="en-IN", synthetic=True,
                                  admin_email=OWNER, admin_name="Synthetic Company Owner", admin_staff_code="OWN-01",
                                  admin_password="Bootstrap@123")
    with tenant_context(tenant.pk):
        seed_working_calendar(tenant, User.objects.get(email=OWNER).human_id)
    plan = Plan(store_code="WRR", store_name="Wardrobe Ratu Road (test)", warehouse_code="RNC-WH",
                warehouse_name="Warehouse Ranchi (test)", entity_code="syn-vas", entity_name="KDPS test entity",
                pan="AAAVT4444D", state_code="20", state_name="Jharkhand (test)", city="Ranchi", prefix="wrr",
                opening_approval=approval)
    _warehouse, built = Loader(tenant, plan, None, lambda _line: None).run()
    with tenant_context(tenant.pk):
        assert sum(read_shelf(built).quantities.values()) == 0  # sell-ready, no stock yet
        assert opening_auto_in_force(tenant.pk) is (not approval)
    assert isinstance(built, Store)
    return built


@pytest.fixture
def store(monkeypatch: pytest.MonkeyPatch) -> Store:
    """The alpha decision: a test store opens with no second person (setup ``opening.approval: false``)."""
    use_auto_decisions(monkeypatch)
    return build_store(approval=False)


@pytest.fixture
def approved_store(monkeypatch: pytest.MonkeyPatch) -> Store:
    """Openings that still need an owner's approval."""
    use_auto_decisions(monkeypatch)
    return build_store(approval=True)


def stage_pt(store: Store, files: list[tuple[str, bytes]], who: str = "wrr.store@synthetic.demo") -> SohImport:
    metadata, rows, _problems = parse_pt_files(files)
    driver = OpeningDriver(store.tenant, store, stager=lambda: access_for(who), preparer=lambda: access_for(who),
                           log=lambda _line: None)
    source = driver.stage(evidence_workbook(rows), "PT opening.xlsx", metadata, rows)
    assert isinstance(source, SohImport)
    return source


def uploader(source: SohImport) -> Any:
    return quick.job_access(source.tenant_id, source.uploaded_by_id)


def load(source: SohImport) -> dict[str, Any]:
    """What the Load button does, then the goods worker."""
    quick.request_load(uploader(source), source)
    drain(source.tenant_id)
    source.refresh_from_db()
    return quick.status(source)


def soh_sync(store: Store, stock: dict[str, int]) -> Any:
    data = soh_book(stock)
    metadata, parsed = parse_soh(data)
    who = access_for("wrr.store@synthetic.demo")
    evidence = stage_upload(who.principal(), command_id=uuid.uuid4(), data=data, filename="SOH.xlsx", kind="manifest",
                            scope={"scope_kind": "sites", "site_ids": [store.pk], "brand_ids": [],
                                   "sensitive_fields": ["cost", "financial"]},
                            expected_sha256=hashlib.sha256(data).hexdigest(), contains_fields=["cost", "financial"])
    sync = act(who, sync_services.ACTION, lambda run: sync_services.create_sync(
        run, site_id=store.pk, evidence=evidence, cutoff_at=timezone.now(), metadata=metadata, parsed=parsed))
    return act(who, sync_services.ACTION, lambda run: sync_services.apply(run, sync))


def test_pt_files_open_a_test_store_and_the_soh_brings_it_to_the_floor(store: Store) -> None:
    with tenant_context(store.tenant_id):
        with pytest.raises(Refusal) as refused:
            soh_sync(store, {"A1": 1})
        assert refused.value.code == "OPENING_REQUIRED"  # the SOH never stands in for the opening

        book = pt_book({"Sheet1": [pt_row("A1", 3), pt_row("A2", 2, size="L"), pt_row("BAD", 1, prate=2000)]})
        source = stage_pt(store, [("ratu road.xlsx", book)])
        state = load(source)
        assert state["state"] == "done", state
        assert state["loaded_qty"] == 5 and state["excluded_rows"] == 1
        assert read_shelf(store).quantities and sum(read_shelf(store).quantities.values()) == 5
        row = source.rows.get(barcode="A1")
        assert row.mapping["basic_paise"] == "109122" and row.mapping["hsn"] == "61091000"  # P RATE, the PT's HSN
        assert source.rows.get(barcode="BAD").exclusion_reason == "PT: P RATE is above MRP"
        assert source.state == "applied"

        with pytest.raises(Refusal) as again:
            quick.request_load(uploader(source), source)
        assert again.value.code == "STATE_CONFLICT"  # loaded once
        assert ApprovalDecision.objects.filter(request__site_id=store.pk).count() == 3  # source, manifest, PT

        sync = soh_sync(store, {"A1": 1, "A2": 4, "N1": 2})  # A1 sold 2, A2 two more found, N1 never in a PT
        assert sync.outcome["removed_qty"] == 2 and sync.outcome["added_qty"] == 2
        assert sync.outcome["new_item_barcodes"] == 1
        assert sum(read_shelf(store).quantities.values()) == 5

        top_up = sync.top_up_import
        assert top_up is not None
        state = load(top_up)  # the same one-click load takes the new items in, at the SOH Rate
        assert state["state"] == "done", state
        assert sum(read_shelf(store).quantities.values()) == 7
        assert top_up.rows.get(barcode="N1").mapping["basic_paise"] == "90000"


def test_the_uploader_loads_and_an_owner_only_while_openings_need_approval(
        store: Store, monkeypatch: pytest.MonkeyPatch) -> None:
    with tenant_context(store.tenant_id):
        source = stage_pt(store, [("pt.xlsx", pt_book({"Sheet1": [pt_row("A1", 1)]}))])
        for who in ("wrr.cashier@synthetic.demo", OWNER):  # not the uploader; no approval to give
            with pytest.raises(Refusal) as denied:
                quick.request_load(access_for(who), source)
            assert denied.value.code == "ACTION_DENIED", who
        assert quick.can_load(access_for("wrr.store@synthetic.demo"), source)

        monkeypatch.setattr(sync_services, "is_real_store", lambda _store: True)
        with pytest.raises(Refusal) as real:
            quick.request_load(uploader(source), source)
        assert real.value.code == "TEST_TENANT_ONLY"


def test_with_approvals_on_the_load_waits_for_the_owner_and_resumes(approved_store: Store) -> None:
    from approvals.goods_models import ApprovalRequest
    from approvals.goods_services import decide

    store = approved_store
    with tenant_context(store.tenant_id):
        source = stage_pt(store, [("pt.xlsx", pt_book({"Sheet1": [pt_row("A1", 2)]}))])
        state = load(source)
        assert state["state"] == "waiting_approval", state  # not a failure
        assert state["step"] == "Waiting for approval of the source"
        job = OutboxIntent.objects.get(kind=quick.JOB, subject_key=quick.subject(source.pk)).job
        assert job.state == JobState.State.CONFIRMED
        owner = access_for(OWNER)
        assert quick.can_load(owner, source)  # the owner may resume while openings need approval
        for _ in range(3):  # the source, the manifest, the opening PT
            pending = ApprovalRequest.objects.get(site_id=store.pk, state="pending")
            act(owner, pending.requested_action, lambda run, r=pending: decide(
                run, access=owner, request_id=r.pk, decision="approve", reviewed_hash=r.reviewed_hash, reason_code=None))
            quick.request_load(owner, source)
            drain(store.tenant_id)
            source.refresh_from_db()
        assert quick.status(source)["state"] == "done", quick.status(source)
        assert sum(read_shelf(store).quantities.values()) == 2


def test_a_second_pt_opening_is_refused_once_the_store_has_stock(store: Store) -> None:
    with tenant_context(store.tenant_id):
        assert load(stage_pt(store, [("first.xlsx", pt_book({"Sheet1": [pt_row("A1", 2)]}))]))["state"] == "done"
        second = load(stage_pt(store, [("second.xlsx", pt_book({"Sheet1": [pt_row("C1", 1)]}))]))
        assert second["state"] == "failed" and second["error"].startswith("SOH_RECONCILIATION_REQUIRED")
        assert sum(read_shelf(store).quantities.values()) == 2


def _alias_queries_per_batch(store: Store, monkeypatch: pytest.MonkeyPatch, barcodes: list[str]) -> int:
    from django.db import connection
    from django.test.utils import CaptureQueriesContext

    from ptmapper import soh_services

    counts: list[int] = []
    real = soh_services.apply_batch

    def counted(*args: Any, **kwargs: Any) -> Any:
        with CaptureQueriesContext(connection) as captured:
            batch = real(*args, **kwargs)
        counts.append(sum('FROM "masters_skualias"' in q["sql"] for q in captured.captured_queries))
        return batch

    monkeypatch.setattr(soh_services, "apply_batch", counted)
    book = pt_book({"Sheet1": [pt_row(code, 1, size=f"S{n}") for n, code in enumerate(barcodes)]})
    assert load(stage_pt(store, [("pt.xlsx", book)]))["state"] == "done"
    assert len(counts) == 1
    return counts[0]


def test_a_batch_looks_its_barcodes_up_together_not_one_by_one(store: Store, monkeypatch: pytest.MonkeyPatch) -> None:
    # One lookup per row was quadratic on the fresh server database (28 Oct 2026).
    with tenant_context(store.tenant_id):
        queries = _alias_queries_per_batch(store, monkeypatch, [f"Q{n}" for n in range(8)])
    assert queries <= 2  # was one per row (8 here, 5,000 in a full batch)


# -- resuming: every step is found from what is stored ----------------------------


def steps_until(source: SohImport, label: str) -> None:
    """Run the job's steps in-process up to and including the first one ending in ``label``."""
    driver = quick.driver_for(source, log=lambda _line: None)
    while True:
        step = driver.next_step(source.pk)
        assert step is not None, f"finished before {label!r}"
        if step.endswith(label):
            return


def finish(source: SohImport) -> dict[str, Any]:
    with quick.source_lock(source.pk) as taken:
        assert taken
        quick.drive(source.pk, budget=None)
    source.refresh_from_db()
    return quick.status(source)


def test_a_draft_opening_pt_is_submitted_on_resume(store: Store) -> None:
    from core.kernel_models import DocumentHead
    from ptmapper.goods_models import GoodsPt

    with tenant_context(store.tenant_id):
        source = stage_pt(store, [("pt.xlsx", pt_book({"Sheet1": [pt_row("A1", 2), pt_row("A2", 1, size="L")]}))])
        steps_until(source, "opening PT drafted")  # the 28 Oct server stopped here
        pt = GoodsPt.objects.get(manifest_version__manifest__site=store)
        assert DocumentHead.objects.get(document_id=pt.document_id).state == "draft"
        state = finish(source)
        assert state["state"] == "done", state
        assert DocumentHead.objects.get(document_id=pt.document_id).state == "official"
        assert sum(read_shelf(store).quantities.values()) == 3


def test_a_half_verified_source_finishes_its_verification(store: Store) -> None:
    from ptmapper import soh_services

    with tenant_context(store.tenant_id):
        source = stage_pt(store, [("pt.xlsx", pt_book({"Sheet1": [pt_row(f"V{n}", 1, size=f"S{n}") for n in range(3)]}))])
        steps_until(source, "Mappings, costs and HSN prepared")
        source.refresh_from_db()
        preparer = quick.preparer_for(source)
        act(preparer, soh_services.STAGE, lambda run: soh_services.verify_rows(
            run, preparer, SohImport.objects.get(pk=source.pk),
            [{"barcode": "V0", "observed_qty": 1, "observed_condition": "good", "reason": "partly counted"}],
            source.revision))
        state = finish(source)
        assert state["state"] == "done", state
        assert not source.rows.filter(quantity__gt=0, exclusion_reason="", verification={}).exists()
        assert sum(read_shelf(store).quantities.values()) == 3


def test_a_crash_before_acceptance_resumes_with_one_manifest_per_batch(
        store: Store, monkeypatch: pytest.MonkeyPatch) -> None:
    from stockledger import goods_acceptance

    real = goods_acceptance.scan
    crashes: list[int] = []

    def crash_once(*args: Any, **kwargs: Any) -> Any:
        if not crashes:
            crashes.append(1)
            raise RuntimeError("worker restarted")
        return real(*args, **kwargs)

    monkeypatch.setattr(goods_acceptance, "scan", crash_once)
    with tenant_context(store.tenant_id):
        source = stage_pt(store, [("pt.xlsx", pt_book({"Sheet1": [pt_row("A1", 2), pt_row("A2", 1, size="L")]}))])
        first = finish(source)
        assert first["state"] == "failed" and first["error"].startswith("RuntimeError"), first
        assert sum(read_shelf(store).quantities.values()) == 0
        state = load(source)
        assert state["state"] == "done", state
        assert source.batches.count() == 1
        assert sum(read_shelf(store).quantities.values()) == 3


def test_a_rolled_back_opening_pt_submit_is_retried_on_a_fresh_head(
        store: Store, monkeypatch: pytest.MonkeyPatch) -> None:
    """A serialization failure rolls back ``submit_opening`` and the command reruns the same handler.

    Tests run inside one transaction, where ``execute_command`` does not retry, so the rollback
    and rerun are made here; the handler must re-lock the head, not reuse the one the failed
    attempt moved to "submitted" (the 28 Oct ``STATE_CONFLICT``).
    """
    from django.db import transaction

    from accounts import test_store_loader as loader
    from core.commands import execute_command as real_execute

    class SerializationFailure(Exception):
        pass

    real_submit = loader._submit_opening
    attempts: list[str] = []

    def submit_then_fail_once(run: Any, document_id: Any, goods_pt_id: Any) -> Any:
        result = real_submit(run, document_id, goods_pt_id)
        attempts.append("submitted")
        if len(attempts) == 1:
            raise SerializationFailure("could not serialize access due to read/write dependencies")
        return result

    def execute_with_retry(principal: Any, spec: Any, handler: Any) -> Any:
        try:
            with transaction.atomic():
                return real_execute(principal, spec, handler)
        except SerializationFailure:
            return real_execute(principal, spec, handler)

    monkeypatch.setattr(loader, "_submit_opening", submit_then_fail_once)
    monkeypatch.setattr("accounts.test_store_loader.execute_command", execute_with_retry)
    with tenant_context(store.tenant_id):
        state = load(stage_pt(store, [("pt.xlsx", pt_book({"Sheet1": [pt_row("A1", 2)]}))]))
    assert attempts == ["submitted", "submitted"]
    assert state["state"] == "done", state


# -- the job --------------------------------------------------------------------


def test_a_second_job_for_the_same_source_is_a_no_op(store: Store) -> None:
    import psycopg
    from django.db import connection

    with tenant_context(store.tenant_id):
        source = stage_pt(store, [("pt.xlsx", pt_book({"Sheet1": [pt_row("A1", 2)]}))])
        intent = quick.request_load(uploader(source), source)
        params = {**connection.get_connection_params(), "dbname": connection.settings_dict["NAME"]}
        with psycopg.connect(**params) as other:  # another worker's session (advisory locks are per database)
            other.execute("SELECT pg_advisory_lock(%s)", [quick._lock_key(source.pk)])
            assert quick.running(source)
            outcome = quick.run_load(intent)
            assert (outcome.state, outcome.provider_ref) == ("confirmed", "already_loading")
            source.refresh_from_db()
            assert source.state == "uploaded" and not source.batches.exists()  # nothing done twice
        outcome = quick.run_load(intent)
        assert outcome.provider_ref == "done"


def test_the_job_queues_its_continuation_when_its_time_is_up(store: Store, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(quick, "BUDGET", 0.0)
    with tenant_context(store.tenant_id):
        source = stage_pt(store, [("pt.xlsx", pt_book({"Sheet1": [pt_row("A1", 2)]}))])
        intent = quick.request_load(uploader(source), source)
        outcome = quick.run_load(intent)
        assert outcome.state == "confirmed" and str(outcome.provider_ref).startswith("continues after")
        follow = OutboxIntent.objects.get(request_key=uuid.uuid5(intent.request_key, "continue"))
        assert follow.kind == quick.JOB and follow.job.state == JobState.State.PENDING
        with pytest.raises(Refusal) as busy:
            quick.request_load(uploader(source), source)  # a second Load while it runs
        assert busy.value.code == "STATE_CONFLICT"
        drain(store.tenant_id, rounds=60)
        source.refresh_from_db()
        jobs = [(j.intent.request_key == intent.request_key, j.state, j.attempts, j.error_code,
                 [d.diagnostic for d in j.intent.deliveries.all()])
                for j in JobState.objects.filter(intent__kind=quick.JOB).select_related("intent")]
        assert quick.status(source)["state"] == "done", jobs
        assert OutboxIntent.objects.filter(kind=quick.JOB, subject_key=quick.subject(source.pk)).count() > 3
        assert sum(read_shelf(store).quantities.values()) == 2


def test_uploading_pt_files_starts_the_opening(store: Store) -> None:
    from django.core.files.uploadedfile import SimpleUploadedFile
    from rest_framework.test import APIRequestFactory, force_authenticate

    from ptmapper.soh_views import SohImportPtUploadView

    book = pt_book({"Sheet1": [pt_row("A1", 2), pt_row("A2", 1, size="L")]})
    with tenant_context(store.tenant_id):
        who = access_for("wrr.store@synthetic.demo")
        request = APIRequestFactory().post("/api/goods-v1/ptmapper/soh-imports/pt-upload", {
            "files": [SimpleUploadedFile("ratu road.xlsx", book)], "site_id": store.pk,
            "command_id": str(uuid.uuid4()), "contract_version": "goods-v1"}, format="multipart")
        force_authenticate(request, who.user, who.session)
        response = SohImportPtUploadView.as_view()(request)
        assert response.status_code == 201, response.data
        assert response.data["data"]["quick_load"]["state"] == "running"
        assert "load" not in response.data["allowed_actions"]  # already queued
        drain(store.tenant_id)
        source = SohImport.objects.get(pk=response.data["id"])
        assert quick.status(source)["state"] == "done", quick.status(source)
        assert sum(read_shelf(store).quantities.values()) == 3


# -- open_test_store (the same load from a server shell) ------------------------


def test_open_test_store_stages_loads_and_never_loads_twice(store: Store, tmp_path: Any) -> None:
    from django.core.management import CommandError, call_command

    path = tmp_path / "pt.xlsx"
    path.write_bytes(pt_book({"Sheet1": [pt_row("A1", 3), pt_row("A2", 2, size="L")]}))
    out = io.StringIO()
    call_command("open_test_store", "--store", store.code, "--pt", str(path), "--uploader", "wrr.store@synthetic.demo",
                 "--foreground", stdout=out)
    assert "Done: 5 pieces on the floor" in out.getvalue() and "all steps" in out.getvalue()
    with pytest.raises(CommandError, match="0 open PT sources"):
        call_command("open_test_store", "--store", store.code, "--uploader", "wrr.store@synthetic.demo")
    with tenant_context(store.tenant_id):
        assert sum(read_shelf(store).quantities.values()) == 5
        assert SohImport.objects.filter(site=store).count() == 1


def test_open_test_store_queues_the_job_and_resumes_a_stopped_source(store: Store) -> None:
    from django.core.management import call_command

    with tenant_context(store.tenant_id):
        source = stage_pt(store, [("pt.xlsx", pt_book({"Sheet1": [pt_row("A1", 2)]}))])
        quick._record(source.pk, state="running", step="Batch 1 of 1")  # the 28 Oct thread that died
        source.refresh_from_db()
        assert quick.status(source)["state"] == "failed"  # no job and no shell holds it
    out = io.StringIO()
    call_command("open_test_store", "--store", store.code, "--uploader", "wrr.store@synthetic.demo", stdout=out)
    assert "Load queued" in out.getvalue()
    with tenant_context(store.tenant_id):
        drain(store.tenant_id)
        source.refresh_from_db()
        assert quick.status(source)["state"] == "done"
        assert sum(read_shelf(store).quantities.values()) == 2


def test_open_test_store_keeps_the_guards(store: Store, monkeypatch: pytest.MonkeyPatch, tmp_path: Any) -> None:
    from django.core.management import CommandError, call_command

    path = tmp_path / "pt.xlsx"
    path.write_bytes(pt_book({"Sheet1": [pt_row("B1", 1)]}))
    with pytest.raises(CommandError, match="No store NOPE"):
        call_command("open_test_store", "--store", "NOPE", "--uploader", "wrr.store@synthetic.demo")
    with pytest.raises(CommandError, match="No active login"):
        call_command("open_test_store", "--store", store.code, "--uploader", "nobody@synthetic.demo")
    with pytest.raises(CommandError, match="0 open PT sources"):
        call_command("open_test_store", "--store", store.code, "--uploader", "wrr.store@synthetic.demo")
    monkeypatch.setattr(sync_services, "is_real_store", lambda _store: True)
    with pytest.raises(CommandError, match="TEST_TENANT_ONLY"):
        call_command("open_test_store", "--store", store.code, "--pt", str(path), "--uploader", "wrr.store@synthetic.demo")
