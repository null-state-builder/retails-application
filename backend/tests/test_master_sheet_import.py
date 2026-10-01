"""Master sheet import: one uploaded KDPS master sheet configures every product list.

The uploader reviews and submits one package; a different owner approves it once,
and the vocabulary, brand and season masters and ITEM suggestion rules are written
through the product's own writers.
"""

from __future__ import annotations

import io
import uuid
from pathlib import Path
from typing import Any

import pytest
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from accounts.principal import AccessContext
from accounts.sessions import issue_session
from approvals.goods_models import ApprovalRequest
from core.canonical import sha256_hex
from core.refusals import Refusal
from core.tenancy import tenant_context
from files.goods_services import stage_upload
from masters import master_sheet as ms
from masters import master_sheet_services as services
from masters.goods_identity_models import SourceCrosswalk
from masters.goods_identity_services import ITEM_ISSUER, vocabulary
from masters.master_sheet_models import MasterSheetImport
from masters.master_sheet_views import MasterSheetImportUploadView
from masters.models import Brand, Season
from ptmapper.goods_rulebook import Rulebook
from tests.first_store_goods import approve, command, live_access
from tests.test_so03_denials import TenantWorld, _assign, _person
from tests.test_so03_denials import worlds as worlds

REAL_SHEET = (
    Path(__file__).resolve().parents[2] / "docs" / "data-from-kdps" / "KDPS PT FILE SHEET.xlsx"
)
HEADER = [
    "SEASON",
    "BRAND",
    "COLOR",
    "GENDER",
    "SUB CATEGORY",
    "TYPE",
    "ITEM",
    "FIT",
    "SIZE",
    "GST %",
    "SUB CATEGORY",
    "TYPE",
]
ROWS: list[list[Any]] = [
    ["SPRING SUMMER(Jan-27)", "PROOF BRAND", "BLACK", "MALE", "CASUAL WEAR", "TOP WEAR",
     "SHIRT", "SLIM", 28.0, 5, "FORMAL / CASUAL/ PARTY WEAR", "TOP WEAR"],
    ["AUTUMN WINTER(Jul-27)", "OTHER\xa0BRAND", "GREEN", "FEMALE", "FORMAL WEAR", "BOTTOM WEAR",
     "JEANS", "REGULAR", "M", 18, "CASUAL WEAR", "BOTTOM WEAR"],
    [None, None, "GREE", None, None, "LUGGAGE", "BACKPACK", None, "28", "TAX FREE",
     "ACCESSORIES", "LUGGAGE/ACCESSORIES"],
]


def sheet_bytes(rows: list[list[Any]] | None = None, header: list[str] | None = None) -> bytes:
    book = Workbook()
    ws = book.active
    assert ws is not None
    ws.title = "Master Sheet"
    ws.append(header or HEADER)
    for row in ROWS if rows is None else rows:
        ws.append(row)
    book.create_sheet('"OWNER" Work Sheet')
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------- reading


@pytest.mark.skipif(not REAL_SHEET.exists(), reason="KDPS workbook not in this checkout")
def test_real_kdps_master_sheet_reads_every_list() -> None:
    parsed = ms.read_master_sheet(REAL_SHEET.read_bytes())
    counts = {name: len(values) for name, values in parsed["columns"].items()}
    assert counts == {
        "season": 22, "colour": 23, "gender": 5, "sub_category": 9, "type": 7,
        "item": 98, "fit": 75, "size": 135, "brand": 592,
    }
    assert parsed["gst"] == ["5", "12", "18", "TAX FREE"]
    assert len(parsed["helper"]) == 98
    sizes = [e["text"] for e in parsed["columns"]["size"]]
    assert sizes[:3] == ["0", "1", "2"] and "28" in sizes


def test_cells_are_normalised_and_duplicates_kept_once() -> None:
    parsed = ms.read_master_sheet(sheet_bytes())
    sizes = [e["text"] for e in parsed["columns"]["size"]]
    assert sizes == ["28", "M"]  # 28.0 and "28" are one value
    assert [e["text"] for e in parsed["columns"]["brand"]] == ["PROOF BRAND", "OTHER BRAND"]
    codes = {(p["column"], p["code"]) for p in parsed["problems"]}
    assert ("size", ms.DUPLICATE_IN_SHEET) in codes
    assert ("brand", ms.WHITESPACE_FIXED) in codes
    assert parsed["gst"] == ["5", "18", "TAX FREE"]


def test_wrong_header_or_missing_sheet_is_refused() -> None:
    with pytest.raises(Refusal) as wrong:
        ms.read_master_sheet(sheet_bytes(header=["SEASON", "COLOR", *HEADER[2:]]))
    assert wrong.value.code == "MASTER_SHEET_INVALID"
    assert "column B should be BRAND" in wrong.value.message
    book = Workbook()
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(Refusal) as missing:
        ms.read_master_sheet(out.getvalue())
    assert missing.value.code == "MASTER_SHEET_INVALID"


def test_helper_options_name_every_real_value_in_order() -> None:
    subs = {"FORMAL WEAR", "CASUAL WEAR", "PARTY WEAR", "SPORTS WEAR", "INNERWEAR"}
    assert ms.helper_options("FORMAL / CASUAL/ PARTY WEAR", subs) == [
        "FORMAL WEAR", "CASUAL WEAR", "PARTY WEAR",
    ]
    assert ms.helper_options("LUGGAGE/ACCESSORIES", {"LUGGAGE", "ACCESSORIES"}) == [
        "LUGGAGE", "ACCESSORIES",
    ]
    assert ms.helper_options("NOTHING", subs) == []


def test_plan_flags_one_letter_twins_but_not_real_neighbours() -> None:
    parsed = ms.read_master_sheet(sheet_bytes())
    empty = {"values": {}, "version_ids": {}, "brands": [], "seasons": [], "item_rules": {}}
    plan = ms.plan_changes(parsed, empty, {})
    colour = next(d for d in plan["dimensions"] if d["dimension"] == "colour")
    assert [p["code"] for p in colour["problems"]] == [ms.LOOKS_LIKE]
    assert not ms._close("KURTA", "KURTI") and not ms._close("SHIRT", "T-SHIRT")
    rules = {r["item"]: r for r in plan["item_rules"]["new"]}
    assert rules["SHIRT"]["sub_category"]["options"] == ["FORMAL WEAR", "CASUAL WEAR"]
    assert rules["BACKPACK"]["type"]["value"] == "LUGGAGE"
    assert ms.season_code("AUTUMN WINTER(Jul-27)") == "aw-jul-27"


# ----------------------------------------------------------------------------- the package


@pytest.fixture
def world(worlds: tuple[TenantWorld, TenantWorld]) -> Any:
    with tenant_context(worlds[0].tenant.pk):
        yield worlds[0]


def _people(world: TenantWorld) -> tuple[AccessContext, AccessContext]:
    maker_user, maker = _person(world, "sheet-maker")
    checker_user, checker = _person(world, "sheet-checker")
    _assign(world, maker, "owner", all_sites=True, all_brands=True)
    _assign(world, checker, "owner", all_sites=True, all_brands=True)
    return live_access(maker_user), live_access(checker_user)


def _upload(access: AccessContext, data: bytes) -> MasterSheetImport:
    evidence = stage_upload(
        access.principal(), command_id=uuid.uuid4(), data=data, filename="kdps.xlsx",
        kind="other",
        scope={"scope_kind": "sites", "site_ids": [], "brand_ids": [], "sensitive_fields": []},
        expected_sha256=sha256_hex(data), contains_fields=[],
    )
    parsed = ms.read_master_sheet(data)
    return command(
        access, "masters.master_sheet.upload",
        lambda run: services.create_import(run, evidence=evidence, parsed=parsed),
    )


def _choose(access: AccessContext, source: MasterSheetImport, **choices: Any) -> MasterSheetImport:
    selections = {**source.selections, **choices}
    return command(
        access, "masters.master_sheet.selections",
        lambda run: services.update_selections(run, access, source, selections, source.revision),
    )


def _submit(access: AccessContext, source: MasterSheetImport) -> MasterSheetImport:
    return command(
        access, "masters.master_sheet.submit",
        lambda run: services.submit(run, access, source, source.revision, source.reviewed_hash),
    )


def _request(source: MasterSheetImport) -> ApprovalRequest:
    source.refresh_from_db()
    assert source.approval_request is not None
    return source.approval_request


def test_approved_package_configures_lists_brands_seasons_and_item_rules(world: TenantWorld) -> None:
    maker, checker = _people(world)
    source = _upload(maker, sheet_bytes())
    assert source.plan["change_count"] > 0
    assert source.plan["brands"]["to_create"][1]["name"] == "OTHER BRAND"
    source = _choose(
        maker, source,
        skip_values={"colour": ["GREE"]},
        rule_choices={"SHIRT": {"sub_category": "CASUAL WEAR"}},
    )
    source = _submit(maker, source)
    request = _request(source)

    with pytest.raises(Refusal) as own:
        approve(maker, request)
    assert own.value.code == "SELF_APPROVAL"

    approve(checker, request)
    source.refresh_from_db()
    assert source.state == "approved" and source.approved_by_id == checker.human_id

    now = timezone.now()
    lists = vocabulary(world.tenant.pk, now)
    assert [v.label for v in lists["colour"]] == ["BLACK", "GREEN"]
    assert [v.label for v in lists["size"]] == ["28", "M"]
    assert {v.label for v in lists["season"]} == {"SPRING SUMMER(Jan-27)", "AUTUMN WINTER(Jul-27)"}
    assert set(Brand.objects.filter(name__in=["PROOF BRAND", "OTHER BRAND"]).values_list(
        "code", flat=True)) == {"proof-brand", "other-brand"}
    assert Season.objects.filter(code="aw-jul-27", name="AUTUMN WINTER(Jul-27)").exists()

    book = Rulebook.load(world.tenant.pk, now)
    sub, _ = book.lookup("sub_category", "SHIRT", [ITEM_ISSUER])
    kind, _ = book.lookup("type", "backpack", [ITEM_ISSUER])
    assert sub is not None and sub.target.label == "CASUAL WEAR"
    assert kind is not None and kind.target.label == "LUGGAGE"

    # The same sheet again changes nothing and cannot be submitted: what was left out
    # stays left out, and the ITEM pick made last time stands.
    again = _upload(maker, sheet_bytes())
    assert again.selections == {"skip_values": {"colour": ["GREE"]}}
    assert again.pk != source.pk and again.plan["change_count"] == 0
    with pytest.raises(Refusal) as nothing:
        _submit(maker, again)
    assert nothing.value.code == "NOTHING_TO_CHANGE"


def test_values_missing_from_the_sheet_retire_only_when_ticked_and_can_return(world: TenantWorld) -> None:
    maker, checker = _people(world)
    first = _submit(maker, _choose(maker, _upload(maker, sheet_bytes()), skip_values={}))
    approve(checker, _request(first))

    shorter = [row[:] for row in ROWS[:1]]
    second = _upload(maker, sheet_bytes(shorter))
    size = next(d for d in second.plan["dimensions"] if d["dimension"] == "size")
    m = next(n for n in size["not_in_sheet"] if n["label"] == "M")
    assert m["retire"] is False
    second = _choose(maker, second, retire=[m["value_id"]])
    approve(checker, _request(_submit(maker, second)))
    sizes = {v.label: v.retired for v in vocabulary(world.tenant.pk, timezone.now())["size"]}
    assert sizes == {"28": False, "M": True}
    jeans = vocabulary(world.tenant.pk, timezone.now())["item"]
    assert {v.label for v in jeans if not v.retired} >= {"JEANS"}  # not ticked: kept

    third = _upload(maker, sheet_bytes())
    size = next(d for d in third.plan["dimensions"] if d["dimension"] == "size")
    assert [b["label"] for b in size["back_in_use"]] == ["M"]
    approve(checker, _request(_submit(maker, third)))
    sizes = {v.label: v.retired for v in vocabulary(world.tenant.pk, timezone.now())["size"]}
    assert sizes == {"28": False, "M": False}


def test_lists_changed_after_submit_refuse_approval_and_write_nothing(world: TenantWorld) -> None:
    maker, checker = _people(world)
    first = _upload(maker, sheet_bytes())
    extra = [None, None, None, None, None, None, None, None, "XL", None, None, None]
    second = _upload(maker, sheet_bytes([*ROWS, extra]))
    assert first.pk != second.pk
    approve(checker, _request(_submit(maker, first)))
    with pytest.raises(Refusal) as stale:
        _submit(maker, second)
    assert stale.value.code == "REVISION_SUPERSEDED"
    rules_before = SourceCrosswalk.objects.filter(tenant_id=world.tenant.pk).count()

    refreshed = command(
        maker, "masters.master_sheet.refresh",
        lambda run: services.refresh(run, maker, second, second.revision),
    )
    submitted = _submit(maker, refreshed)
    command(maker, "proof.brand", lambda run: Brand.objects.create(code="late", name="LATE"))
    with pytest.raises(Refusal) as late:
        approve(checker, _request(submitted))
    assert late.value.code == "REVISION_SUPERSEDED"
    assert SourceCrosswalk.objects.filter(tenant_id=world.tenant.pk).count() == rules_before


def test_editing_choices_after_submit_replaces_the_pending_approval(world: TenantWorld) -> None:
    maker, checker = _people(world)
    source = _submit(maker, _upload(maker, sheet_bytes()))
    request = _request(source)
    _choose(maker, source, skip_brands=["OTHER BRAND"])
    request.refresh_from_db()
    assert request.state == ApprovalRequest.State.SUPERSEDED
    with pytest.raises(Refusal):
        approve(checker, request)


def test_only_the_uploader_changes_the_package_and_store_staff_cannot_upload(world: TenantWorld) -> None:
    maker, checker = _people(world)
    source = _upload(maker, sheet_bytes())
    with pytest.raises(Refusal) as other:
        _choose(checker, source, skip_brands=["OTHER BRAND"])
    assert other.value.code == "ACTION_DENIED"

    staff_user, staff = _person(world, "sheet-staff")
    _assign(world, staff, "store_person", sites=(world.sites[0],), all_brands=True)
    session = issue_session(staff_user).session
    data = sheet_bytes()
    request = APIRequestFactory().post(
        "/api/goods-v1/masters/master-sheet-imports/upload",
        {
            "file": io.BytesIO(data),
            "command_id": str(uuid.uuid4()),
            "contract_version": "goods-v1",
            "expected_sha256": sha256_hex(data),
        },
        format="multipart",
    )
    force_authenticate(request, staff_user, session)
    response = MasterSheetImportUploadView.as_view()(request)
    assert response.status_code == 403


def test_another_tenants_package_is_not_found(worlds: tuple[TenantWorld, TenantWorld]) -> None:
    with tenant_context(worlds[0].tenant.pk):
        maker, _checker = _people(worlds[0])
        source = _upload(maker, sheet_bytes())
    with tenant_context(worlds[1].tenant.pk):
        stranger, _ = _people(worlds[1])
        with pytest.raises(Refusal) as missing:
            _choose(stranger, source, skip_brands=["OTHER BRAND"])
    assert missing.value.code == "NOT_FOUND"


@pytest.mark.skipif(not REAL_SHEET.exists(), reason="KDPS workbook not in this checkout")
def test_real_kdps_sheet_configures_a_new_tenant_in_one_approval(world: TenantWorld) -> None:
    maker, checker = _people(world)
    source = _upload(maker, REAL_SHEET.read_bytes())
    brands = len(source.plan["brands"]["to_create"])
    assert brands == 590  # 592 names; S F / TOM BOY twins are kept once
    approve(checker, _request(_submit(maker, source)))
    lists = vocabulary(world.tenant.pk, timezone.now())
    assert {name: len(values) for name, values in lists.items()} == {
        "season": 22, "colour": 23, "gender": 5, "sub_category": 9, "type": 7,
        "item": 98, "fit": 75, "size": 135,
    }
    assert Brand.objects.filter(code__in=[b["code"] for b in source.plan["brands"]["to_create"]]).count() == brands
    book = Rulebook.load(world.tenant.pk, timezone.now())
    rule, _ = book.lookup("type", "BACKPACK", [ITEM_ISSUER])
    assert rule is not None and rule.target.label == "LUGGAGE"
    assert _upload(maker, REAL_SHEET.read_bytes()).plan["change_count"] == 0


# ----------------------------------------------------------------------------- PT side


def test_generated_pt_file_round_trips_and_carries_the_sheet_dropdowns(world: TenantWorld) -> None:
    from decimal import Decimal

    from openpyxl import load_workbook

    from masters.master_sheet_template import template_bytes

    maker, checker = _people(world)
    source = _choose(maker, _upload(maker, sheet_bytes()), skip_values={"colour": ["GREE"]})
    approve(checker, _request(_submit(maker, source)))
    data = template_bytes(world.tenant.pk, timezone.now(), Decimal("1.1"))
    again = _upload(maker, data)
    assert again.plan["change_count"] == 0
    book = load_workbook(io.BytesIO(data))
    work = book["Work Sheet"]
    assert [c.value for c in work[2]][:9] == HEADER[:9]
    assert work["P3"].value == '=IF(O3<>"",O3*1.1,"")'
    ranges = {str(dv.sqref): dv.formula1 for dv in work.data_validations.dataValidation}
    assert ranges["C3:C1002"] == "'Master Sheet'!$C$2:$C$3"  # BLACK, GREEN
    master = book["Master Sheet"]
    shirt = next(r for r in master.iter_rows(min_row=2, values_only=True) if r[6] == "SHIRT")
    assert shirt[10:12] == ("FORMAL WEAR", "TOP WEAR")


def _work_book(*filled: str) -> bytes:
    from ptmapper.profiles import KDPS_COLUMNS

    book = Workbook()
    master = book.active
    assert master is not None
    master.title = "Master Sheet"
    master.append(HEADER)
    header = [*KDPS_COLUMNS[:20], None, *KDPS_COLUMNS[20:]]
    for name in ("OWNER", "NARESH", "ANKIT"):
        ws = book.create_sheet(f'"{name}" Work Sheet')
        ws.append([None] * 12 + ["=SUBTOTAL(9,M3:M9)"])
        ws.append(header)
        row = ["AUTUMN WINTER(Jul-27)", "PROOF BRAND", "BLACK", "MALE", "", "TOP WEAR", "SHIRT",
               "SLIM", 28.0, f"89{name[:2]}", "D-1", 62052000, 2, 1999, 1000, 1100]
        ws.append(row if name in filled else [None] * 16)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_one_filled_work_sheet_maps_every_kdps_column_from_its_own_cells() -> None:
    from masters.goods_identity_services import vocabulary_value_id
    from ptmapper import goods_mapper as gm
    from ptmapper.goods_rulebook import Rule, RuleValue

    parsed = ms.read_master_sheet(sheet_bytes())
    values = {
        dim: tuple(
            RuleValue(str(vocabulary_value_id(dim, e["text"])), dim, e["text"], e["text"])
            for e in parsed["columns"][dim]
        )
        for dim in ms.DIMENSIONS
    }
    casual = next(v for v in values["sub_category"] if v.label == "CASUAL WEAR")
    book = Rulebook(
        values=values,
        rules=(Rule("r1", "sub_category", ITEM_ISSUER, "SHIRT", casual),),
    )
    brand_file = gm.read_brand_file(_work_book("NARESH"), "kdps.xlsx")
    assert brand_file.sheet == '"NARESH" Work Sheet'
    assert brand_file.profile["code"] == "kdps_work_sheet"
    row = gm.map_records(brand_file, book).rows[0]
    cells = {name: (cell.value, cell.origin) for name, cell in row.cells.items()}
    assert cells["ITEM"] == ("SHIRT", "file")
    assert cells["TYPE"] == ("TOP WEAR", "file")
    assert cells["SIZE"] == ("28", "file")
    assert cells["SUB CATEGORY"] == ("CASUAL WEAR", "rule")  # blank cell: ITEM suggestion
    assert cells["SUGGESTED SUB CATEGORY"] == ("CASUAL WEAR", "rule")
    assert cells["DESIGN"] == ("D-1", "file") and cells["P RATE"] == ("1100", "file")

    with pytest.raises(gm.UnsupportedFormat) as two:
        gm.read_brand_file(_work_book("NARESH", "ANKIT"), "kdps.xlsx")
    assert "2 filled work sheets" in str(two.value)


def test_import_view_names_the_uploader(world: TenantWorld) -> None:
    from masters.master_sheet_views import import_dto

    maker, _checker = _people(world)
    source = _upload(maker, sheet_bytes())
    dto = import_dto(maker, source)
    assert dto["data"]["uploaded_by"] and dto["data"]["approved_by"] is None
