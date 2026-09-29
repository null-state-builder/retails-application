"""Each brand's Sale and SOH reports, in the brand's own layout (store operations ST-BRD-2).

Ticket 29. A report is one brand at one store for one month:

* **Sale** - every goods line of the brand's accepted bills at the store in the
  month, a piece given back negative (``reporting.brand_sale_facts``);
* **SOH** - the brand's good, accepted pieces standing at the store at the
  month's last stock snapshot (or the newest, in the running month), one row
  per barcode and MRP (``InventoryItemFact``, ticket 43's snapshot).

Both read only the reporting copies, never the billing or goods tables, so
making one never slows a bill. Each is laid out by the brand's saved layout
(``reporting.brand_layouts``) and written as ``.xlsx`` or ``.csv``. An ``.xlsx``
file carries a second sheet stating its basis, as-of time and missing data, so
the brand's own sheet stays exactly in its format; the page states the same.

Reports are made **on demand** (``make``, the Brand Reports page) and **monthly**
(``make_month``, from the worker's clock once the copies have passed the month's
end, and from ``manage.py make_brand_reports``). A monthly file is kept exactly
as made (``BrandReportFile``); each monthly run is one command, so the audit log
records what it made. An on-demand file is not kept; taking it is a
``reports.export`` audit entry, as every report export is.
"""

from __future__ import annotations

import csv
import hashlib
import io
import logging
import time
import uuid
from collections import defaultdict
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from core.commands import (
    CommandResult,
    CommandRun,
    CommandSpec,
    LockRank,
    Principal,
    execute_command,
)
from core.goods_money import amount_text
from core.refusals import Refusal
from masters.models import Brand, Store
from masters.store_feature_registry import BRAND_REPORTS
from masters.store_features import feature, switch_states
from offers.resolution import normalise
from reporting import brand_layouts
from reporting.base import XLSX, Missing, freshness
from reporting.brand_layouts import FIXED, SALE, SOH
from reporting.brand_sale_facts import KEY as SALES_KEY
from reporting.inventory_facts import KEY as STOCK_KEY
from reporting.margin_share import month_bounds
from reporting.models import (
    BrandReportFile,
    BrandReportKind,
    BrandReportLayout,
    BrandReportRun,
    BrandSaleLineFact,
    InventoryItemFact,
    InventorySnapshot,
    ReportRefresh,
)

logger = logging.getLogger(__name__)

CSV = "text/csv; charset=utf-8"
#: The audit action one store's monthly run is recorded under.
MONTHLY_ACTION = "reports.brand_reports.monthly"
KIND_WORDS = dict(BrandReportKind.choices)

SALE_BASIS = [
    "Every goods line of the brand's accepted, numbered bills at this store in the period; "
    "cancelled bills are left out.",
    "A piece given back (an exchange, or an old return) is negative, at the MRP of the bill "
    "it came off and what the customer was paid back.",
    "Price is MRP times pieces; Total is what the customer paid, GST inside; Discount is the "
    "difference.",
    "A line is the brand's when the brand on the bill matches the brand's name or short "
    "code, ignoring case, spaces and punctuation.",
    "Alteration charges are a service, not the brand's goods, and are left out.",
]
SOH_BASIS = [
    "The brand's good, accepted pieces standing at this store, as the stock snapshot taken "
    "at the end of the day named below found them; one row per barcode and MRP.",
    "Pieces in transit, awaiting acceptance, damaged or in quarantine are not stock on hand.",
    "A past month shows its last snapshot (its closing stock); the running month its newest.",
]


# -- the period --------------------------------------------------------------------


@dataclass(frozen=True)
class Period:
    label: str
    first: date
    last: date
    #: The last day the report can cover: the month's last day, or today while it runs.
    upto: date

    @property
    def running(self) -> bool:
        return self.upto < self.last


def period(raw: str | None) -> Period:
    """``YYYY-MM`` (default: this month); a month not yet begun is refused."""
    label, first, last = month_bounds(raw)
    today = timezone.localdate()
    if first > today:
        raise Refusal("INVALID_REQUEST", "That month has not started yet.")
    return Period(label=label, first=first, last=last, upto=min(last, today))


# -- which lines are a brand's --------------------------------------------------------


def keys_of(brand: Brand) -> set[str]:
    """What a bill's brand reduces to when it names this brand: its name or its code."""
    return {key for key in (normalise(brand.name), normalise(brand.code)) if key}


def brand_index(brands: Iterable[Brand]) -> dict[str, list[Brand]]:
    index: dict[str, list[Brand]] = defaultdict(list)
    for brand in brands:
        for key in keys_of(brand):
            index[key].append(brand)
    return index


# -- rows --------------------------------------------------------------------------


def _pct(disc: int, gross: int) -> Decimal | None:
    if not gross:
        return None
    exact = (Decimal(disc) * 100 / Decimal(gross)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    # A return's nought is nought, never "-0".
    return exact if exact else Decimal("0.00")


def sale_rows(facts: Iterable[BrandSaleLineFact], brand: Brand) -> list[dict[str, Any]]:
    return [
        {
            "day": f.day,
            "bill_number": f.doc_number,
            "item": f.item,
            "brand_on_bill": f.brand,
            "brand_name": brand.name,
            "brand_code": brand.code.upper(),
            "size": f.size,
            "design": f.design,
            "barcode": f.barcode,
            "colour": f.color,
            "season": f.season,
            "qty": f.qty,
            "mrp": int(f.mrp_paise) if f.mrp_paise is not None else None,
            "gross": int(f.gross_paise),
            "disc_pct": _pct(int(f.disc_paise), int(f.gross_paise)),
            "disc": int(f.disc_paise),
            "value": int(f.value_paise),
            "gst": int(f.gst_paise),
        }
        for f in facts
    ]


def soh_rows(
    facts: Iterable[InventoryItemFact], brand: Brand, when: Period, stock_day: date
) -> list[dict[str, Any]]:
    return [
        {
            "month_start": when.first,
            "stock_day": stock_day,
            "item": f.item,
            "brand_name": brand.name,
            "brand_code": brand.code.upper(),
            "size": f.size,
            "design": f.design,
            "barcode": f.barcode,
            "colour": f.color,
            "season": f.season,
            "qty": f.pieces,
            "mrp": int(f.mrp_paise) if f.mrp_paise is not None else None,
            "mrp_value": int(f.mrp_paise) * f.pieces if f.mrp_paise is not None else None,
        }
        for f in facts
    ]


# -- what one store and month holds -------------------------------------------------------


@dataclass
class StoreMonth:
    """Everything the copies hold for one store in one month, read once."""

    store: Store
    when: Period
    sales: dict[str, list[BrandSaleLineFact]]
    stock: dict[str, list[InventoryItemFact]]
    #: The day of the snapshot the SOH shows; None when none was taken in the month.
    stock_day: date | None
    goods_records: bool
    #: False when that snapshot was taken before stock by barcode was kept.
    items_kept: bool = True


def read_store_month(store: Store, when: Period, keys: set[str] | None = None) -> StoreMonth:
    lines = BrandSaleLineFact.objects.filter(
        store=store, day__gte=when.first, day__lte=when.upto
    ).order_by("day", "doc_number", "doc_id", "line_no", "id")
    if keys is not None:
        lines = lines.filter(brand_key__in=sorted(keys))
    sales: dict[str, list[BrandSaleLineFact]] = defaultdict(list)
    for line in lines:
        sales[line.brand_key].append(line)
    snapshot = (
        InventorySnapshot.objects.filter(store=store, day__gte=when.first, day__lte=when.upto)
        .order_by("-day")
        .first()
    )
    stock: dict[str, list[InventoryItemFact]] = defaultdict(list)
    if snapshot is not None:
        items = InventoryItemFact.objects.filter(store=store, day=snapshot.day).order_by(
            "item", "design", "size", "barcode", "mrp_paise", "id"
        )
        if keys is not None:
            items = items.filter(brand_key__in=sorted(keys))
        for item in items:
            stock[item.brand_key].append(item)
    return StoreMonth(
        store=store,
        when=when,
        sales=dict(sales),
        stock=dict(stock),
        stock_day=snapshot.day if snapshot else None,
        goods_records=snapshot.goods_records if snapshot else True,
        items_kept=snapshot.items_kept if snapshot else True,
    )


def _of(groups: dict[str, list[Any]], brand: Brand) -> list[Any]:
    found: list[Any] = []
    for key in sorted(keys_of(brand)):
        found += groups.get(key, [])
    return found


# -- one report -------------------------------------------------------------------------


@dataclass
class Made:
    """One report file and everything it states."""

    store: Store
    brand: Brand
    kind: str
    when: Period
    layout: BrandReportLayout
    content: bytes
    content_type: str
    file_name: str
    rows: int
    as_of: datetime | None
    basis: list[str]
    missing: list[dict[str, str]] = field(default_factory=list)

    def about(self) -> dict[str, Any]:
        return {
            "store": self.store.code,
            "brand": self.brand.code,
            "kind": self.kind,
            "month": self.when.label,
            "date_from": self.when.first.isoformat(),
            "date_to": self.when.upto.isoformat(),
            "as_of": self.as_of.isoformat() if self.as_of else None,
            "basis": self.basis,
            "missing": self.missing,
            "layout": {
                "id": self.layout.pk,
                "name": self.layout.name,
                "revision": self.layout.revision,
            },
            "rows": self.rows,
        }


def build(
    store_month: StoreMonth,
    brand: Brand,
    kind: str,
    *,
    as_of: datetime | None,
    missing: Missing,
    layout: BrandReportLayout | None = None,
) -> Made:
    """One brand's report of ``kind`` from what ``store_month`` holds."""
    when, store = store_month.when, store_month.store
    chosen = layout or brand_layouts.effective(brand, kind)
    if kind == SALE:
        rows = sale_rows(_of(store_month.sales, brand), brand)
        basis = [*SALE_BASIS]
        if when.running:
            basis.append(f"The month is still running: bills up to {when.upto:%d %b %Y}.")
        if any(row["mrp"] is None for row in rows):
            missing.add(
                "NO_ORIGINAL",
                "Some pieces were given back on an old return whose bill is not held: their "
                "MRP is empty and their price is what was paid back.",
            )
    else:
        stock_day = store_month.stock_day
        if stock_day is None:
            missing.add(
                "NO_SNAPSHOT",
                f"No stock snapshot was taken at {store_month.store.name} in "
                f"{when.first:%B %Y}, so its stock is not known and nothing is listed.",
            )
            rows = []
        else:
            rows = soh_rows(_of(store_month.stock, brand), brand, when, stock_day)
        if not store_month.items_kept:
            missing.add(
                "NO_ITEMS",
                f"The stock snapshot of {stock_day:%d %b %Y} was taken before stock was kept "
                "by barcode, so nothing is listed.",
            )
        if not store_month.goods_records:
            missing.add(
                "NOT_GOODS_RECORDS",
                f"{store_month.store.name} keeps its stock in the older system, which records "
                "no pieces by barcode, so nothing is listed.",
            )
        basis = [*SOH_BASIS]
        if stock_day is not None:
            basis.append(f"Stock as at the end of {stock_day:%d %b %Y}.")
        if any(row["mrp"] is None for row in rows):
            missing.add(
                "NO_MRP", "Some pieces have no MRP on their receipt: their MRP cells are empty."
            )
    values = {
        "company": store_month.store.gstin.legal_entity.name,
        "brand": brand.name,
        "brand_code": brand.code.upper(),
        "store": store_month.store.name,
        "store_code": store_month.store.code,
        "from": when.first.strftime("%d-%m-%Y"),
        "to": when.upto.strftime("%d-%m-%Y"),
        "month": when.first.strftime("%B %Y"),
    }
    layout_data = brand_layouts.clean(kind, chosen.layout)
    about = [
        ("Report", f"{brand.name} {KIND_WORDS[kind]} - {store.name} ({store.code})"),
        ("Period", f"{when.first:%d %b %Y} to {when.upto:%d %b %Y}"),
        (
            "As of",
            timezone.localtime(as_of).strftime("%Y-%m-%d %H:%M") if as_of else "Not built yet",
        ),
        (
            "Layout",
            f"{chosen.name}, "
            + ("built in" if chosen.pk is None else f"revision {chosen.revision}"),
        ),
        *[("Basis", line) for line in basis],
        *([("Missing data", item["text"]) for item in missing.items] or [("Missing data", "None")]),
    ]
    stem = f"{brand.code.upper()}-{KIND_WORDS[kind]}-{store_month.store.code}-{when.label}"
    if layout_data["file_type"] == "csv":
        content, content_type, ext = render_csv(kind, layout_data, rows, values), CSV, "csv"
    else:
        content, content_type, ext = (
            render_xlsx(kind, layout_data, rows, values, about),
            XLSX,
            "xlsx",
        )
    return Made(
        store=store_month.store,
        brand=brand,
        kind=kind,
        when=when,
        layout=chosen,
        content=content,
        content_type=content_type,
        file_name=f"{stem}.{ext}",
        rows=len(rows),
        as_of=as_of,
        basis=basis,
        missing=missing.items,
    )


# -- writing a layout ---------------------------------------------------------------------


def _totals(columns: Sequence[dict[str, Any]], rows: Sequence[dict[str, Any]]) -> dict[str, int]:
    out: dict[str, int] = {}
    for index, column in enumerate(columns):
        if column["total"]:
            out[str(index)] = sum(int(row.get(column["field"]) or 0) for row in rows)
    return out


def _value(column: dict[str, Any], row: dict[str, Any]) -> Any:
    if column["field"] == FIXED.key:
        return column.get("value", "")
    return row.get(column["field"])


def _kind(column: dict[str, Any], kind: str) -> str:
    found = brand_layouts.field_of(kind, column["field"])
    return found.kind if found else brand_layouts.TEXT


def _rupees(paise: int) -> str:
    return ("-" if paise < 0 else "") + amount_text(abs(paise))


def _plain(value: Any, column: dict[str, Any], field_kind: str) -> str:
    """One cell as text, for a ``.csv`` file."""
    if value is None:
        return ""
    fmt = column["format"]
    if field_kind == brand_layouts.DATE:
        day: date = value
        return f"{day.day}-{day:%b-%y}" if fmt == "d-mmm-yy" else day.strftime("%d-%m-%Y")
    if field_kind == brand_layouts.MONEY:
        rupees = Decimal(int(value)) / 100
        if fmt == "0":
            return str(rupees.quantize(Decimal(1), rounding=ROUND_HALF_UP))
        return _rupees(int(value))
    if field_kind == brand_layouts.PERCENT:
        places = Decimal(1) if fmt == "0" else Decimal("0.01")
        return str(Decimal(value).quantize(places, rounding=ROUND_HALF_UP))
    return str(value)


def render_csv(
    kind: str, layout: dict[str, Any], rows: Sequence[dict[str, Any]], values: dict[str, str]
) -> bytes:
    columns = layout["columns"]
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    for line in layout["title_lines"]:
        writer.writerow([brand_layouts.fill(line["text"], values)])
    writer.writerow([column["header"] for column in columns])
    kinds = [_kind(column, kind) for column in columns]
    for row in rows:
        writer.writerow([_plain(_value(c, row), c, k) for c, k in zip(columns, kinds, strict=True)])
    if any(column["total"] for column in columns):
        # The sum of the values as printed, so a column of whole rupees adds up.
        writer.writerow(
            [
                _plain_total([_plain(_value(c, row), c, k) for row in rows]) if c["total"] else ""
                for c, k in zip(columns, kinds, strict=True)
            ]
        )
    return out.getvalue().encode("utf-8-sig")


def _plain_total(shown: list[str]) -> str:
    total = sum((Decimal(text) for text in shown if text), Decimal(0))
    return str(total)


THIN = Side(style="thin", color="000000")
MEDIUM = Side(style="medium", color="000000")


def render_xlsx(
    kind: str,
    layout: dict[str, Any],
    rows: Sequence[dict[str, Any]],
    values: dict[str, str],
    about: Sequence[tuple[str, str]],
) -> bytes:
    """The brand's sheet exactly as laid out, and a second sheet saying what it is."""
    columns = layout["columns"]
    kinds = [_kind(column, kind) for column in columns]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = layout["sheet_name"] or KIND_WORDS[kind]
    number = _title_block(sheet, layout, values)
    number = _header_row(sheet, layout, number + 1)
    body = Font(name=layout["font"], size=10)
    grid = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    for row in rows:
        number += 1
        for index, (column, field_kind) in enumerate(zip(columns, kinds, strict=True), start=1):
            cell = sheet.cell(row=number, column=index)
            _write(cell, _value(column, row), column, field_kind)
            cell.font, cell.alignment, cell.border = body, CENTRED, grid
    totals = _totals(columns, rows)
    if totals:
        number += 1
        bold = Font(name=layout["font"], bold=True, size=10)
        for index, (column, field_kind) in enumerate(zip(columns, kinds, strict=True), start=1):
            cell = sheet.cell(row=number, column=index)
            if str(index - 1) in totals:
                _write(cell, totals[str(index - 1)], column, field_kind)
            cell.font, cell.alignment, cell.border = bold, CENTRED, Border(top=THIN)
    for index, column in enumerate(columns, start=1):
        seen = [len(str(_value(column, row) or "")) + 2 for row in rows[:200]]
        widest = max([len(column["header"]) + 2, 10, *seen])
        sheet.column_dimensions[get_column_letter(index)].width = min(widest, 40)
    _about_sheet(workbook, about)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


CENTRED = Alignment(horizontal="center", vertical="center")


def _title_block(sheet: Any, layout: dict[str, Any], values: dict[str, str]) -> int:
    """The title lines, each merged across the table, boxed; the last row used."""
    width = len(layout["columns"])
    titles = layout["title_lines"]
    for number, line in enumerate(titles, start=1):
        cell = sheet.cell(row=number, column=1, value=brand_layouts.fill(line["text"], values))
        cell.font = Font(name=layout["font"], bold=True, size=line["size"])
        cell.alignment = CENTRED
        if width > 1:
            sheet.merge_cells(start_row=number, start_column=1, end_row=number, end_column=width)
    last = len(titles)
    for row_number in range(1, last + 1):
        for column_number in range(1, width + 1):
            sheet.cell(row=row_number, column=column_number).border = Border(
                left=MEDIUM if column_number == 1 else None,
                right=MEDIUM if column_number == width else None,
                top=MEDIUM if row_number == 1 else None,
                bottom=MEDIUM if row_number == last else None,
            )
    return last


def _header_row(sheet: Any, layout: dict[str, Any], number: int) -> int:
    """The headers: a thin grid with a heavy top, or one heavy box round the row."""
    font = Font(name=layout["font"], bold=True, underline="single", size=layout["header_size"])
    width = len(layout["columns"])
    boxed = layout.get("header_border") == "box"
    for index, column in enumerate(layout["columns"], start=1):
        cell = sheet.cell(row=number, column=index, value=column["header"])
        if boxed:
            border = Border(
                left=MEDIUM if index == 1 else None,
                right=MEDIUM if index == width else None,
                top=MEDIUM,
                bottom=MEDIUM,
            )
        else:
            border = Border(left=THIN, right=THIN, top=MEDIUM, bottom=THIN)
        cell.font, cell.alignment, cell.border = font, CENTRED, border
    return number


def _about_sheet(workbook: Any, about: Sequence[tuple[str, str]]) -> None:
    notes = workbook.create_sheet("About this report")
    for row_number, (label, text) in enumerate(about, start=1):
        notes.cell(row=row_number, column=1, value=label).font = Font(bold=True)
        notes.cell(row=row_number, column=2, value=text)
    notes.column_dimensions["A"].width = 16
    notes.column_dimensions["B"].width = 110


def _write(cell: Any, value: Any, column: dict[str, Any], field_kind: str) -> None:
    if value is None:
        return
    fmt = column["format"]
    if field_kind == brand_layouts.MONEY:
        cell.value = 0  # a numeric cell ...
        cell._value = _rupees(int(value))  # ... whose stored text is exact
        cell.number_format = fmt
    elif field_kind == brand_layouts.DATE:
        cell.value = value
        cell.number_format = fmt
    elif field_kind in (brand_layouts.COUNT, brand_layouts.PERCENT):
        cell.value = value if field_kind == brand_layouts.COUNT else float(value)
        cell.number_format = fmt
    else:
        cell.value = str(value)
        cell.number_format = "@"


# -- as-of ---------------------------------------------------------------------------------


def as_of_for(kind: str, missing: Missing) -> datetime | None:
    return freshness(SALES_KEY if kind == SALE else STOCK_KEY, missing)


# -- on demand -------------------------------------------------------------------------------


def make(store: Store, brand: Brand, kind: str, when: Period) -> Made:
    """One report now, from the copies as they stand. Nothing is kept."""
    if kind not in brand_layouts.KINDS:
        raise Refusal("INVALID_REQUEST", "kind must be sale or soh.")
    missing = Missing()
    as_of = as_of_for(kind, missing)
    store_month = read_store_month(store, when, keys_of(brand))
    return build(store_month, brand, kind, as_of=as_of, missing=missing)


# -- monthly -----------------------------------------------------------------------------


def month_start_moment(day: date) -> datetime:
    """Midnight at the start of ``day``, India time."""
    return timezone.make_aware(datetime.combine(day, datetime.min.time()))


def copies_past(moment: datetime) -> bool:
    """Have both copies been brought up to date since ``moment``?"""
    for key in (SALES_KEY, STOCK_KEY):
        state = ReportRefresh.objects.filter(key=key).first()
        if state is None or state.as_of is None or state.as_of < moment:
            return False
    return True


def brands_active(store_month: StoreMonth, brands: Sequence[Brand]) -> list[Brand]:
    """The brands with a bill line or stock at the store in the month, in name order."""
    held = set(store_month.sales) | set(store_month.stock)
    return [brand for brand in brands if keys_of(brand) & held]


@dataclass(frozen=True)
class MonthResult:
    store: Store
    month: date
    made: bool
    files: int = 0
    took_ms: int = 0


def make_month(store: Store, month: date) -> MonthResult:
    """The month's reports for every brand active at ``store``, kept; once per store and month.

    One command: its audit record lists every file made. A second call for the
    same store and month makes nothing and says so.
    """
    clock = time.monotonic()
    first = month.replace(day=1)
    when = Period(
        label=first.strftime("%Y-%m"),
        first=first,
        last=month_bounds(first.strftime("%Y-%m"))[2],
        upto=month_bounds(first.strftime("%Y-%m"))[2],
    )
    if BrandReportRun.objects.filter(store=store, month=first).exists():
        return MonthResult(store=store, month=first, made=False)
    brands = list(Brand.objects.filter(tenant_id=store.tenant_id, is_active=True).order_by("name"))
    store_month = read_store_month(store, when)
    outcome: dict[str, Any] = {}

    def handler(run: CommandRun) -> CommandResult:
        # The worker and the command can both reach here: one store's month at a time.
        run.advisory_lock(LockRank.CHAIN, [f"brand-reports:{store.pk}:{when.label}"])
        if BrandReportRun.objects.filter(store=store, month=first).exists():
            raise Refusal("ALREADY_MADE", "This month's brand reports are already made here.")
        made_at = timezone.now()
        record = BrandReportRun.objects.create(store=store, month=first, made_at=made_at)
        listed = []
        stock_known = store_month.stock_day is not None
        for brand in brands_active(store_month, brands):
            for kind in brand_layouts.KINDS:
                if kind == SOH and not stock_known:
                    continue
                missing = Missing()
                as_of = as_of_for(kind, missing)
                made = build(store_month, brand, kind, as_of=as_of, missing=missing)
                BrandReportFile.objects.create(
                    run=record,
                    store=store,
                    brand=brand,
                    kind=kind,
                    month=first,
                    file_name=made.file_name,
                    content_type=made.content_type,
                    content=made.content,
                    size=len(made.content),
                    rows=made.rows,
                    layout_id=made.layout.pk,
                    layout_revision=made.layout.revision,
                    as_of=made.as_of,
                    about=made.about(),
                    made_at=made_at,
                )
                listed.append(
                    {
                        "brand": brand.code,
                        "kind": kind,
                        "file_name": made.file_name,
                        "rows": made.rows,
                        "sha256": hashlib.sha256(made.content).hexdigest(),
                        "layout_id": made.layout.pk,
                        "layout_revision": made.layout.revision,
                    }
                )
        record.files = len(listed)
        record.took_ms = int((time.monotonic() - clock) * 1000)
        record.save(update_fields=["files", "took_ms"])
        run.audit_before = None
        run.audit_after = {
            "store": store.code,
            "month": when.label,
            "files": listed,
            "stock_snapshot": (
                store_month.stock_day.isoformat() if store_month.stock_day else None
            ),
        }
        outcome["files"] = len(listed)
        outcome["took_ms"] = record.took_ms
        return CommandResult(resource_type="brand_report_run", resource_id=str(record.pk))

    try:
        execute_command(
            Principal(tenant_id=store.tenant_id, service_code="reports"),
            CommandSpec(
                MONTHLY_ACTION,
                uuid.uuid4(),
                {"store": store.code, "month": when.label},
                subject_key=f"brand_reports:{store.code}:{when.label}",
                site_id=store.pk,
            ),
            handler,
        )
    except Refusal as refusal:
        if refusal.code == "ALREADY_MADE":
            return MonthResult(store=store, month=first, made=False)
        raise
    return MonthResult(
        store=store,
        month=first,
        made=True,
        files=int(outcome["files"]),
        took_ms=int(outcome["took_ms"]),
    )


def stores_switched_on(tenant_id: Any) -> list[Store]:
    stores = list(
        Store.objects.filter(
            tenant_id=tenant_id, is_active=True, store_type=Store.StoreType.STORE
        ).order_by("code")
    )
    on = {
        state.site_id for state in switch_states(stores, [feature(BRAND_REPORTS)]) if state.enabled
    }
    return [store for store in stores if store.pk in on]


def scheduled_monthly(tenant_id: uuid.UUID, today: date | None = None) -> list[MonthResult]:
    """The worker's call each tick: last month's reports, once the copies have passed
    its end, at every store where the feature is on and they are not made yet."""
    day = today or timezone.localdate()
    this_month = day.replace(day=1)
    last_month = (this_month - timedelta(days=1)).replace(day=1)
    stores = stores_switched_on(tenant_id)
    made = set(
        BrandReportRun.objects.filter(month=last_month, store__in=stores).values_list(
            "store_id", flat=True
        )
    )
    waiting = [store for store in stores if store.pk not in made]
    if not waiting or not copies_past(month_start_moment(this_month)):
        return []
    results = []
    for store in waiting:
        # One store's failure is logged, and the others still get their files.
        try:
            results.append(make_month(store, last_month))
        except Exception:  # noqa: BLE001 - logged; this store is tried again next tick
            logger.exception("brand reports: %s %s failed", store.code, last_month)
    return results
